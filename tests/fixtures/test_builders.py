"""Tests for the generated OOXML fixtures.

Each fixture must open with python-docx, rebuild byte-for-byte identically, and carry
the OOXML constructs it is named after. The markers below double as documentation of
what downstream parts (J01-P5, J02, J03) are entitled to find in each fixture.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pytest

from tests.fixtures.builders import ALL_FIXTURES, build

#: Literal XML substrings every fixture must contain, searched across all XML parts.
EXPECTED_MARKERS: dict[str, tuple[str, ...]] = {
    "simple": (
        '<w:pStyle w:val="Heading1"/>',
        '<w:pStyle w:val="Heading2"/>',
        "<w:t>First paragraph of the simple fixture.</w:t>",
    ),
    "mixed_runs": (
        'w:rsidR="00A10002"',
        'w:rsidR="00A10006"',
        "<w:rPr><w:b/><w:i/></w:rPr>",
        "<w:tab/>",
        "<w:br/>",
        "<w:noProof/>",
        '<w:lang w:val="fr-FR"/>',
    ),
    "paragraph_styles": (
        'w:styleId="FixtureBody"',
        'w:styleId="FixtureNote"',
        '<w:pStyle w:val="FixtureBody"/>',
        '<w:pStyle w:val="Quote"/>',
    ),
    "character_styles": (
        'w:type="character" w:styleId="FixtureEmphasis"',
        '<w:rStyle w:val="FixtureEmphasis"/>',
        '<w:rStyle w:val="FixtureCode"/>',
        '<w:rStyle w:val="Strong"/>',
    ),
    "style_inheritance": (
        '<w:basedOn w:val="FixtureRoot"/>',
        '<w:basedOn w:val="FixtureBranch"/>',
        '<w:next w:val="FixtureRoot"/>',
        '<w:link w:val="FixtureBranchChar"/>',
        '<w:link w:val="FixtureBranch"/>',
    ),
    "themes": (
        'w:asciiTheme="majorHAnsi"',
        'w:cstheme="majorBidi"',
        'w:themeColor="accent1"',
        'w:themeTint="66"',
        'w:themeShade="BF"',
        'w:themeFill="accent1"',
    ),
    "complex_numbering": (
        '<w:abstractNum w:abstractNumId="900">',
        '<w:abstractNum w:abstractNumId="901">',
        '<w:numFmt w:val="lowerRoman"/>',
        '<w:lvlRestart w:val="1"/>',
        '<w:startOverride w:val="5"/>',
        '<w:numId w:val="902"/>',
    ),
    "tables": (
        "<w:tbl>",
        '<w:gridSpan w:val="2"/>',
        '<w:vMerge w:val="restart"/>',
        "<w:vMerge/>",
        '<w:tblStyle w:val="TableGrid"/>',
    ),
    "comments": (
        '<w:commentRangeStart w:id="1"/>',
        '<w:commentRangeEnd w:id="3"/>',
        '<w:commentReference w:id="2"/>',
        "<w:annotationRef/>",
        '<w:comment w:id="2"',
        'w14:paraId="10000002"',
        'w15:paraIdParent="10000001"',
    ),
    "tracked_changes": (
        '<w:ins w:id="201"',
        '<w:del w:id="202"',
        "<w:delText",
        '<w:ins w:id="203"',
        '<w:del w:id="204"',
        '<w:rPrChange w:id="205"',
        '<w:pPrChange w:id="206"',
        '<w:del w:id="207"',
        '<w:ins w:id="208"',
    ),
    "hyperlinks": (
        "<w:hyperlink r:id=",
        'w:anchor="FixtureAnchor"',
        'TargetMode="External"',
        '<w:rStyle w:val="Hyperlink"/>',
    ),
    "bookmarks": (
        '<w:bookmarkStart w:id="310" w:name="FixtureInline"/>',
        '<w:bookmarkStart w:id="311" w:name="FixtureSpanning"/>',
        '<w:bookmarkStart w:id="312" w:name="FixtureEmpty"/><w:bookmarkEnd w:id="312"/>',
        '<w:bookmarkEnd w:id="313"/>',
    ),
    "fields": (
        '<w:fldSimple w:instr=" PAGE   \\* MERGEFORMAT ">',
        '<w:fldChar w:fldCharType="begin"/>',
        '<w:fldChar w:fldCharType="separate"/>',
        '<w:fldChar w:fldCharType="end"/>',
        "<w:instrText",
        " TOC ",
        " REF FixtureRefTarget ",
        'w:dirty="true"',
    ),
    "footnotes": (
        '<w:footnoteReference w:id="2"/>',
        '<w:footnoteReference w:id="3"/>',
        "<w:footnoteRef/>",
        '<w:footnote w:type="separator" w:id="-1">',
        '<w:endnoteReference w:id="2"/>',
        "<w:endnoteRef/>",
    ),
    "headers_footers": (
        "<w:hdr ",
        "<w:ftr ",
        "Fixture header ",
        "Fixture first-page header",
        '<w:fldSimple w:instr=" NUMPAGES ">',
        "<w:drawing>",
        "<w:titlePg/>",
    ),
    "sections": (
        '<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>',
        "<w:pPr><w:sectPr>",
    ),
    "content_controls": (
        "<w:sdt>",
        "<w:sdtPr>",
        "<w:sdtContent>",
        '<w:tag w:val="fixture-block"/>',
        '<w:tag w:val="fixture-inline"/>',
        '<w:listItem w:displayText="Beta" w:value="beta"/>',
    ),
    "drawings": (
        "<w:drawing>",
        "<wp:inline",
        "<pic:pic>",
        "<a:blip",
    ),
    "combined": (
        '<w:commentRangeStart w:id="1"/>',
        '<w:ins w:id="201"',
        "<w:tbl>",
        '<w:footnoteReference w:id="2"/>',
        "<w:sdtContent>",
        "<w:drawing>",
        '<w:abstractNum w:abstractNumId="900">',
        '<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>',
    ),
}

FIXTURE_NAMES = sorted(ALL_FIXTURES)


def _all_xml(blob: bytes) -> str:
    """Concatenate every XML member of the package into one searchable string."""
    archive = zipfile.ZipFile(io.BytesIO(blob))
    return "\n".join(
        archive.read(name).decode("utf-8")
        for name in archive.namelist()
        if name.endswith((".xml", ".rels"))
    )


def test_registry_is_large_enough_and_named_consistently() -> None:
    assert len(ALL_FIXTURES) >= 18
    assert set(ALL_FIXTURES) == set(EXPECTED_MARKERS)
    for name, builder in ALL_FIXTURES.items():
        assert builder.__name__ == f"build_{name}"


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_fixture_opens_with_python_docx(name: str) -> None:
    blob = build(name)
    assert blob[:2] == b"PK"

    archive = zipfile.ZipFile(io.BytesIO(blob))
    assert archive.testzip() is None
    members = set(archive.namelist())
    assert {"[Content_Types].xml", "word/document.xml", "word/styles.xml"} <= members

    document = docx.Document(io.BytesIO(blob))
    assert document.element.body is not None


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_fixture_is_deterministic(name: str) -> None:
    assert build(name) == build(name)


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_fixture_contains_expected_markers(name: str) -> None:
    xml = _all_xml(build(name))
    missing = [marker for marker in EXPECTED_MARKERS[name] if marker not in xml]
    assert not missing, f"{name} is missing {missing}"


@pytest.mark.parametrize("name", ["sections", "combined"])
def test_two_sections_are_present(name: str) -> None:
    document = docx.Document(io.BytesIO(build(name)))
    assert len(document.sections) == 2


def test_comment_thread_is_wired_through_comments_extended() -> None:
    archive = zipfile.ZipFile(io.BytesIO(build("comments")))
    comments = archive.read("word/comments.xml").decode("utf-8")
    extended = archive.read("word/commentsExtended.xml").decode("utf-8")
    # The reply's paraId must be declared as a child of the root comment's paraId.
    assert 'w14:paraId="10000001"' in comments
    assert 'w14:paraId="10000002"' in comments
    assert 'w15:paraId="10000002" w15:paraIdParent="10000001"' in extended


def test_footnotes_and_endnotes_are_separate_parts() -> None:
    members = set(zipfile.ZipFile(io.BytesIO(build("footnotes"))).namelist())
    assert {"word/footnotes.xml", "word/endnotes.xml"} <= members


def test_build_rejects_unknown_name() -> None:
    with pytest.raises(KeyError, match="unknown fixture"):
        build("does_not_exist")


def test_fixture_docx_writes_a_readable_file(fixture_docx) -> None:
    path = fixture_docx("simple")
    assert path.name == "simple.docx"
    assert path.read_bytes() == build("simple")
    docx.Document(str(path))
