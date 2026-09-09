"""Tests for ``add_table_of_contents``, now backed by the OOXML engine.

The tool used to rebuild the whole document inside a blank ``Document()`` and
save that over the original: only ``paragraph.text``, ``cell.text`` and style
names survived, and the "table of contents" it produced was a list of plain
paragraphs that no reader could click and that Word could never refresh.

What it writes now is what Word writes: a ``w:sdt`` of the "Table of Contents"
gallery holding a ``TOC \\o "1-N" \\h \\z \\u`` complex field, whose cached
result -- the runs between ``separate`` and ``end`` -- carries the headings
found, plus a ``w:updateFields`` in ``settings.xml`` so the field is refreshed,
with real page numbers, the first time the document is opened.

What is pinned here is both halves of the contract: the API a caller sees (the
parameters, the two return messages) is unchanged, and the document the tool was
given comes back with everything it had.
"""

from __future__ import annotations

import asyncio
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document as PDocument
from lxml import etree

from tests.support.package_check import validate_package
from tests.support.snapshot import Snapshot, snapshot
from word_document_server.tools.content_tools import add_table_of_contents

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _w(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _run(coro: Any) -> Any:
    """Drive an ``async`` tool call to completion."""
    return asyncio.run(coro)


def _document_root(path: Path) -> etree._Element:
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _part_root(path: Path, name: str) -> etree._Element | None:
    with zipfile.ZipFile(path) as archive:
        if name not in archive.namelist():
            return None
        return etree.fromstring(archive.read(name))


def _toc_control(root: etree._Element) -> etree._Element:
    """The one ``w:sdt`` of the "Table of Contents" gallery, or fail."""
    found = [
        control
        for control in root.iter(_w("sdt"))
        if any(
            gallery.get(_w("val")) == "Table of Contents"
            for gallery in control.iter(_w("docPartGallery"))
        )
    ]
    assert len(found) == 1, f"expected exactly one TOC content control, got {len(found)}"
    return found[0]


def _field_chars(element: etree._Element) -> list[str]:
    return [
        char.get(_w("fldCharType")) for char in element.iter(_w("fldChar"))
    ]


def _instructions(element: etree._Element) -> list[str]:
    return [node.text or "" for node in element.iter(_w("instrText"))]


def _entry_texts(control: etree._Element) -> list[str]:
    """Visible text of every paragraph of the control, in document order."""
    texts = []
    for paragraph in control.iter(_w("p")):
        texts.append("".join(node.text or "" for node in paragraph.iter(_w("t"))))
    return texts


def _headingless(tmp_path: Path) -> Path:
    """A document made of plain paragraphs only."""
    path = tmp_path / "headingless.docx"
    document = PDocument()
    document.add_paragraph("Just a paragraph.")
    document.add_paragraph("And another one.")
    document.save(str(path))
    return path


def _counters(snap: Snapshot, *names: str) -> dict[str, int]:
    return {name: snap.counters[name] for name in names}


# --------------------------------------------------------------------------------------
# The message a caller gets
# --------------------------------------------------------------------------------------


def test_the_message_reports_the_number_of_entries(fixture_docx) -> None:
    path = fixture_docx("simple")

    result = _run(add_table_of_contents(str(path)))

    # "Fixture: simple" (level 1) and "Second heading" (level 2).
    assert result == f"Table of contents with 2 entries added to {path}"


def test_a_document_without_headings_is_left_alone(tmp_path) -> None:
    path = _headingless(tmp_path)
    before = path.read_bytes()

    result = _run(add_table_of_contents(str(path)))

    assert result == (
        f"No headings found in document {path}. Table of contents not created."
    )
    assert path.read_bytes() == before


def test_a_missing_document_is_reported_not_created(tmp_path) -> None:
    path = tmp_path / "absent.docx"

    assert _run(add_table_of_contents(str(path))) == f"Document {path} does not exist"
    assert not path.exists()


# --------------------------------------------------------------------------------------
# What the tool writes
# --------------------------------------------------------------------------------------


def test_the_toc_is_a_content_control_holding_a_toc_field(fixture_docx) -> None:
    path = fixture_docx("simple")

    _run(add_table_of_contents(str(path)))

    control = _toc_control(_document_root(path))
    assert _instructions(control) == [' TOC \\o "1-3" \\h \\z \\u ']
    # A complex field is begin / instruction / separate / cached result / end;
    # a field missing its `separate` shows its own instruction to the reader.
    assert _field_chars(control) == ["begin", "separate", "end"]


def test_the_cached_result_carries_the_headings_it_found(fixture_docx) -> None:
    path = fixture_docx("simple")

    _run(add_table_of_contents(str(path)))

    control = _toc_control(_document_root(path))
    assert _entry_texts(control) == [
        "Table of Contents",
        "Fixture: simple",
        "Second heading",
    ]


def test_max_level_limits_both_the_entries_and_the_switch(fixture_docx) -> None:
    path = fixture_docx("simple")

    result = _run(add_table_of_contents(str(path), max_level=1))

    control = _toc_control(_document_root(path))
    assert _instructions(control) == [' TOC \\o "1-1" \\h \\z \\u ']
    # "Second heading" is a Heading 2 and is out of the requested range.
    assert _entry_texts(control) == ["Table of Contents", "Fixture: simple"]
    assert "with 1 entries" in result


def test_an_empty_title_writes_no_title_paragraph(fixture_docx) -> None:
    path = fixture_docx("simple")

    _run(add_table_of_contents(str(path), title=""))

    control = _toc_control(_document_root(path))
    assert _entry_texts(control) == ["Fixture: simple", "Second heading"]


def test_the_toc_goes_after_the_first_heading(fixture_docx) -> None:
    path = fixture_docx("simple")
    root = _document_root(path)
    body = root.find(_w("body"))
    assert body[0].tag == _w("p")  # sanity: the fixture opens on its title

    _run(add_table_of_contents(str(path)))

    body = _document_root(path).find(_w("body"))
    blocks = [child for child in body if child.tag in (_w("p"), _w("tbl"), _w("sdt"))]
    assert blocks[0].tag == _w("p")
    assert "Fixture: simple" in "".join(
        node.text or "" for node in blocks[0].iter(_w("t"))
    )
    assert blocks[1].tag == _w("sdt")


def test_the_toc_goes_first_when_the_body_opens_on_a_paragraph(tmp_path) -> None:
    path = tmp_path / "late-heading.docx"
    document = PDocument()
    document.add_paragraph("An abstract, before any heading.")
    document.add_heading("The first heading", level=1)
    document.save(str(path))

    _run(add_table_of_contents(str(path)))

    body = _document_root(path).find(_w("body"))
    blocks = [child for child in body if child.tag in (_w("p"), _w("tbl"), _w("sdt"))]
    assert blocks[0].tag == _w("sdt")


def test_the_fields_are_marked_for_update_in_schema_order(fixture_docx) -> None:
    path = fixture_docx("simple")

    _run(add_table_of_contents(str(path)))

    settings = _part_root(path, "word/settings.xml")
    assert settings is not None
    update = settings.findall(_w("updateFields"))
    assert len(update) == 1
    assert update[0].get(_w("val")) == "true"
    # w:settings is an xsd:sequence: w:updateFields comes before w:compat and
    # w:rsids, and Word offers to repair a document that gets that order wrong.
    localnames = [
        etree.QName(child).localname for child in settings if isinstance(child.tag, str)
    ]
    for follower in ("compat", "rsids", "mathPr", "themeFontLang"):
        if follower in localnames:
            assert localnames.index("updateFields") < localnames.index(follower)


def test_adding_a_toc_twice_leaves_one_update_flag(fixture_docx) -> None:
    path = fixture_docx("simple")

    _run(add_table_of_contents(str(path)))
    _run(add_table_of_contents(str(path)))

    settings = _part_root(path, "word/settings.xml")
    assert len(settings.findall(_w("updateFields"))) == 1
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# What the tool must not destroy
# --------------------------------------------------------------------------------------


def test_no_part_is_dropped_from_a_document_that_has_everything(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())
    assert "word/comments.xml" in before.parts  # sanity: the fixture has one

    _run(add_table_of_contents(str(path)))

    after = snapshot(path.read_bytes())
    assert set(before.parts) <= set(after.parts)
    assert before.content_types <= after.content_types
    assert validate_package(path) == []


def test_the_annotations_of_the_document_survive(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())
    watched = (
        "comments",
        "comment_references",
        "revisions",
        "bookmarks",
        "hyperlinks",
        "footnotes",
        "footnote_references",
        "endnotes",
        "endnote_references",
    )

    _run(add_table_of_contents(str(path)))

    after = snapshot(path.read_bytes())
    assert _counters(after, *watched) == _counters(before, *watched)
    # One field more, and one only: the TOC field the tool was asked to add.
    assert after.counters["fields"] == before.counters["fields"] + 1


def test_the_body_text_and_its_tables_are_untouched(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())

    _run(add_table_of_contents(str(path)))

    after = snapshot(path.read_bytes())
    # The table of contents is inserted after the first heading, so the paragraph index
    # of everything below it shifts; what must hold is that every paragraph of
    # every story is still there, in order, with the same text once the
    # table of contents' own paragraphs are taken back out.
    body_after = [sig.text for sig in after.story_paragraphs("document")]
    inserted = len(body_after) - len(before.story_paragraphs("document"))
    assert inserted > 0
    assert [sig.text for sig in before.story_paragraphs("document")] == (
        body_after[:1] + body_after[1 + inserted :]
    )
    for story in before.stories:
        if story == "document":
            continue
        assert [sig.text for sig in before.story_paragraphs(story)] == [
            sig.text for sig in after.story_paragraphs(story)
        ]
    assert before.tables == after.tables
    assert before.sect_pr == after.sect_pr


def test_the_relationships_of_the_document_are_untouched(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())

    _run(add_table_of_contents(str(path)))

    after = snapshot(path.read_bytes())
    for owner, relationships in before.relationships.items():
        assert relationships <= after.relationships[owner], owner


def test_an_entry_carries_the_visible_text_of_its_heading(tmp_path) -> None:
    """Deleted runs and field instructions never leak into a TOC entry."""
    path = tmp_path / "rich-heading.docx"
    document = PDocument()
    heading = document.add_heading("Kept", level=1)
    for fragment in (
        (
            f'<w:del xmlns:w="{W}" w:id="900" w:author="a" w:date="2024-01-01T00:00:00Z">'
            f'<w:r><w:delText xml:space="preserve"> gone</w:delText></w:r></w:del>'
        ),
        f'<w:r xmlns:w="{W}"><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r xmlns:w="{W}"><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>',
        f'<w:r xmlns:w="{W}"><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r xmlns:w="{W}"><w:t>7</w:t></w:r>',
        f'<w:r xmlns:w="{W}"><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        heading._p.append(etree.fromstring(fragment))
    document.save(str(path))

    _run(add_table_of_contents(str(path), title=""))

    control = _toc_control(_document_root(path))
    assert _entry_texts(control) == ["Kept7"]
    # The heading's own PAGE instruction is not copied into the entry: the only
    # instruction inside the control is the TOC field's.
    assert _instructions(control) == [' TOC \\o "1-3" \\h \\z \\u ']
    assert validate_package(path) == []


@pytest.mark.parametrize("name", ["simple", "combined", "sections"])
def test_the_package_stays_valid(fixture_docx, name: str) -> None:
    path = fixture_docx(name)

    _run(add_table_of_contents(str(path)))

    assert validate_package(path) == []
    # A malformed sdt or an unbalanced field would still parse; reopening the
    # package with python-docx is the cheapest check that the body is walkable.
    assert PDocument(str(path)).paragraphs
