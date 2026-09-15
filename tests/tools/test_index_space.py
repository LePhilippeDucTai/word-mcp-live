"""One index space for every public ``paragraph_index``.

``find_text`` numbers the paragraphs of the body in document order, the content
of a block ``w:sdt`` included.  ``delete_paragraph`` and ``add_bookmark`` were
moved onto that space; the reading tool, the spacing tool and the three
``insert_*_near_text`` anchors were still counting python-docx's direct body
children, which is a *different* space as soon as a block content control sits
in the document -- and ``add_table_of_contents`` inserts exactly one.

Nothing here is destructive, which is why the drift survived: each tool simply
reads, spaces or anchors at a paragraph several positions away from the one the
search tool pointed at, and reports the index it was given as if it had.  The
tests below pin the only property that makes an index usable across two calls:
the paragraph a tool acts on is the paragraph ``find_text`` reported.

The helpers are copied from ``tests/tools/test_layout_blocks.py`` rather than
imported: a test module is not a library, and both files must keep the right to
describe the index space on their own terms.
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document as PDocument

from tests.support.package_check import validate_package
from tests.support.snapshot import assert_unchanged_except, snapshot
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import visible_text
from word_document_server.engine.xmlns import qn
from word_document_server.tools.content_tools import add_table_of_contents
from word_document_server.tools.extended_document_tools import (
    get_paragraph_text_from_document,
)
from word_document_server.tools.layout_tools import add_bookmark, set_paragraph_spacing
from word_document_server.utils.document_utils import (
    insert_header_near_text,
    insert_line_or_paragraph_near_text,
    insert_numbered_list_near_text,
)
from word_document_server.utils.extended_document_utils import find_text

BLOCK_TAGS = (qn("w:p"), qn("w:tbl"), qn("w:sdt"))

#: Text of the paragraph every test targets, and of the one they insert.
ANCHOR = "Charlie"
INSERTED = "Inserted by the test"


def _run(coro: Any) -> Any:
    """Drive an ``async`` tool call to completion."""
    return asyncio.run(coro)


def _body(path: Path):
    return DocxPackage.open(str(path)).document.find(qn("w:body"))


def _blocks(path: Path) -> list:
    return [child for child in _body(path) if child.tag in BLOCK_TAGS]


def _all_paragraphs(path: Path) -> list:
    """Every ``w:p`` of the body, nesting included, in document order.

    Deliberately independent of the engine's own filter: on a document holding
    no table and no text box this is exactly the V2 index space the tools
    address, so a test can state what the space contains without asking the
    code under test what it thinks the space is.
    """
    return list(_body(path).iter(qn("w:p")))


def _all_paragraph_texts(path: Path) -> list[str]:
    """Visible text of every ``w:p`` of the body, nesting included."""
    return [visible_text(paragraph) for paragraph in _all_paragraphs(path)]


def _v2_index_of(path: Path, text: str) -> int:
    """Index ``find_text`` reports for the single paragraph holding `text`."""
    occurrences = find_text(str(path), text)["occurrences"]
    assert len(occurrences) == 1, occurrences
    return occurrences[0]["paragraph_index"]


BODY_TEXTS = ("Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot")


def _document_with_a_table_of_contents(path: Path) -> Path:
    document = PDocument()
    document.add_heading("Titre", level=1)
    for text in BODY_TEXTS:
        document.add_paragraph(text)
    document.save(str(path))
    assert "Table of contents" in _run(add_table_of_contents(str(path)))
    assert any(block.tag == qn("w:sdt") for block in _blocks(path))  # sanity
    return path


def _spaced_paragraph_texts(path: Path) -> list[str]:
    """Text of every paragraph carrying a direct ``w:spacing``."""
    return [
        visible_text(paragraph)
        for paragraph in _all_paragraphs(path)
        if paragraph.find(f"{qn('w:pPr')}/{qn('w:spacing')}") is not None
    ]


# --------------------------------------------------------------------------------------
# get_paragraph_text_from_document
# --------------------------------------------------------------------------------------


def test_the_paragraph_read_is_the_one_find_text_reported(tmp_path) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    index = _v2_index_of(path, ANCHOR)

    result = json.loads(_run(get_paragraph_text_from_document(str(path), index)))

    assert result["text"] == ANCHOR
    assert result["index"] == index
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# set_paragraph_spacing
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("form", ["single", "range"])
def test_the_paragraph_spaced_is_the_one_find_text_reported(tmp_path, form) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    index = _v2_index_of(path, ANCHOR)
    assert _spaced_paragraph_texts(path) == []  # sanity: nothing is spaced yet
    before = snapshot(path.read_bytes())

    if form == "single":
        result = _run(
            set_paragraph_spacing(str(path), paragraph_index=index, space_before_pt=18)
        )
    else:
        result = _run(
            set_paragraph_spacing(
                str(path),
                start_paragraph=index,
                end_paragraph=index,
                space_before_pt=18,
            )
        )

    assert json.loads(result)["paragraphs_affected"] == 1
    assert _spaced_paragraph_texts(path) == [ANCHOR]
    # No table here, so the snapshot space and the V2 space are the same one.
    assert_unchanged_except(before, snapshot(path.read_bytes()), paragraphs=[index])
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# insert_header_near_text / insert_line_or_paragraph_near_text / insert_numbered_list_near_text
# --------------------------------------------------------------------------------------


def _insert_header(path: Path, index: int) -> str:
    return insert_header_near_text(
        str(path), header_title=INSERTED, position="after", target_paragraph_index=index
    )


def _insert_line(path: Path, index: int) -> str:
    return insert_line_or_paragraph_near_text(
        str(path), line_text=INSERTED, position="after", target_paragraph_index=index
    )


def _insert_list(path: Path, index: int) -> str:
    return insert_numbered_list_near_text(
        str(path), list_items=[INSERTED], position="after", target_paragraph_index=index
    )


@pytest.mark.parametrize(
    "insert",
    [_insert_header, _insert_line, _insert_list],
    ids=["header", "line", "list"],
)
def test_the_anchor_of_each_insert_near_text_is_the_one_find_text_reported(
    tmp_path, insert
) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    index = _v2_index_of(path, ANCHOR)

    result = insert(path, index)

    # A substring: insert_line_or_paragraph_near_text interpolates the style it
    # inherited from the anchor, and the shape of that tail is not what is
    # under test here.
    assert f"(index {index})" in result, result
    texts = _all_paragraph_texts(path)
    assert texts[texts.index(ANCHOR) + 1] == INSERTED
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# add_bookmark: the lower bound of the space
# --------------------------------------------------------------------------------------


def test_a_negative_bookmark_index_is_refused(tmp_path) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    before = path.read_bytes()

    result = _run(add_bookmark(str(path), -1, "Neg"))

    assert "success" not in result, result
    assert not [
        marker
        for marker in _body(path).iter(qn("w:bookmarkStart"))
        if marker.get(qn("w:name")) == "Neg"
    ]
    assert path.read_bytes() == before, "a refused call writes nothing"
