"""Tests for the two text-search tools now backed by the OOXML engine.

``search_and_replace`` (``tools/content_tools.py`` ->
``utils/document_utils.replace_text_everywhere``) and ``find_text_in_document``
(``tools/extended_document_tools.py`` -> ``utils/extended_document_utils.find_text``)
both delegate to :func:`word_document_server.engine.find.find`. What is pinned
here is the behaviour a caller sees: the message and the JSON shape are the ones
the tools always had, and the new reach -- runs, hyperlinks, tracked insertions,
content controls, tables, headers, footers, notes -- comes with the guarantee
that nothing is destroyed on the way.

Paragraph identity is established by text, never by a hard-coded index: several
paragraph-index spaces coexist in this repository, and the one these tools now
report is the V2 space of ``Match.index`` (top-level paragraphs of a story,
``None`` inside a table cell or a text box).
"""

from __future__ import annotations

import asyncio
import json
import zipfile
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from tests.support.package_check import validate_package
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import visible_text
from word_document_server.engine.xmlns import qn
from word_document_server.tools.content_tools import search_and_replace
from word_document_server.tools.extended_document_tools import find_text_in_document
from word_document_server.utils.document_utils import (
    ReplacementReport,
    find_and_replace_text,
    replace_text_everywhere,
)
from word_document_server.utils.extended_document_utils import find_text


def _run(coro: Any) -> Any:
    """Drive an ``async`` tool call to completion."""
    return asyncio.run(coro)


def _texts(path, story: str = "document") -> list[str]:
    """Visible text of every paragraph of `story`, in document order."""
    pkg = DocxPackage.open(str(path))
    root = dict(pkg.stories())[story]
    return [visible_text(p) for p in root.iter(qn("w:p"))]


def _occurrences(result: dict) -> list[dict]:
    assert "error" not in result, result
    return result["occurrences"]


# --------------------------------------------------------------------------
# search_and_replace: reach
# --------------------------------------------------------------------------


def test_replaces_a_match_split_across_three_runs(fixture_docx):
    path = fixture_docx("mixed_runs")

    result = _run(search_and_replace(str(path), "sentence cut in three", "phrase in three"))

    assert result == (
        "Replaced 1 occurrence(s) of 'sentence cut in three' with 'phrase in three'."
    )
    assert "One phrase in three runs." in _texts(path)
    assert validate_package(path) == []


def test_counts_occurrences_not_runs(fixture_docx):
    # "paragraph" appears once in each of two body paragraphs. The previous
    # per-run implementation counted one replacement per *run* it touched.
    path = fixture_docx("simple")

    result = _run(search_and_replace(str(path), "paragraph", "section"))

    assert result.startswith("Replaced 2 occurrence(s)")
    assert "First section of the simple fixture." in _texts(path)
    assert "Second section, with a trailing sentence." in _texts(path)


def test_replacement_inside_a_hyperlink_keeps_the_link(fixture_docx):
    path = fixture_docx("hyperlinks")
    with zipfile.ZipFile(path) as archive:
        before_rels = archive.read("word/_rels/document.xml.rels")

    result = _run(search_and_replace(str(path), "example.org", "sample.test"))

    assert result.startswith("Replaced 1 occurrence(s)")
    pkg = DocxPackage.open(str(path))
    root = dict(pkg.stories())["document"]
    links = root.findall(f".//{qn('w:hyperlink')}")
    external = [link for link in links if link.get(qn("r:id"))]
    assert len(external) == 1, "the external hyperlink must survive as one element"
    assert visible_text(external[0].getparent()) == "An external link: sample.test."
    # The replacement text lives inside the hyperlink, not beside it.
    assert "sample.test" in "".join(external[0].itertext())
    with zipfile.ZipFile(path) as archive:
        assert archive.read("word/_rels/document.xml.rels") == before_rels
    assert validate_package(path) == []


def test_replacement_inside_a_tracked_insertion_stays_inside_it(fixture_docx):
    path = fixture_docx("tracked_changes")

    result = _run(search_and_replace(str(path), "inserted text", "added text"))

    assert result.startswith("Replaced 1 occurrence(s)")
    pkg = DocxPackage.open(str(path))
    root = dict(pkg.stories())["document"]
    insertions = root.findall(f".//{qn('w:ins')}")
    assert any("added text" in "".join(ins.itertext()) for ins in insertions)
    # The neighbouring deletion is untouched: its w:delText still carries its text.
    deletions = root.findall(f".//{qn('w:del')}")
    assert any("deleted text, " in "".join(d.itertext()) for d in deletions)
    assert validate_package(path) == []


def test_text_hidden_under_a_deletion_is_never_matched(fixture_docx):
    path = fixture_docx("tracked_changes")

    result = _run(search_and_replace(str(path), "deleted text", "REWRITTEN"))

    assert result == "No occurrences of 'deleted text' found."
    assert "REWRITTEN" not in "".join(_texts(path))


def test_replacement_inside_a_content_control_keeps_the_sdt(fixture_docx):
    path = fixture_docx("content_controls")
    before_controls = len(
        dict(DocxPackage.open(str(path)).stories())["document"].findall(
            f".//{qn('w:sdt')}"
        )
    )

    result = _run(search_and_replace(str(path), "inline control content", "inline text"))

    assert result.startswith("Replaced 1 occurrence(s)")
    root = dict(DocxPackage.open(str(path)).stories())["document"]
    assert len(root.findall(f".//{qn('w:sdt')}")) == before_controls
    assert "Before the control, inline text, after the control." in _texts(path)
    assert validate_package(path) == []


def test_replacement_reaches_table_cells_headers_and_footnotes(fixture_docx):
    path = fixture_docx("combined")

    assert _run(search_and_replace(str(path), "Row two, column two", "Cell rewritten"))
    assert _run(search_and_replace(str(path), "Fixture header", "Rewritten header"))
    assert _run(search_and_replace(str(path), "First footnote body", "Rewritten note"))

    assert "Cell rewritten" in "".join(_texts(path))
    assert "Rewritten header" in "".join(_texts(path, "header1"))
    assert "Rewritten note" in "".join(_texts(path, "footnotes"))
    assert validate_package(path) == []


# --------------------------------------------------------------------------
# search_and_replace: what it refuses to touch
# --------------------------------------------------------------------------


def test_occurrence_overlapping_a_field_is_skipped_and_reported(fixture_docx):
    # "Fixture: fields" appears twice: as a plain heading, and as the cached
    # result of the TOC field, which a replacement must not cut into.
    path = fixture_docx("fields")

    result = _run(search_and_replace(str(path), "Fixture: fields", "Fixture: champs"))

    assert result == (
        "Replaced 1 occurrence(s) of 'Fixture: fields' with 'Fixture: champs'."
        " 1 skipped (inside fields)."
    )
    texts = _texts(path)
    assert "Fixture: champs" in texts
    assert "Fixture: fields\t1" in texts, "the TOC field's result must be untouched"
    assert validate_package(path) == []


def test_a_wholly_skipped_search_still_reports_the_skip(fixture_docx):
    # "Page 1" swallows the PAGE field that renders the "1".
    path = fixture_docx("fields")
    before = path.read_bytes()

    result = _run(search_and_replace(str(path), "Page 1", "Page one"))

    assert result == "No occurrences of 'Page 1' found. 1 skipped (inside fields)."
    assert path.read_bytes() == before, "nothing is saved when nothing is replaced"


def test_toc_styled_paragraphs_are_skipped(fixture_docx):
    path = fixture_docx("simple")
    doc = Document(str(path))
    doc.add_paragraph("Contents entry marker")
    toc_style = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Contents entry marker").style = toc_style
    doc.save(str(path))

    result = _run(search_and_replace(str(path), "entry marker", "entry mark"))

    assert result.startswith("Replaced 1 occurrence(s)")
    texts = _texts(path)
    assert texts.count("Contents entry mark") == 1
    assert texts.count("Contents entry marker") == 1


def test_missing_document_is_reported(tmp_path):
    missing = tmp_path / "nope.docx"
    assert _run(search_and_replace(str(missing), "a", "b")) == (
        f"Document {missing} does not exist"
    )


def test_empty_search_text_is_refused_rather_than_matching_everywhere(fixture_docx):
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = _run(search_and_replace(str(path), "", "X"))

    assert result.startswith("Failed to search and replace:")
    assert path.read_bytes() == before


# --------------------------------------------------------------------------
# The utils layer directly
# --------------------------------------------------------------------------


def test_replace_text_everywhere_returns_a_report(fixture_docx):
    pkg = DocxPackage.open(str(fixture_docx("fields")))

    report = replace_text_everywhere(pkg, "Fixture: fields", "Fixture: champs")

    assert report == ReplacementReport(replaced=1, skipped=1)


def test_find_and_replace_text_returns_the_replacement_count(fixture_docx):
    pkg = DocxPackage.open(str(fixture_docx("simple")))

    assert find_and_replace_text(pkg, "paragraph", "section") == 2


def test_find_and_replace_text_rejects_a_python_docx_document(fixture_docx):
    doc = Document(str(fixture_docx("simple")))

    with pytest.raises(TypeError, match="DocxPackage"):
        find_and_replace_text(doc, "paragraph", "section")


# --------------------------------------------------------------------------
# find_text
# --------------------------------------------------------------------------


def test_find_text_reports_story_v2_index_and_offsets(fixture_docx):
    path = fixture_docx("simple")

    occurrences = _occurrences(find_text(str(path), "simple fixture"))

    assert len(occurrences) == 1
    occurrence = occurrences[0]
    assert occurrence["story"] == "document"
    assert occurrence["paragraph_index"] == 1
    assert occurrence["position"] == occurrence["start"] == 23
    assert occurrence["end"] == 37
    assert occurrence["text"] == "simple fixture"
    assert occurrence["context"] == "First paragraph of the simple fixture."
    assert "location" not in occurrence


def test_find_text_keeps_the_historical_result_keys(fixture_docx):
    path = fixture_docx("simple")

    result = find_text(str(path), "paragraph")

    assert result["query"] == "paragraph"
    assert result["match_case"] is True
    assert result["whole_word"] is False
    assert result["total_count"] == len(result["occurrences"]) == 2
    for occurrence in result["occurrences"]:
        assert set(occurrence) >= {"paragraph_index", "position", "context"}


def test_find_text_reports_none_index_and_a_location_in_a_table_cell(fixture_docx):
    path = fixture_docx("tables")

    occurrences = _occurrences(find_text(str(path), "Row two, column two"))

    assert len(occurrences) == 1
    assert occurrences[0]["paragraph_index"] is None
    assert occurrences[0]["location"] == "Table 0, Row 1, Column 1"
    assert occurrences[0]["story"] == "document"


def test_find_text_indexes_the_paragraph_after_a_table_in_v2_space(fixture_docx):
    path = fixture_docx("tables")

    occurrences = _occurrences(find_text(str(path), "Paragraph after the table."))

    # The table's own paragraphs consume no V2 slot: this is the second
    # top-level paragraph of the story, so index 1.
    assert occurrences[0]["paragraph_index"] == 1


def test_find_text_names_the_real_story_never_the_body_alias(fixture_docx):
    path = fixture_docx("combined")

    stories = {occ["story"] for occ in _occurrences(find_text(str(path), "Fixture"))}

    assert "body" not in stories
    assert "document" in stories
    assert any(story.startswith("header") for story in stories)


def test_find_text_searches_footnotes(fixture_docx):
    path = fixture_docx("footnotes")

    occurrences = _occurrences(find_text(str(path), "First footnote body"))

    assert [occ["story"] for occ in occurrences] == ["footnotes"]
    assert occurrences[0]["paragraph_index"] == 2


def test_find_text_finds_a_match_split_across_runs_and_containers(fixture_docx):
    path = fixture_docx("hyperlinks")

    occurrences = _occurrences(find_text(str(path), "link: example.org"))

    assert len(occurrences) == 1
    assert occurrences[0]["text"] == "link: example.org"


def test_find_text_honours_match_case_and_whole_word(fixture_docx):
    path = fixture_docx("simple")

    # "simple" reads twice: "Fixture: simple" and "the simple fixture".
    assert find_text(str(path), "SIMPLE", match_case=True)["total_count"] == 0
    assert find_text(str(path), "SIMPLE", match_case=False)["total_count"] == 2

    # "paragraph" is a whole word twice; "aragraph" never is.
    assert find_text(str(path), "paragraph", whole_word=True)["total_count"] == 2
    assert find_text(str(path), "aragraph", whole_word=True)["total_count"] == 0
    assert find_text(str(path), "aragraph", whole_word=False)["total_count"] == 2


def test_find_text_whole_word_positions_are_character_offsets(fixture_docx):
    path = fixture_docx("simple")

    occurrences = _occurrences(find_text(str(path), "paragraph", whole_word=True))

    for occurrence in occurrences:
        assert occurrence["end"] - occurrence["start"] == len("paragraph")
        assert occurrence["position"] == occurrence["start"]


def test_find_text_ellipsises_a_long_paragraph_context(fixture_docx):
    path = fixture_docx("simple")
    doc = Document(str(path))
    long_text = "needle " + "filler " * 40
    doc.add_paragraph(long_text)
    doc.save(str(path))

    occurrence = _occurrences(find_text(str(path), "needle"))[0]

    assert occurrence["context"].endswith("...")
    assert len(occurrence["context"]) == 103
    assert len(occurrence["match_context"]) < len(occurrence["context"])


def test_find_text_does_not_see_deleted_text(fixture_docx):
    path = fixture_docx("tracked_changes")

    assert find_text(str(path), "deleted text")["total_count"] == 0
    assert find_text(str(path), "inserted text")["total_count"] == 1


def test_find_text_reports_errors_as_before(tmp_path, fixture_docx):
    missing = tmp_path / "nope.docx"
    assert find_text(str(missing), "a") == {
        "error": f"Document {missing} does not exist"
    }
    assert find_text(str(fixture_docx("simple")), "") == {
        "error": "Search text cannot be empty"
    }


def test_find_text_in_document_returns_the_json_the_tool_always_returned(fixture_docx):
    path = fixture_docx("simple")

    payload = json.loads(_run(find_text_in_document(str(path), "simple fixture")))

    assert payload["query"] == "simple fixture"
    assert payload["total_count"] == 1
    assert payload["occurrences"][0]["paragraph_index"] == 1
    assert payload["occurrences"][0]["position"] == 23
