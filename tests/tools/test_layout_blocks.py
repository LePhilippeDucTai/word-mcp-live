"""Tests for the destructive block and layout tools, now backed by the engine.

Five tools share one failure mode, and it is the reason they are pinned
together here: each of them used to remove more than it was asked to.

``delete_paragraph`` took the section break of the paragraph it removed with
it, silently changing the page size, the margins and the header references of
everything above.  ``add_header_footer`` called ``p.clear()`` on every
paragraph of the story before writing, throwing away images and page-number
fields that had nothing to do with the new text.  ``add_bookmark`` drew its id
from ``random.randint(1000, 99999)`` -- an id space shared with revisions and
comment anchors -- and inserted its markers at index 0, in front of ``w:pPr``.
``replace_paragraph_block_below_header`` walked ``doc.paragraphs``, so a table
inside the block outlived the replacement, and
``replace_block_between_manual_anchors`` compared ``el.tag == CT_P.tag``, which
reads an unbound lxml attribute descriptor and never equals a string, so it
reported the start anchor as missing whatever the document held.

What is pinned is both halves of the contract: the API a caller sees -- the
parameters and the shape of every return message -- is unchanged, and what the
document held outside the block comes back untouched, section properties and
markers first.
"""

from __future__ import annotations

import asyncio
import copy
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document as PDocument

from tests.support.package_check import validate_package
from tests.support.snapshot import assert_unchanged_except, snapshot
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import visible_text
from word_document_server.engine.xmlns import qn
from word_document_server.tools.content_tools import (
    add_table_of_contents,
    delete_paragraph,
)
from word_document_server.tools.layout_tools import add_bookmark, add_header_footer
from word_document_server.utils.document_utils import (
    delete_block_under_header,
    replace_block_between_manual_anchors,
    replace_paragraph_block_below_header,
)
from word_document_server.utils.extended_document_utils import find_text

BLOCK_TAGS = (qn("w:p"), qn("w:tbl"), qn("w:sdt"))


def _run(coro: Any) -> Any:
    """Drive an ``async`` tool call to completion."""
    return asyncio.run(coro)


def _body(path: Path):
    return DocxPackage.open(str(path)).document.find(qn("w:body"))


def _blocks(path: Path) -> list:
    return [child for child in _body(path) if child.tag in BLOCK_TAGS]


def _texts(path: Path) -> list[str]:
    """Visible text of every block child of the body, tables included."""
    return [
        visible_text(block) if block.tag == qn("w:p") else "".join(block.itertext())
        for block in _blocks(path)
    ]


def _paragraph_texts(path: Path) -> list[str]:
    """Visible text of the paragraphs that are children of the body."""
    return [
        visible_text(child) for child in _body(path) if child.tag == qn("w:p")
    ]


def _all_paragraphs(path: Path) -> list:
    """Every ``w:p`` of the body, nesting included, in document order.

    Deliberately independent of the engine's own filter: on a document holding
    no table and no text box this is exactly the V2 index space the tools
    address, so a test can state what the space contains without asking the
    code under test what it thinks the space is.
    """
    return list(_body(path).iter(qn("w:p")))


def _all_paragraph_texts(path: Path) -> list[str]:
    """Visible text of every ``w:p`` of the body, nesting included."""
    return [visible_text(paragraph) for paragraph in _all_paragraphs(path)]


def _v2_index_of(path: Path, text: str) -> int:
    """Index ``find_text`` reports for the single paragraph holding `text`."""
    occurrences = find_text(str(path), text)["occurrences"]
    assert len(occurrences) == 1, occurrences
    return occurrences[0]["paragraph_index"]


def _section_break(paragraph) -> Any:
    return paragraph.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")


def _canonical(element) -> bytes:
    """Canonical form of an element, independent of where it now sits."""
    from lxml import etree

    return etree.tostring(element, method="c14n", exclusive=True, with_comments=False)


def _break_after(paragraph, body) -> None:
    """Give `paragraph` a section break copied from the body's own ``w:sectPr``."""
    properties = paragraph._p.get_or_add_pPr()
    properties.append(copy.deepcopy(body.find(qn("w:sectPr"))))


def _two_sections(path: Path) -> Path:
    """A document whose anchor and whose block both end a section."""
    document = PDocument()
    anchor = document.add_paragraph("Anchor paragraph.")
    middle = document.add_paragraph("Middle paragraph.")
    document.add_paragraph("Tail paragraph.")
    body = document.element.body
    _break_after(anchor, body)
    _break_after(middle, body)
    document.save(str(path))
    return path


# --------------------------------------------------------------------------------------
# delete_paragraph
# --------------------------------------------------------------------------------------


def test_deleting_a_paragraph_removes_that_paragraph_and_no_other(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = _paragraph_texts(path)

    result = _run(delete_paragraph(str(path), 1))

    assert result == "Paragraph at index 1 deleted successfully."
    assert _paragraph_texts(path) == before[:1] + before[2:]
    assert validate_package(path) == []


def test_a_deleted_paragraph_hands_its_section_break_to_the_one_above(
    fixture_docx,
) -> None:
    path = fixture_docx("sections")
    before = snapshot(path.read_bytes())
    carried = _canonical(_section_break(_blocks(path)[2]))
    assert _section_break(_blocks(path)[1]) is None  # sanity

    _run(delete_paragraph(str(path), 2))

    blocks = _blocks(path)
    moved = _section_break(blocks[1])
    assert moved is not None, "the section break must not go with the paragraph"
    assert _canonical(moved) == carried
    # The break stays a break: the body's own final w:sectPr is a different
    # element and describes the last section, which nothing here touched.
    assert snapshot(path.read_bytes()).sect_pr == before.sect_pr
    assert validate_package(path) == []


def test_the_section_break_lands_in_schema_order_inside_the_pPr(fixture_docx) -> None:
    path = fixture_docx("sections")
    # Give the paragraph above the break a w:pPr holding a w:jc, which the
    # schema puts *before* w:sectPr: appending would put the two in the wrong
    # order and Word offers to repair a document that gets that wrong.
    document = PDocument(str(path))
    document.paragraphs[1].alignment = 1
    document.save(str(path))

    _run(delete_paragraph(str(path), 2))

    properties = _blocks(path)[1].find(qn("w:pPr"))
    names = [child.tag for child in properties]
    assert names.index(qn("w:jc")) < names.index(qn("w:sectPr"))
    assert validate_package(path) == []


def test_the_body_section_properties_survive_deleting_the_last_paragraph(
    fixture_docx,
) -> None:
    path = fixture_docx("sections")
    before = snapshot(path.read_bytes())

    _run(delete_paragraph(str(path), 3))

    assert snapshot(path.read_bytes()).sect_pr == before.sect_pr
    assert validate_package(path) == []


def test_an_out_of_range_index_is_reported_and_changes_nothing(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = _run(delete_paragraph(str(path), 99))

    assert result == "Invalid paragraph index. Document has 5 paragraphs (0-4)."
    assert path.read_bytes() == before


def test_a_missing_document_is_reported_not_created(tmp_path) -> None:
    path = tmp_path / "absent.docx"

    assert _run(delete_paragraph(str(path), 0)) == f"Document {path} does not exist"
    assert not path.exists()


def test_a_paragraph_inside_a_content_control_is_in_the_index_space(
    fixture_docx,
) -> None:
    """The index space is the V2 one: block content control content counts.

    Index 1 of ``content_controls`` is the paragraph *inside* the block
    ``w:sdt``, not the second child of the body; deleting it must reach into
    the control rather than remove the body child that used to sit there.
    """
    path = fixture_docx("content_controls")
    nested = "Paragraph inside a block content control."
    children_before = _paragraph_texts(path)
    assert nested not in children_before  # sanity: it is not a child of the body

    _run(delete_paragraph(str(path), 1))

    assert _paragraph_texts(path) == children_before, "no body child may be removed"
    assert nested not in _all_paragraph_texts(path)
    assert any(block.tag == qn("w:sdt") for block in _blocks(path))
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# delete_paragraph / add_bookmark: one index space, shared with find_text
# --------------------------------------------------------------------------------------
#
# add_table_of_contents puts a block w:sdt in the body, so the milestone itself
# produces the documents where the two spaces part company: every paragraph of
# the table of contents counts in the V2 index find_text reports and none of
# them is a child of the body.  A tool that kept indexing body children would
# act on a paragraph that many positions further up, silently.


BODY_TEXTS = ("Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot")


def _document_with_a_table_of_contents(path: Path) -> Path:
    document = PDocument()
    document.add_heading("Titre", level=1)
    for text in BODY_TEXTS:
        document.add_paragraph(text)
    document.save(str(path))
    assert "Table of contents" in _run(add_table_of_contents(str(path)))
    assert any(block.tag == qn("w:sdt") for block in _blocks(path))  # sanity
    return path


def test_the_paragraph_deleted_is_the_one_find_text_reported(tmp_path) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    before = _all_paragraph_texts(path)
    index = _v2_index_of(path, "Charlie")

    result = _run(delete_paragraph(str(path), index))

    assert result == f"Paragraph at index {index} deleted successfully."
    assert _all_paragraph_texts(path) == [t for t in before if t != "Charlie"]
    assert validate_package(path) == []


def test_the_paragraph_bookmarked_is_the_one_find_text_reported(tmp_path) -> None:
    path = _document_with_a_table_of_contents(tmp_path / "toc.docx")
    index = _v2_index_of(path, "Charlie")

    result = _run(add_bookmark(str(path), index, "Marked"))

    assert json.loads(result)["success"] is True
    marked = [
        visible_text(paragraph)
        for paragraph in _body(path).iter(qn("w:p"))
        for marker in paragraph.findall(qn("w:bookmarkStart"))
        if marker.get(qn("w:name")) == "Marked"
    ]
    assert marked == ["Charlie"]
    assert validate_package(path) == []


def test_a_nested_paragraph_hands_its_section_break_outside_the_control(
    fixture_docx,
) -> None:
    """The break lands on the previous paragraph of the index space.

    Inside a content control the deleted paragraph has no preceding sibling
    ``w:p`` at all, so a carry-over that looked at siblings would drop the
    section break and take the page setup of everything above with it.
    """
    from lxml import etree

    path = fixture_docx("content_controls")
    pkg = DocxPackage.open(str(path))
    body = pkg.document.find(qn("w:body"))
    nested = body.find(f"{qn('w:sdt')}/{qn('w:sdtContent')}/{qn('w:p')}")
    assert nested is not None  # sanity
    assert next(iter(nested.itersiblings(qn("w:p"), preceding=True)), None) is None
    properties = etree.SubElement(nested, qn("w:pPr"))
    properties.append(copy.deepcopy(body.find(qn("w:sectPr"))))
    nested.insert(0, properties)
    pkg.save(str(path))
    carried = _canonical(_section_break(_all_paragraphs(path)[1]))

    _run(delete_paragraph(str(path), 1))

    above = _all_paragraphs(path)[0]
    moved = _section_break(above)
    assert moved is not None, "the section break must not go with the paragraph"
    assert _canonical(moved) == carried
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# add_bookmark
# --------------------------------------------------------------------------------------


def _bookmarks(path: Path) -> list[tuple[str, str]]:
    return [
        (element.get(qn("w:id")), element.get(qn("w:name")))
        for element in _body(path).iter(qn("w:bookmarkStart"))
    ]


def test_the_bookmark_id_comes_from_the_package_id_space(tmp_path) -> None:
    """``random.randint(1000, 99999)`` could not even reach past this document."""
    from lxml import etree

    path = tmp_path / "high-ids.docx"
    document = PDocument()
    paragraph = document.add_paragraph("Bookmarked already.")
    for tag, attrs in (
        ("w:bookmarkStart", {"w:id": "99999", "w:name": "High"}),
        ("w:bookmarkEnd", {"w:id": "99999"}),
    ):
        marker = etree.SubElement(paragraph._p, qn(tag))
        for name, value in attrs.items():
            marker.set(qn(name), value)
    document.save(str(path))

    _run(add_bookmark(str(path), 0, "Fresh"))

    added = [int(bm_id) for bm_id, name in _bookmarks(path) if name == "Fresh"]
    assert added == [100000]
    assert validate_package(path) == []


def test_the_bookmark_id_does_not_collide_with_the_ids_in_use(fixture_docx) -> None:
    path = fixture_docx("bookmarks")
    taken = {int(bm_id) for bm_id, _ in _bookmarks(path)}
    assert taken  # sanity: the fixture already uses the id space

    _run(add_bookmark(str(path), 1, "Fresh"))

    added = [int(bm_id) for bm_id, name in _bookmarks(path) if name == "Fresh"]
    assert len(added) == 1
    assert added[0] not in taken


def test_the_markers_frame_the_paragraph_and_leave_the_pPr_first(fixture_docx) -> None:
    path = fixture_docx("paragraph_styles")
    assert _blocks(path)[1].find(qn("w:pPr")) is not None  # sanity: it is styled

    _run(add_bookmark(str(path), 1, "Framed"))

    paragraph = _blocks(path)[1]
    tags = [child.tag for child in paragraph]
    assert tags[0] == qn("w:pPr"), "w:pPr must stay the first child of a w:p"
    assert tags[1] == qn("w:bookmarkStart")
    assert tags[-1] == qn("w:bookmarkEnd")
    assert validate_package(path) == []


def test_a_name_holding_xml_metacharacters_is_stored_intact(fixture_docx) -> None:
    path = fixture_docx("simple")
    name = 'Quote"And&Ampersand<Less'

    result = _run(add_bookmark(str(path), 1, name))

    assert json.loads(result)["success"] is True
    assert (name in [label for _, label in _bookmarks(path)]), _bookmarks(path)
    assert validate_package(path) == []


def test_the_bookmark_changes_nothing_but_its_own_paragraph(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())

    _run(add_bookmark(str(path), 1, "Probe"))

    assert_unchanged_except(
        before,
        snapshot(path.read_bytes()),
        paragraphs={1},
        counters={"bookmarks"},
    )


def test_the_bookmark_payload_keeps_its_shape(fixture_docx) -> None:
    path = fixture_docx("simple")

    result = _run(add_bookmark(str(path), 2, "Shape"))

    assert json.loads(result) == {
        "success": True,
        "bookmark_name": "Shape",
        "paragraph_index": 2,
    }


def test_a_paragraph_that_does_not_exist_is_reported(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = path.read_bytes()

    assert _run(add_bookmark(str(path), 99, "Nowhere")) == "Paragraph 99 does not exist."
    assert path.read_bytes() == before


# --------------------------------------------------------------------------------------
# add_header_footer
# --------------------------------------------------------------------------------------


def _story_runs(path: Path, story: str, index: int = 0):
    return snapshot(path.read_bytes()).paragraphs[(story, index)]


def test_the_header_text_is_rewritten_and_the_picture_survives(fixture_docx) -> None:
    path = fixture_docx("headers_footers")
    before = _story_runs(path, "header1")
    assert len(before.runs) >= 2  # sanity: text run + picture run

    result = _run(add_header_footer(str(path), header_text="New header text"))

    assert json.loads(result) == {"success": True, "added": ["header"], "section": 0}
    after = _story_runs(path, "header1")
    assert after.text == "New header text"
    assert len(after.runs) == len(before.runs), "the picture run must survive"
    assert validate_package(path) == []


def test_a_footer_holding_fields_is_refused_with_an_explicit_message(
    fixture_docx,
) -> None:
    path = fixture_docx("headers_footers")
    before = path.read_bytes()

    result = _run(add_header_footer(str(path), footer_text="Plain footer"))

    error = json.loads(result)["error"]
    assert "PAGE" in error, error
    assert "replace_content=False" in error
    assert path.read_bytes() == before, "a refused call must write nothing"


def test_a_refusal_discards_the_header_written_in_the_same_call(fixture_docx) -> None:
    path = fixture_docx("headers_footers")
    before = path.read_bytes()

    result = _run(
        add_header_footer(
            str(path), header_text="New header text", footer_text="Plain footer"
        )
    )

    assert "error" in json.loads(result)
    assert path.read_bytes() == before


def test_replace_content_false_adds_a_paragraph_and_keeps_the_fields(
    fixture_docx,
) -> None:
    path = fixture_docx("headers_footers")
    before = snapshot(path.read_bytes())
    fields_before = before.counters["fields"]

    result = _run(
        add_header_footer(str(path), footer_text="Plain footer", replace_content=False)
    )

    assert json.loads(result)["added"] == ["footer"]
    after = snapshot(path.read_bytes())
    assert after.counters["fields"] == fields_before
    assert after.paragraphs[("footer1", 0)].text == before.paragraphs[("footer1", 0)].text
    assert after.paragraphs[("footer1", 1)].text == "Plain footer"
    assert validate_package(path) == []


def test_the_mcp_tool_exposes_replace_content(fixture_docx) -> None:
    """The escape hatch the refusal names has to be reachable from a client.

    A footer holding a ``PAGE`` field -- what Word puts there by default -- is
    refused with "Pass replace_content=False", so a wrapper that did not
    forward the parameter turned that advice into advice no MCP client could
    follow, and made the footer impossible to write at all.
    """
    from word_document_server import main

    main.register_tools()
    tool = asyncio.run(main.mcp.get_tool("add_header_footer"))
    assert "replace_content" in tool.parameters["properties"]

    path = fixture_docx("headers_footers")
    result = _run(tool.fn(str(path), footer_text="Plain footer", replace_content=False))

    assert json.loads(result)["added"] == ["footer"]
    assert validate_package(path) == []


def test_the_header_call_touches_no_other_story(fixture_docx) -> None:
    path = fixture_docx("headers_footers")
    before = snapshot(path.read_bytes())

    _run(add_header_footer(str(path), header_text="New header text"))

    after = snapshot(path.read_bytes())
    for story in before.stories:
        if story == "header1":
            continue
        assert [sig.text for sig in before.story_paragraphs(story)] == [
            sig.text for sig in after.story_paragraphs(story)
        ], story


def test_an_out_of_range_section_is_reported(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = _run(add_header_footer(str(path), section_index=9, header_text="x"))

    assert result == "Section 9 does not exist."
    assert path.read_bytes() == before


# --------------------------------------------------------------------------------------
# replace_paragraph_block_below_header
# --------------------------------------------------------------------------------------


def test_a_table_under_the_header_is_removed_with_the_paragraphs(fixture_docx) -> None:
    path = fixture_docx("combined")
    before = snapshot(path.read_bytes())
    assert any(key[0] == "document" for key in before.tables)  # sanity

    result = replace_paragraph_block_below_header(
        str(path), "Fixture: tables", ["Only this now."]
    )

    assert "removed 2 elements" in result
    after = snapshot(path.read_bytes())
    assert [key for key in after.tables if key[0] == "document"] == []
    assert "Paragraph after the table." not in after.text()
    assert "Only this now." in after.text()
    assert validate_package(path) == []


def test_the_block_below_header_message_keeps_its_shape(fixture_docx) -> None:
    path = fixture_docx("simple")

    result = replace_paragraph_block_below_header(
        str(path), "Fixture: simple", ["One.", "Two."]
    )

    assert result == (
        "Replaced content under 'Fixture: simple' with 2 paragraph(s), "
        "style: Normal, removed 2 elements."
    )
    assert _paragraph_texts(path) == [
        "Fixture: simple",
        "One.",
        "Two.",
        "Second heading",
        "Body text under the second heading.",
    ]


def test_a_missing_header_is_reported_and_nothing_is_written(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = replace_paragraph_block_below_header(str(path), "No such header", ["x"])

    assert result == "Header 'No such header' not found in document."
    assert path.read_bytes() == before


def test_a_missing_file_is_reported(tmp_path) -> None:
    path = tmp_path / "absent.docx"

    assert (
        replace_paragraph_block_below_header(str(path), "h", ["x"])
        == f"Document {path} not found."
    )


def test_an_unknown_style_is_reported_and_nothing_is_written(fixture_docx) -> None:
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = replace_paragraph_block_below_header(
        str(path), "Fixture: simple", ["x"], new_paragraph_style="NoSuchStyle"
    )

    assert result == "Style 'NoSuchStyle' not found in document."
    assert path.read_bytes() == before


def test_the_style_named_is_the_style_written(fixture_docx) -> None:
    path = fixture_docx("paragraph_styles")

    replace_paragraph_block_below_header(
        str(path), "Fixture: paragraph styles", ["Styled."], new_paragraph_style="FixtureBody"
    )

    styles = [
        paragraph.style
        for paragraph in snapshot(path.read_bytes()).story_paragraphs("document")
    ]
    assert styles[1] == "FixtureBody"


def test_the_default_style_writes_no_pStyle(fixture_docx) -> None:
    """``Normal`` is the default paragraph style: python-docx writes no w:pStyle."""
    path = fixture_docx("simple")

    replace_paragraph_block_below_header(str(path), "Fixture: simple", ["Plain."])

    assert _blocks(path)[1].find(f"{qn('w:pPr')}/{qn('w:pStyle')}") is None


def test_the_section_break_of_the_block_moves_onto_the_replacement(
    fixture_docx,
) -> None:
    path = fixture_docx("sections")
    before = snapshot(path.read_bytes())
    carried = _canonical(_section_break(_blocks(path)[2]))

    result = replace_paragraph_block_below_header(
        str(path), "Fixture: sections", ["First.", "Second."]
    )

    assert "removed 3 elements" in result
    blocks = _blocks(path)
    assert [visible_text(block) for block in blocks] == [
        "Fixture: sections",
        "First.",
        "Second.",
    ]
    # The break still ends a section, and it still ends it *last*: on the last
    # paragraph before what followed the block.
    assert _canonical(_section_break(blocks[-1])) == carried
    assert snapshot(path.read_bytes()).sect_pr == before.sect_pr
    assert validate_package(path) == []


def test_a_heading_closes_the_block(fixture_docx) -> None:
    path = fixture_docx("simple")

    replace_paragraph_block_below_header(str(path), "Fixture: simple", [])

    assert _paragraph_texts(path) == [
        "Fixture: simple",
        "Second heading",
        "Body text under the second heading.",
    ]


def test_the_block_below_header_is_written_once(monkeypatch, fixture_docx) -> None:
    """The old implementation saved, reopened and saved again."""
    path = fixture_docx("simple")
    saves = []
    original = DocxPackage.save

    def counting(self, target):
        saves.append(str(target))
        return original(self, target)

    monkeypatch.setattr(DocxPackage, "save", counting)
    replace_paragraph_block_below_header(str(path), "Fixture: simple", ["One."])

    assert saves == [str(path)]


# --------------------------------------------------------------------------------------
# replace_block_between_manual_anchors
# --------------------------------------------------------------------------------------


def test_the_start_anchor_is_found(fixture_docx) -> None:
    """The bug of record: the anchor was never found, whatever the document."""
    path = fixture_docx("simple")

    result = replace_block_between_manual_anchors(
        str(path),
        "Second paragraph, with a trailing sentence.",
        ["Replacement."],
        end_anchor_text="Second heading",
    )

    assert "not found" not in result
    assert _paragraph_texts(path) == [
        "Fixture: simple",
        "First paragraph of the simple fixture.",
        "Second paragraph, with a trailing sentence.",
        "Replacement.",
        "Second heading",
        "Body text under the second heading.",
    ]


def test_the_between_anchors_message_keeps_its_shape(fixture_docx) -> None:
    path = fixture_docx("simple")

    result = replace_block_between_manual_anchors(
        str(path),
        "Fixture: simple",
        ["One."],
        end_anchor_text="Second heading",
    )

    assert result == (
        "Replaced content between 'Fixture: simple' and 'Second heading' with "
        "1 paragraph(s), style: Normal, removed 2 elements."
    )


def test_without_an_end_anchor_the_block_stops_at_the_next_heading(
    fixture_docx,
) -> None:
    path = fixture_docx("simple")

    result = replace_block_between_manual_anchors(
        str(path), "Fixture: simple", ["One."]
    )

    assert "next logical header" in result
    assert _paragraph_texts(path) == [
        "Fixture: simple",
        "One.",
        "Second heading",
        "Body text under the second heading.",
    ]


def test_an_unknown_end_anchor_is_reported_and_nothing_is_removed(
    fixture_docx,
) -> None:
    """The old code deleted to the end of the document instead."""
    path = fixture_docx("simple")
    before = path.read_bytes()

    result = replace_block_between_manual_anchors(
        str(path), "Fixture: simple", ["One."], end_anchor_text="No such anchor"
    )

    assert result == "End anchor 'No such anchor' not found."
    assert path.read_bytes() == before


def test_a_table_between_the_anchors_is_removed(fixture_docx) -> None:
    path = fixture_docx("combined")

    result = replace_block_between_manual_anchors(
        str(path),
        "Fixture: tables",
        ["Only this now."],
        end_anchor_text="Fixture: comments",
    )

    assert "removed 2 elements" in result
    after = snapshot(path.read_bytes())
    assert [key for key in after.tables if key[0] == "document"] == []
    assert validate_package(path) == []


def test_the_body_section_properties_are_never_removed(fixture_docx) -> None:
    path = fixture_docx("sections")
    before = snapshot(path.read_bytes())

    replace_block_between_manual_anchors(
        str(path), "First section, portrait, default margins.", ["Tail."]
    )

    after = snapshot(path.read_bytes())
    assert after.sect_pr == before.sect_pr
    assert validate_package(path) == []


def test_a_section_break_with_no_paragraph_to_carry_it_is_refused(tmp_path) -> None:
    path = _two_sections(tmp_path / "two-sections.docx")
    before = path.read_bytes()

    result = replace_block_between_manual_anchors(
        str(path), "Anchor paragraph.", [], end_anchor_text="Tail paragraph."
    )

    assert result == (
        "Refusing to replace the block: it holds 1 section break(s) and only "
        "0 replacement paragraph(s) to carry them."
    )
    assert path.read_bytes() == before


def test_a_section_break_lands_on_the_replacement_when_the_anchor_is_taken(
    tmp_path,
) -> None:
    path = _two_sections(tmp_path / "two-sections.docx")
    anchor_break = _canonical(_section_break(_blocks(path)[0]))

    result = replace_block_between_manual_anchors(
        str(path), "Anchor paragraph.", ["New middle."], end_anchor_text="Tail paragraph."
    )

    assert "removed 1 elements" in result
    blocks = _blocks(path)
    assert [visible_text(block) for block in blocks] == [
        "Anchor paragraph.",
        "New middle.",
        "Tail paragraph.",
    ]
    assert _canonical(_section_break(blocks[0])) == anchor_break
    assert _section_break(blocks[1]) is not None
    assert validate_package(path) == []


def test_a_match_fn_drives_both_anchors(fixture_docx) -> None:
    path = fixture_docx("simple")

    def match(text, element, is_end=False):
        return text.startswith("Second heading") if is_end else text.startswith("Fixture:")

    result = replace_block_between_manual_anchors(
        str(path), "ignored", ["One."], match_fn=match
    )

    assert "removed 2 elements" in result
    assert _paragraph_texts(path) == [
        "Fixture: simple",
        "One.",
        "Second heading",
        "Body text under the second heading.",
    ]


def test_the_block_between_anchors_is_written_once(monkeypatch, fixture_docx) -> None:
    path = fixture_docx("simple")
    saves = []
    original = DocxPackage.save

    def counting(self, target):
        saves.append(str(target))
        return original(self, target)

    monkeypatch.setattr(DocxPackage, "save", counting)
    replace_block_between_manual_anchors(
        str(path), "Fixture: simple", ["One."], end_anchor_text="Second heading"
    )

    assert saves == [str(path)]


@pytest.mark.parametrize("name", ["simple", "combined", "sections"])
def test_the_package_stays_valid(fixture_docx, name: str) -> None:
    path = fixture_docx(name)

    replace_paragraph_block_below_header(str(path), "Fixture: simple", ["One.", "Two."])
    replace_block_between_manual_anchors(str(path), "Fixture: simple", ["Three."])

    assert validate_package(path) == []
    assert PDocument(str(path)).paragraphs


# --------------------------------------------------------------------------------------
# delete_block_under_header
# --------------------------------------------------------------------------------------


def test_delete_block_under_header_removes_the_blocks_in_place(fixture_docx) -> None:
    """Kept for callers holding an open python-docx document; it does not save."""
    path = fixture_docx("combined")
    before = path.read_bytes()
    document = PDocument(str(path))

    header, removed = delete_block_under_header(document, "Fixture: tables")

    assert removed == 2, "the table under the header counts as a block"
    assert header is not None
    assert visible_text(header) == "Fixture: tables"
    assert path.read_bytes() == before, "the helper mutates in place and saves nothing"
    document.save(str(path))
    assert [key for key in snapshot(path.read_bytes()).tables if key[0] == "document"] == []


def test_delete_block_under_header_reports_a_missing_header(fixture_docx) -> None:
    path = fixture_docx("simple")
    document = PDocument(str(path))

    assert delete_block_under_header(document, "No such header") == (None, 0)
