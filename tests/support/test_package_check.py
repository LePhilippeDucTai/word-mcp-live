"""Tests for tests/support/package_check.py.

Two kinds of coverage:
- structural checks exercised directly on hand-built (and hand-corrupted)
  .docx packages, which never require LibreOffice;
- a sanity check that LibreOffice-produced fixtures are themselves clean,
  guarded by `requires_libreoffice` since LibreOffice stays optional.
"""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from tests.fixtures.builders import build
from tests.support.libreoffice import fodt_to_docx, requires_libreoffice
from tests.support.package_check import CT_NS, REL_NS, W_NS, validate_package

_FIXTURE_NAMES = ["rich", "tables_lists", "notes_fields"]


def _make_base_docx(tmp_path: Path, name: str = "base.docx") -> Path:
    doc = Document()
    doc.add_paragraph("Hello world.")
    doc.add_paragraph("Second paragraph.")
    path = tmp_path / name
    doc.save(path)
    return path


def _read_zip(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _parts_from_bytes(blob: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _remove_element(
    parts: dict[str, bytes], part_name: str, tag: str, id_attr: str, id_value: str
) -> None:
    """Remove the first `tag` element whose `id_attr` equals `id_value` from `part_name`."""
    root = etree.fromstring(parts[part_name])
    for el in root.iter(f"{{{W_NS}}}{tag}"):
        if el.get(f"{{{W_NS}}}{id_attr}") == id_value:
            el.getparent().remove(el)
            parts[part_name] = _serialize(root)
            return
    raise AssertionError(f"no <{tag}> with {id_attr}={id_value!r} found in {part_name}")


def _write_zip(path: Path, parts: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)


def _serialize(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_bookmark(parts: dict[str, bytes], bookmark_id: str, name: str) -> None:
    """Insert a bookmarkStart/bookmarkEnd pair into the first paragraph."""
    root = etree.fromstring(parts["word/document.xml"])
    first_p = root.find(f"{{{W_NS}}}body/{{{W_NS}}}p")
    start = etree.SubElement(first_p, f"{{{W_NS}}}bookmarkStart")
    start.set(f"{{{W_NS}}}id", bookmark_id)
    start.set(f"{{{W_NS}}}name", name)
    end = etree.SubElement(first_p, f"{{{W_NS}}}bookmarkEnd")
    end.set(f"{{{W_NS}}}id", bookmark_id)
    parts["word/document.xml"] = _serialize(root)


def _add_comment(parts: dict[str, bytes], comment_id: str) -> None:
    """Add a commentReference in the body plus a matching word/comments.xml part."""
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find(f"{{{W_NS}}}body")
    p = etree.SubElement(body, f"{{{W_NS}}}p")
    r = etree.SubElement(p, f"{{{W_NS}}}r")
    ref = etree.SubElement(r, f"{{{W_NS}}}commentReference")
    ref.set(f"{{{W_NS}}}id", comment_id)
    sect_pr = body.find(f"{{{W_NS}}}sectPr")
    if sect_pr is not None:
        body.remove(p)
        sect_pr.addprevious(p)
    parts["word/document.xml"] = _serialize(root)

    comments_root = etree.Element(f"{{{W_NS}}}comments", nsmap={"w": W_NS})
    comment_el = etree.SubElement(comments_root, f"{{{W_NS}}}comment")
    comment_el.set(f"{{{W_NS}}}id", comment_id)
    comment_el.set(f"{{{W_NS}}}author", "Tester")
    cp = etree.SubElement(comment_el, f"{{{W_NS}}}p")
    cr = etree.SubElement(cp, f"{{{W_NS}}}r")
    ct_el = etree.SubElement(cr, f"{{{W_NS}}}t")
    ct_el.text = "A comment."
    parts["word/comments.xml"] = _serialize(comments_root)

    ct_root = etree.fromstring(parts["[Content_Types].xml"])
    override = etree.SubElement(ct_root, f"{{{CT_NS}}}Override")
    override.set("PartName", "/word/comments.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    parts["[Content_Types].xml"] = _serialize(ct_root)

    rels_root = etree.fromstring(parts["word/_rels/document.xml.rels"])
    rel = etree.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
    rel.set("Id", "rIdComments")
    rel.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    )
    rel.set("Target", "comments.xml")
    parts["word/_rels/document.xml.rels"] = _serialize(rels_root)


def test_validate_package_clean_document_has_no_issues(tmp_path):
    path = _make_base_docx(tmp_path)
    assert validate_package(path) == []


def test_validate_package_bookmark_and_comment_are_coherent(tmp_path):
    path = _make_base_docx(tmp_path)
    parts = _read_zip(path)
    _add_bookmark(parts, bookmark_id="0", name="MyBookmark")
    _add_comment(parts, comment_id="0")
    _write_zip(path, parts)

    assert validate_package(path) == []


def test_validate_package_libreoffice_note_disposition_is_clean(tmp_path):
    """LibreOffice reserves ids 0/1 for its separator notes; real notes start at 2."""
    path = tmp_path / "libreoffice_notes.docx"
    parts = _parts_from_bytes(build("footnotes"))
    footnotes_root = etree.fromstring(parts["word/footnotes.xml"])
    for el in footnotes_root.iter(f"{{{W_NS}}}footnote"):
        note_id = el.get(f"{{{W_NS}}}id")
        if note_id == "-1":
            el.set(f"{{{W_NS}}}id", "0")
        elif note_id == "0":
            el.set(f"{{{W_NS}}}id", "1")
    parts["word/footnotes.xml"] = _serialize(footnotes_root)
    _write_zip(path, parts)

    assert validate_package(path) == []


def test_validate_package_word_note_disposition_is_clean(tmp_path):
    """Word reserves ids -1/0 for its separator notes; real notes start at 1."""
    path = tmp_path / "word_notes.docx"
    parts = _parts_from_bytes(build("footnotes"))
    footnotes_root = etree.fromstring(parts["word/footnotes.xml"])
    for el in footnotes_root.iter(f"{{{W_NS}}}footnote"):
        note_id = el.get(f"{{{W_NS}}}id")
        if note_id == "2":
            el.set(f"{{{W_NS}}}id", "1")
    parts["word/footnotes.xml"] = _serialize(footnotes_root)
    document_root = etree.fromstring(parts["word/document.xml"])
    for el in document_root.iter(f"{{{W_NS}}}footnoteReference"):
        if el.get(f"{{{W_NS}}}id") == "2":
            el.set(f"{{{W_NS}}}id", "1")
    parts["word/document.xml"] = _serialize(document_root)
    _write_zip(path, parts)

    assert validate_package(path) == []


def test_validate_package_detects_missing_relationship_target(tmp_path):
    path = _make_base_docx(tmp_path)
    parts = _read_zip(path)
    rels_root = etree.fromstring(parts["word/_rels/document.xml.rels"])
    styles_rel = next(
        rel
        for rel in rels_root.findall(f"{{{REL_NS}}}Relationship")
        if rel.get("Target") == "styles.xml"
    )
    styles_rel.set("Target", "styles-does-not-exist.xml")
    parts["word/_rels/document.xml.rels"] = _serialize(rels_root)
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(i.code == "REL-TARGET-MISSING" for i in issues), issues


def test_validate_package_detects_duplicate_id(tmp_path):
    path = _make_base_docx(tmp_path)
    parts = _read_zip(path)
    _add_bookmark(parts, bookmark_id="0", name="BM1")
    _add_bookmark(parts, bookmark_id="0", name="BM2")
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(i.code == "ID-DUPLICATE" for i in issues), issues


def test_validate_package_detects_part_without_content_type(tmp_path):
    path = _make_base_docx(tmp_path)
    parts = _read_zip(path)
    parts["word/extra.bin"] = b"\x00\x01\x02"
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(i.code == "CT-MISSING" and i.part == "word/extra.bin" for i in issues), issues


def test_validate_package_detects_note_missing(tmp_path):
    """A footnoteReference pointing at a note that no longer exists in footnotes.xml."""
    path = tmp_path / "note_missing.docx"
    parts = _parts_from_bytes(build("footnotes"))
    _remove_element(parts, "word/footnotes.xml", "footnote", "id", "3")
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "NOTE-MISSING" and i.part == "word/footnotes.xml" for i in issues
    ), issues


def test_validate_package_detects_note_orphan(tmp_path):
    """A footnote body that no footnoteReference in the document points to."""
    path = tmp_path / "note_orphan.docx"
    parts = _parts_from_bytes(build("footnotes"))
    _remove_element(parts, "word/document.xml", "footnoteReference", "id", "3")
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "NOTE-ORPHAN" and i.part == "word/footnotes.xml" for i in issues
    ), issues


def test_validate_package_detects_comment_missing(tmp_path):
    """A commentReference pointing at a comment that no longer exists in comments.xml."""
    path = tmp_path / "comment_missing.docx"
    parts = _parts_from_bytes(build("comments"))
    _remove_element(parts, "word/comments.xml", "comment", "id", "3")
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "COMMENT-MISSING" and i.part == "word/comments.xml" for i in issues
    ), issues


def test_validate_package_detects_comment_orphan(tmp_path):
    """A comment that no commentReference in the document points to."""
    path = tmp_path / "comment_orphan.docx"
    parts = _parts_from_bytes(build("comments"))
    _remove_element(parts, "word/document.xml", "commentReference", "id", "3")
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "COMMENT-ORPHAN" and i.part == "word/comments.xml" for i in issues
    ), issues


def test_validate_package_detects_numid_unresolved(tmp_path):
    """A numId used by a paragraph with no matching <w:num> definition."""
    path = tmp_path / "numid_unresolved.docx"
    parts = _parts_from_bytes(build("complex_numbering"))
    numbering_root = etree.fromstring(parts["word/numbering.xml"])
    for el in numbering_root.findall(f"{{{W_NS}}}num"):
        if el.get(f"{{{W_NS}}}numId") == "902":
            numbering_root.remove(el)
            break
    else:
        raise AssertionError("no <w:num numId='902'> found in fixture")
    parts["word/numbering.xml"] = _serialize(numbering_root)
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "NUMID-UNRESOLVED" and i.part == "word/numbering.xml" for i in issues
    ), issues


def test_validate_package_detects_main_part_missing(tmp_path):
    """A package whose root .rels carries no officeDocument relationship."""
    path = tmp_path / "main_part_missing.docx"
    parts = _parts_from_bytes(build("simple"))
    rels_root = etree.fromstring(parts["_rels/.rels"])
    for rel in rels_root.findall(f"{{{REL_NS}}}Relationship"):
        if (rel.get("Type") or "").endswith("/officeDocument"):
            rels_root.remove(rel)
            break
    else:
        raise AssertionError("no officeDocument relationship found in fixture")
    parts["_rels/.rels"] = _serialize(rels_root)
    _write_zip(path, parts)

    issues = validate_package(path)
    assert any(
        i.code == "MAIN-PART-MISSING" and i.part == "_rels/.rels" for i in issues
    ), issues


def test_validate_package_missing_file_reports_issue(tmp_path):
    issues = validate_package(tmp_path / "does-not-exist.docx")
    assert len(issues) == 1
    assert issues[0].code == "ZIP-UNREADABLE"


def test_validate_package_flags_stray_temp_file(tmp_path):
    path = _make_base_docx(tmp_path)
    (tmp_path / "~$base.docx").write_bytes(b"lock file")

    issues = validate_package(path)
    assert any(i.code == "TEMP-FILE-PRESENT" for i in issues), issues


@requires_libreoffice
@pytest.mark.libreoffice
@pytest.mark.parametrize("name", _FIXTURE_NAMES)
def test_validate_package_libreoffice_fixtures_are_clean(name, tmp_path_factory):
    cache_dir = tmp_path_factory.mktemp("odt_docx_cache_pkgcheck")
    docx_path = fodt_to_docx(name, cache_dir)

    issues = validate_package(docx_path)

    assert issues == [], issues
