"""Tests for :mod:`word_document_server.engine.textmodel`.

Two layers, on purpose.

The *tables* below spell out, fixture by fixture and paragraph by paragraph, the
text the engine must read.  They are written from the XML in
``tests/fixtures/builders.py``, not from the engine's output: a policy change --
a container forgotten, an opaque element that starts leaking its text -- has to
show up as a diff against a literal somebody wrote down, otherwise the tests
only pin whatever the code happens to do.

The *properties* then run over every paragraph of every story of every fixture:
segments tile the paragraph, offsets round-trip through ``position``, and no
function mutates the tree.  Those hold whatever the tables say, and they are what
``ranges`` will rely on.
"""

from __future__ import annotations

import io
import itertools
from pathlib import Path
from typing import get_args

import pytest
from docx import Document
from lxml import etree

from tests.fixtures.builders import ALL_FIXTURES, build
from word_document_server.engine.errors import LocatorError
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import (
    SEGMENT_KINDS,
    Segment,
    SegmentKind,
    fields,
    position,
    segments,
    visible_text,
)
from word_document_server.engine.xmlns import NAMESPACES, qn

# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------

W_P = qn("w:p")


def paragraph_from(children: str) -> etree._Element:
    """Return a standalone ``w:p`` whose children are the given XML fragment."""
    declarations = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in NAMESPACES.items())
    return etree.fromstring(f"<w:p {declarations}>{children}</w:p>")


def story_paragraphs(name: str) -> dict[str, list[etree._Element]]:
    """Return every ``w:p`` of every story of fixture `name`, in document order."""
    package = DocxPackage.open(build(name))
    return {story: list(root.iter(W_P)) for story, root in package.stories()}


def body_paragraphs(name: str) -> list[etree._Element]:
    """Return every ``w:p`` of the main document part of fixture `name`."""
    return story_paragraphs(name)["document"]


def all_paragraphs(name: str) -> list[tuple[str, int, etree._Element]]:
    """Return ``(story, index, paragraph)`` for every paragraph of fixture `name`."""
    return [
        (story, index, paragraph)
        for story, paragraphs in story_paragraphs(name).items()
        for index, paragraph in enumerate(paragraphs)
    ]


def texts(paragraphs: list[etree._Element]) -> list[str]:
    return [visible_text(paragraph) for paragraph in paragraphs]


def kinds(paragraph: etree._Element) -> list[str]:
    return [segment.kind for segment in segments(paragraph)]


def marks(paragraph: etree._Element, kind: str) -> list[tuple[str, int]]:
    """Return ``(local tag name, offset)`` for every segment of `kind`."""
    return [
        (etree.QName(segment.element).localname, segment.start)
        for segment in segments(paragraph)
        if segment.kind == kind
    ]


def field_tuples(paragraph: etree._Element) -> list[tuple[str, int, int]]:
    return [(entry.instr, entry.start, entry.end) for entry in fields(paragraph)]


# --------------------------------------------------------------------------------------
# The expected reading, fixture by fixture
# --------------------------------------------------------------------------------------

#: Visible text of every ``w:p`` of the main document part, in document order
#: (table cells and content controls included).  Transcribed from the fragments in
#: ``tests/fixtures/builders.py``.
EXPECTED_BODY_TEXT: dict[str, list[str]] = {
    "simple": [
        "Fixture: simple",
        "First paragraph of the simple fixture.",
        "Second paragraph, with a trailing sentence.",
        "Second heading",
        "Body text under the second heading.",
    ],
    "mixed_runs": [
        "Fixture: mixed runs",
        # Four runs split by run properties, read as one flow.
        "Plain then bold bold italic et la fin.",
        # Three runs split by rsid only.
        "One sentence cut in three runs.",
        # w:tab -> \t, w:br -> \n, and a second run appended with no separator.
        "Before\tafter tab\nafter breakno proofing",
    ],
    "paragraph_styles": [
        "Fixture: paragraph styles",
        "Custom body style.",
        "Custom note style.",
        "Built-in Quote style.",
        "Built-in ListParagraph style.",
        "Styled and directly centered.",
    ],
    "character_styles": [
        "Fixture: character styles",
        "Plain, then emphasised, then code_span(), then built-in Strong.",
        "Style plus direct bold underline.",
    ],
    "style_inheritance": [
        "Fixture: style inheritance",
        "Root of the chain.",
        "Branch, based on root, linked to a character style.",
        "Leaf, based on branch.",
        "Linked character style used on its own.",
    ],
    "themes": [
        "Fixture: themes",
        "Major theme font. Accent 1. Accent 1 tinted. Accent 1 shaded.",
        "Paragraph shaded with a theme fill.",
    ],
    "complex_numbering": [
        "Fixture: complex numbering",
        # The list label lives in numbering.xml and is not paragraph content.
        "First top level item",
        "Nested letter item",
        "Nested roman item",
        "Second top level item",
        "Overridden list restarting at five",
        "Bulleted item",
    ],
    "tables": [
        "Fixture: tables",
        "Merged header across two columns",
        "Third column",
        "Vertically merged cell",
        "Row two, column two",
        "Nested A",
        "Nested B",
        "",  # the paragraph that must follow a nested table inside its cell
        "",  # the continuation cell of the vertical merge
        "Row three, column two",
        "Row three, column three",
        "Paragraph after the table.",
    ],
    "comments": [
        "Fixture: comments",
        # commentRangeStart/End and the commentReference runs are zero width.
        "This sentence carries a comment thread. The rest of the paragraph is not commented.",
        "The third comment starts here",
        "and ends in the next paragraph.",
    ],
    "tracked_changes": [
        "Fixture: tracked changes",
        # The insertion is read, the deletion is not.
        "Kept text, inserted text, and kept tail.",
        # Inserted then deleted: still deleted, so the two spaces meet.
        "Before.  After.",
        "This run became bold, it used to be italic.",
        "This paragraph became centered.",
        # A deleted paragraph mark lives in w:pPr/w:rPr and is not content.
        "This paragraph mark is deleted, so this merges with the next one.",
        "This paragraph mark is inserted.",
    ],
    "hyperlinks": [
        "Fixture: hyperlinks",
        "An external link: example.org.",
        "An internal link: jump to the anchor.",
        "Anchor target paragraph.",
    ],
    "bookmarks": [
        "Fixture: bookmarks",
        "Before. Bookmarked span. After.",
        "A bookmark opens here",
        "and closes in this paragraph.",
        # An empty bookmark leaves the two surrounding spaces adjacent.
        "An empty bookmark sits here:  and text resumes.",
        "Bookmark whose name follows the reserved underscore convention.",
    ],
    "fields": [
        "Fixture: fields",
        # fldSimple: the cached result shows, the instruction does not.
        "Page 1 of the document.",
        # Complex field: fldChar and instrText are zero width, the result shows.
        "Fixture: fields\t1",
        # Nested field: the inner cached result sits inside the outer instruction
        # and is stored as an ordinary w:t, so the engine reports it.
        "1first page",
        "Referenced paragraph.",
        "Referenced paragraph.",
    ],
    "footnotes": [
        "Fixture: footnotes",
        # footnoteReference and endnoteReference are opaque.
        "A sentence with a footnote and a second one.",
        "A sentence with an endnote.",
    ],
    "headers_footers": [
        "Fixture: headers and footers",
        "Body text; the header and footer carry the interesting parts.",
    ],
    "sections": [
        "Fixture: sections",
        "First section, portrait, default margins.",
        "",  # the paragraph carrying the section break
        "Second section, landscape, narrow margins.",
    ],
    "content_controls": [
        "Fixture: content controls",
        "Paragraph inside a block content control.",
        "Before the control, inline control content, after the control.",
        "Alpha",
    ],
    "drawings": [
        "Fixture: drawings",
        "An inline image follows.",
        "",  # a paragraph holding nothing but the picture
        # w:drawing is opaque, so the two surrounding spaces meet.
        "A second inline image, inside a text paragraph:  and text after it.",
    ],
}

#: Visible text of the stories other than the body, for the fixtures that have any.
EXPECTED_OTHER_STORY_TEXT: dict[str, dict[str, list[str]]] = {
    "footnotes": {
        # w:separator and w:continuationSeparator are opaque; w:footnoteRef and
        # w:endnoteRef -- the auto-numbering mark -- too.
        "footnotes": ["", "", " First footnote body.", " Second footnote body, with bold text."],
        "endnotes": ["", "", " Endnote body."],
    },
    "headers_footers": {
        "header1": ["Fixture header "],  # the inline image contributes nothing
        "header2": ["Fixture first-page header"],
        "footer1": ["Page 1 of 1"],  # two PAGE/NUMPAGES fldSimple results
    },
}

#: ``combined`` runs every other populator on one document, in this order.
COMBINED_CONSTITUENTS = [
    "simple",
    "mixed_runs",
    "paragraph_styles",
    "character_styles",
    "style_inheritance",
    "themes",
    "complex_numbering",
    "tables",
    "comments",
    "tracked_changes",
    "hyperlinks",
    "bookmarks",
    "fields",
    "footnotes",
    "content_controls",
    "drawings",
    "headers_footers",
    "sections",
]

#: ``(fixture, story, paragraph index)`` -> the fields the engine must report as
#: ``(instruction, start, end)``.  Every other paragraph must report none.
EXPECTED_FIELDS: dict[tuple[str, str, int], list[tuple[str, int, int]]] = {
    ("fields", "document", 1): [(r" PAGE   \* MERGEFORMAT ", 5, 6)],
    ("fields", "document", 2): [(r' TOC \o "1-3" \h \z \u ', 0, 17)],
    # The IF field's own instruction is the concatenation of its two instrText
    # runs; the PAGE field embedded between them is reported separately.
    ("fields", "document", 3): [
        (r' IF  = 1 "first page" "later page" ', 0, 11),
        (" PAGE ", 0, 1),
    ],
    ("fields", "document", 5): [(r" REF FixtureRefTarget \h ", 0, 21)],
    ("headers_footers", "footer1", 0): [(" PAGE ", 5, 6), (" NUMPAGES ", 10, 11)],
}


# --------------------------------------------------------------------------------------
# The tables
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("name", sorted(EXPECTED_BODY_TEXT))
def test_body_text_matches_the_expected_table(name: str) -> None:
    assert texts(body_paragraphs(name)) == EXPECTED_BODY_TEXT[name]


@pytest.mark.parametrize("name", sorted(EXPECTED_OTHER_STORY_TEXT))
def test_other_stories_match_the_expected_table(name: str) -> None:
    paragraphs = story_paragraphs(name)
    for story, expected in EXPECTED_OTHER_STORY_TEXT[name].items():
        assert texts(paragraphs[story]) == expected, story


def test_the_table_covers_every_fixture() -> None:
    """A new fixture must not slip in without an expected reading."""
    assert set(EXPECTED_BODY_TEXT) | {"combined"} == set(ALL_FIXTURES)


def test_combined_reads_as_its_constituents_concatenated() -> None:
    """``combined`` is the other populators run on one document, so its reading
    is theirs in order -- give or take the empty paragraphs sections introduce."""
    expected = [
        text
        for name in COMBINED_CONSTITUENTS
        for text in EXPECTED_BODY_TEXT[name]
        if text
    ]
    assert [text for text in texts(body_paragraphs("combined")) if text] == expected


def test_engine_reads_what_python_docx_drops() -> None:
    """The reason this module exists, pinned on the two constructs that break
    ``Paragraph.text``: a tracked insertion and an inline content control."""
    tracked = Document(io.BytesIO(build("tracked_changes"))).paragraphs[1]
    assert tracked.text == "Kept text, and kept tail."
    assert visible_text(tracked) == "Kept text, inserted text, and kept tail."

    controlled = Document(io.BytesIO(build("content_controls"))).paragraphs[1]
    assert controlled.text == "Before the control, , after the control."
    assert visible_text(controlled) == "Before the control, inline control content, after the control."


# --------------------------------------------------------------------------------------
# Fields
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("name", sorted(set(ALL_FIXTURES) - {"combined"}))
def test_fields_match_the_expected_table(name: str) -> None:
    for story, paragraphs in story_paragraphs(name).items():
        for index, paragraph in enumerate(paragraphs):
            expected = EXPECTED_FIELDS.get((name, story, index), [])
            assert field_tuples(paragraph) == expected, f"{name}/{story}[{index}]"


def test_combined_reports_every_instruction_its_constituents_carry() -> None:
    """Paragraph indices shift in ``combined``, but no field may be lost or
    invented by the surrounding markup."""
    found = sorted(
        entry.instr
        for _, _, paragraph in all_paragraphs("combined")
        for entry in fields(paragraph)
    )
    assert found == sorted(instr for row in EXPECTED_FIELDS.values() for instr, _, _ in row)


@pytest.mark.parametrize("name", sorted(ALL_FIXTURES))
def test_field_spans_stay_inside_the_paragraph(name: str) -> None:
    for paragraphs in story_paragraphs(name).values():
        for paragraph in paragraphs:
            length = len(visible_text(paragraph))
            for entry in fields(paragraph):
                assert entry.atomic is True
                assert 0 <= entry.start <= entry.end <= length


def test_nested_field_spans_nest() -> None:
    outer, inner = fields(body_paragraphs("fields")[3])
    assert outer.start <= inner.start and inner.end <= outer.end


def test_field_instruction_text_is_never_visible() -> None:
    """The TOC switches must not leak into the paragraph's text."""
    paragraph = body_paragraphs("fields")[2]
    assert "TOC" not in visible_text(paragraph)
    instructions = [
        segment
        for segment in segments(paragraph)
        if segment.element.tag == qn("w:instrText")
    ]
    assert [segment.kind for segment in instructions] == ["opaque"]
    assert all(segment.start == segment.end for segment in instructions)


def test_field_opened_in_an_earlier_paragraph_is_reported_from_zero() -> None:
    paragraph = paragraph_from(
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>7</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    assert field_tuples(paragraph) == [(" PAGE ", 0, 1)]


def test_field_left_open_is_reported_up_to_the_paragraph_end() -> None:
    paragraph = paragraph_from(
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>result</w:t></w:r>"
    )
    assert field_tuples(paragraph) == [(" TOC ", 0, 6)]


def test_nested_fld_simple_reports_outer_first() -> None:
    paragraph = paragraph_from(
        '<w:fldSimple w:instr=" IF ">'
        '  <w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
        "</w:fldSimple>"
    )
    assert field_tuples(paragraph) == [(" IF ", 0, 1), (" PAGE ", 0, 1)]


def test_deleted_instruction_stays_out_of_the_instruction() -> None:
    paragraph = paragraph_from(
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:del w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">'
        '  <w:r><w:delInstrText xml:space="preserve"> NUMPAGES </w:delInstrText></w:r>'
        "</w:del>"
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>1</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    assert field_tuples(paragraph) == [(" PAGE ", 0, 1)]
    assert visible_text(paragraph) == "1"


# --------------------------------------------------------------------------------------
# Tracked changes
# --------------------------------------------------------------------------------------


def test_deleted_text_is_hidden_but_still_reachable() -> None:
    paragraph = body_paragraphs("tracked_changes")[1]
    assert "deleted text" not in visible_text(paragraph)
    hidden = [segment for segment in segments(paragraph) if segment.kind == "hidden"]
    assert [segment.source_text for segment in hidden] == ["deleted text, "]
    assert all(segment.start == segment.end and segment.text == "" for segment in hidden)


def test_inserted_text_is_visible() -> None:
    paragraph = body_paragraphs("tracked_changes")[1]
    inserted = [
        segment
        for segment in segments(paragraph)
        if any(container.tag == qn("w:ins") for container in segment.containers)
    ]
    assert [segment.text for segment in inserted] == ["inserted text, "]
    assert [segment.kind for segment in inserted] == ["text"]


def test_text_inserted_then_deleted_stays_hidden() -> None:
    paragraph = body_paragraphs("tracked_changes")[2]
    assert visible_text(paragraph) == "Before.  After."
    hidden = [segment for segment in segments(paragraph) if segment.kind == "hidden"]
    assert [segment.source_text for segment in hidden] == ["inserted then deleted"]


def test_deleted_paragraph_mark_is_not_content() -> None:
    """``w:pPr/w:rPr/w:del`` marks the paragraph mark, not any text."""
    paragraph = body_paragraphs("tracked_changes")[5]
    assert visible_text(paragraph) == (
        "This paragraph mark is deleted, so this merges with the next one."
    )
    assert "hidden" not in kinds(paragraph)


def test_moved_text_leaves_only_the_destination_visible() -> None:
    stamp = 'w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z"'
    paragraph = paragraph_from(
        f"<w:moveFrom {stamp}><w:r><w:t>moved away</w:t></w:r></w:moveFrom>"
        f'<w:moveTo w:id="2" w:author="A" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:t>moved here</w:t></w:r></w:moveTo>"
    )
    assert visible_text(paragraph) == "moved here"
    hidden = [segment for segment in segments(paragraph) if segment.kind == "hidden"]
    assert [segment.source_text for segment in hidden] == ["moved away"]


def test_non_text_content_under_a_deletion_keeps_its_kind() -> None:
    """A bookmark inside a deletion is still a marker: it must survive whatever
    happens to the deletion around it."""
    paragraph = paragraph_from(
        '<w:del w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">'
        '  <w:bookmarkStart w:id="9" w:name="Inside"/>'
        "  <w:r><w:delText>gone</w:delText><w:drawing/></w:r>"
        '  <w:bookmarkEnd w:id="9"/>'
        "</w:del>"
    )
    assert kinds(paragraph) == ["marker", "hidden", "opaque", "marker"]


# --------------------------------------------------------------------------------------
# Markers
# --------------------------------------------------------------------------------------


def test_bookmark_markers_sit_at_their_offsets() -> None:
    paragraph = body_paragraphs("bookmarks")[1]
    assert visible_text(paragraph) == "Before. Bookmarked span. After."
    assert marks(paragraph, "marker") == [("bookmarkStart", 8), ("bookmarkEnd", 23)]


def test_an_empty_bookmark_is_two_markers_at_one_offset() -> None:
    paragraph = body_paragraphs("bookmarks")[4]
    offset = len("An empty bookmark sits here: ")
    assert marks(paragraph, "marker") == [("bookmarkStart", offset), ("bookmarkEnd", offset)]


def test_comment_range_and_reference_are_zero_width() -> None:
    paragraph = body_paragraphs("comments")[1]
    end = len("This sentence carries a comment thread.")
    assert marks(paragraph, "marker") == [
        ("commentRangeStart", 0),
        ("commentRangeStart", 0),
        ("commentRangeEnd", end),
        ("commentRangeEnd", end),
    ]
    assert marks(paragraph, "opaque") == [("commentReference", end), ("commentReference", end)]


def test_permissions_and_proofing_markers_are_recognised() -> None:
    paragraph = paragraph_from(
        '<w:proofErr w:type="spellStart"/>'
        '<w:permStart w:id="1" w:edGrp="everyone"/>'
        "<w:r><w:t>text</w:t></w:r>"
        '<w:permEnd w:id="1"/>'
        '<w:moveFromRangeStart w:id="2" w:name="m"/>'
        '<w:moveFromRangeEnd w:id="2"/>'
        '<w:moveToRangeStart w:id="3" w:name="m"/>'
        '<w:moveToRangeEnd w:id="3"/>'
    )
    assert kinds(paragraph) == ["marker"] * 2 + ["text"] + ["marker"] * 5
    assert visible_text(paragraph) == "text"


# --------------------------------------------------------------------------------------
# Character-level policy
# --------------------------------------------------------------------------------------


def test_tabs_and_breaks_become_characters() -> None:
    paragraph = paragraph_from(
        "<w:r><w:t>a</w:t><w:tab/><w:t>b</w:t><w:br/><w:t>c</w:t><w:cr/><w:t>d</w:t></w:r>"
    )
    assert visible_text(paragraph) == "a\tb\nc\nd"
    assert kinds(paragraph) == ["text", "tab", "text", "break", "text", "break", "text"]


def test_hyphens_and_symbols_are_one_character_each() -> None:
    paragraph = paragraph_from(
        "<w:r>"
        "<w:t>a</w:t><w:noBreakHyphen/><w:t>b</w:t><w:softHyphen/>"
        '<w:sym w:font="Symbol" w:char="F0B7"/>'
        "</w:r>"
    )
    assert visible_text(paragraph) == "a" + chr(0x2011) + "b" + chr(0x00AD) + chr(0xF0B7)
    assert kinds(paragraph) == ["text"] * 5


@pytest.mark.parametrize("attributes", ['w:font="Symbol"', 'w:char="not-hex"', 'w:char="0007"'])
def test_an_unusable_symbol_still_takes_exactly_one_position(attributes: str) -> None:
    """Whatever the code says, a ``w:sym`` is one character wide -- otherwise a
    malformed attribute would shift every offset after it."""
    paragraph = paragraph_from(f"<w:r><w:t>a</w:t><w:sym {attributes}/><w:t>b</w:t></w:r>")
    assert visible_text(paragraph) == "a" + chr(0xFFFD) + "b"


def test_tab_stops_declared_in_ppr_are_not_content() -> None:
    """``w:pPr/w:tabs/w:tab`` declares a stop; only a ``w:tab`` in a run is a tab."""
    paragraph = paragraph_from(
        '<w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
        "<w:r><w:t>x</w:t></w:r>"
    )
    assert visible_text(paragraph) == "x"
    assert kinds(paragraph) == ["text"]


def test_property_revisions_are_not_content() -> None:
    """``w:pPrChange`` and ``w:rPrChange`` keep the previous properties, which
    may themselves contain a ``w:tabs``; none of it is text."""
    paragraph = paragraph_from(
        "<w:pPr>"
        '  <w:pPrChange w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">'
        '    <w:pPr><w:tabs><w:tab w:val="left" w:pos="360"/></w:tabs></w:pPr>'
        "  </w:pPrChange>"
        "</w:pPr>"
        "<w:r>"
        '  <w:rPr><w:rPrChange w:id="2" w:author="A" w:date="2024-01-01T00:00:00Z">'
        "    <w:rPr><w:i/></w:rPr></w:rPrChange></w:rPr>"
        "  <w:t>x</w:t>"
        "</w:r>"
    )
    assert kinds(paragraph) == ["text"]
    assert visible_text(paragraph) == "x"


@pytest.mark.parametrize(
    "container",
    [
        '<w:hyperlink r:id="rId9">{}</w:hyperlink>',
        '<w:ins w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">{}</w:ins>',
        '<w:moveTo w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">{}</w:moveTo>',
        "<w:sdt><w:sdtPr><w:id w:val=\"1\"/></w:sdtPr><w:sdtContent>{}</w:sdtContent></w:sdt>",
        '<w:smartTag w:element="place"><w:smartTagPr/>{}</w:smartTag>',
        '<w:customXml w:element="e"><w:customXmlPr/>{}</w:customXml>',
        '<w:dir w:val="rtl">{}</w:dir>',
        '<w:bdo w:val="ltr">{}</w:bdo>',
        '<w:fldSimple w:instr=" PAGE ">{}</w:fldSimple>',
    ],
)
def test_wrappers_are_traversed(container: str) -> None:
    paragraph = paragraph_from(container.format("<w:r><w:t>inside</w:t></w:r>"))
    assert visible_text(paragraph) == "inside"


def test_an_unknown_element_is_opaque_rather_than_transparent() -> None:
    """``w:ruby`` holds two texts and shows one; the engine reports neither and
    marks the construct opaque, so ``ranges`` refuses to cut through it."""
    paragraph = paragraph_from(
        "<w:r><w:t>a</w:t>"
        "<w:ruby><w:rt><w:r><w:t>gloss</w:t></w:r></w:rt>"
        "<w:rubyBase><w:r><w:t>base</w:t></w:r></w:rubyBase></w:ruby>"
        "<w:t>b</w:t></w:r>"
    )
    assert visible_text(paragraph) == "ab"
    assert kinds(paragraph) == ["text", "opaque", "text"]


@pytest.mark.parametrize(
    "tag",
    ["w:drawing", "w:pict", "w:object", "w:footnoteReference", "w:endnoteReference"],
)
def test_named_opaque_objects_take_no_width(tag: str) -> None:
    paragraph = paragraph_from(f"<w:r><w:t>a</w:t><{tag}/><w:t>b</w:t></w:r>")
    assert visible_text(paragraph) == "ab"
    assert kinds(paragraph) == ["text", "opaque", "text"]


# --------------------------------------------------------------------------------------
# Segment shape
# --------------------------------------------------------------------------------------


def test_containers_are_the_ancestor_chain_outermost_first() -> None:
    paragraph = paragraph_from(
        '<w:hyperlink r:id="rId9"><w:r><w:t>linked</w:t></w:r></w:hyperlink>'
    )
    (segment,) = segments(paragraph)
    hyperlink, run = segment.containers
    assert hyperlink.tag == qn("w:hyperlink")
    assert run.tag == qn("w:r")
    assert hyperlink.getparent() is paragraph
    assert run.getparent() is hyperlink
    assert segment.element.getparent() is run
    assert segment.run is run


def test_a_marker_has_no_run() -> None:
    paragraph = paragraph_from('<w:bookmarkStart w:id="1" w:name="B"/>')
    (segment,) = segments(paragraph)
    assert segment.containers == ()
    assert segment.run is None


def test_source_text_of_visible_and_opaque_segments() -> None:
    paragraph = paragraph_from("<w:r><w:t>a</w:t><w:tab/><w:drawing/></w:r>")
    assert [segment.source_text for segment in segments(paragraph)] == ["a", "\t", ""]


def test_hidden_tab_reports_the_character_it_stands_for() -> None:
    paragraph = paragraph_from(
        '<w:del w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:tab/></w:r></w:del>"
    )
    (segment,) = segments(paragraph)
    assert (segment.kind, segment.text, segment.source_text) == ("hidden", "", "\t")


# --------------------------------------------------------------------------------------
# position()
# --------------------------------------------------------------------------------------


def test_position_prefers_the_segment_opening_at_the_offset() -> None:
    paragraph = paragraph_from(
        "<w:r><w:t>ab</w:t></w:r>"
        '<w:bookmarkStart w:id="1" w:name="B"/>'
        "<w:r><w:t>cd</w:t></w:r>"
    )
    first, marker, second = segments(paragraph)
    assert position(paragraph, 0) == (first, 0)
    assert position(paragraph, 1) == (first, 1)
    # Offset 2 opens the second run; the marker sits there but is zero width.
    assert position(paragraph, 2) == (second, 0)
    assert marker.start == marker.end == 2
    # The end of the visible text has no segment opening at it, so it resolves to
    # the last one ending there.
    assert position(paragraph, 4) == (second, 2)


def test_position_on_an_empty_visible_text_returns_the_last_marker() -> None:
    paragraph = paragraph_from(
        '<w:bookmarkStart w:id="1" w:name="B"/><w:bookmarkEnd w:id="1"/>'
    )
    _, end_marker = segments(paragraph)
    assert position(paragraph, 0) == (end_marker, 0)


@pytest.mark.parametrize("offset", [-1, 3])
def test_position_rejects_offsets_outside_the_text(offset: int) -> None:
    paragraph = paragraph_from("<w:r><w:t>ab</w:t></w:r>")
    with pytest.raises(LocatorError) as excinfo:
        position(paragraph, offset)
    assert excinfo.value.code == "out-of-range"


def test_position_rejects_a_paragraph_with_no_content() -> None:
    paragraph = paragraph_from("<w:pPr/>")
    with pytest.raises(LocatorError) as excinfo:
        position(paragraph, 0)
    assert excinfo.value.code == "empty-paragraph"


# --------------------------------------------------------------------------------------
# Input handling
# --------------------------------------------------------------------------------------


def test_a_python_docx_paragraph_is_accepted(tmp_path: Path) -> None:
    path = tmp_path / "simple.docx"
    path.write_bytes(build("simple"))
    paragraph = Document(str(path)).paragraphs[0]
    assert visible_text(paragraph) == "Fixture: simple"


@pytest.mark.parametrize("bad", ["<w:r><w:t>x</w:t></w:r>", "<w:tbl/>"])
def test_anything_that_is_not_a_paragraph_is_refused(bad: str) -> None:
    element = paragraph_from(bad)[0]
    with pytest.raises(TypeError):
        segments(element)


def test_a_plain_object_is_refused() -> None:
    with pytest.raises(TypeError):
        visible_text("not a paragraph")


# --------------------------------------------------------------------------------------
# Properties, over every paragraph of every story of every fixture
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("name", sorted(ALL_FIXTURES))
def test_segments_tile_the_paragraph(name: str) -> None:
    for story, index, paragraph in all_paragraphs(name):
        where = f"{name}/{story}[{index}]"
        text = visible_text(paragraph)
        found = segments(paragraph)
        assert "".join(segment.text for segment in found) == text, where
        offset = 0
        for segment in found:
            assert segment.kind in SEGMENT_KINDS, where
            assert segment.start == offset, where
            assert segment.end - segment.start == len(segment.text), where
            if segment.kind in ("opaque", "marker", "hidden"):
                assert segment.text == "", where
            offset = segment.end
        assert offset == len(text), where
        assert found or text == "", where


@pytest.mark.parametrize("name", sorted(ALL_FIXTURES))
def test_containers_describe_the_real_ancestry(name: str) -> None:
    for story, index, paragraph in all_paragraphs(name):
        where = f"{name}/{story}[{index}]"
        for segment in segments(paragraph):
            chain = (*segment.containers, segment.element)
            assert chain[0].getparent() is paragraph, where
            for parent, child in itertools.pairwise(chain):
                assert child.getparent() is parent, where


@pytest.mark.parametrize("name", sorted(ALL_FIXTURES))
def test_every_offset_round_trips_through_position(name: str) -> None:
    for story, index, paragraph in all_paragraphs(name):
        where = f"{name}/{story}[{index}]"
        found = segments(paragraph)
        if not found:
            continue
        for offset in range(len(visible_text(paragraph)) + 1):
            segment, local = position(paragraph, offset)
            assert segment.start + local == offset, f"{where}@{offset}"
            assert 0 <= local <= len(segment.text), f"{where}@{offset}"
            assert segment in found, f"{where}@{offset}"


@pytest.mark.parametrize("name", sorted(ALL_FIXTURES))
def test_reading_never_mutates_the_tree(name: str) -> None:
    """D-004: ``textmodel`` is read-only, so the layers above can call it freely
    -- including in the middle of an edit -- without disturbing the document."""
    for story, index, paragraph in all_paragraphs(name):
        where = f"{name}/{story}[{index}]"
        before = etree.tostring(paragraph)
        text = visible_text(paragraph)
        segments(paragraph)
        fields(paragraph)
        for offset in range(len(text) + 1):
            try:
                position(paragraph, offset)
            except LocatorError:
                pass
        assert etree.tostring(paragraph) == before, where


def test_visible_text_agrees_with_the_segments_on_every_fixture() -> None:
    for name in sorted(ALL_FIXTURES):
        for story, index, paragraph in all_paragraphs(name):
            assert visible_text(paragraph) == "".join(
                segment.text for segment in segments(paragraph)
            ), f"{name}/{story}[{index}]"


def test_segment_kinds_constant_matches_the_type() -> None:
    assert set(get_args(SegmentKind)) == SEGMENT_KINDS
    assert Segment.__dataclass_fields__.keys() == {
        "kind",
        "element",
        "containers",
        "start",
        "end",
        "text",
    }
