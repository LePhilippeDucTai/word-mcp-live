"""Tests for :mod:`word_document_server.engine.numbering`.

What a list is made of cannot be read from the text of a document: the bullet
and the number are not in the paragraph, they are computed by Word from a
definition the paragraph points at.  So most cases here read the XML the module
writes -- the ids it allocates, where it puts the elements, what each level
says -- and the last sections check the two things a caller can observe: that a
package stays valid and that the older tools no longer point at a list they do
not own.

The last two sections leave the engine layer on purpose.  ``doc_apply_list``
(``tools/v2/numbering.py``) and the two ``utils/document_utils`` functions are
this module's only callers, and both are thin: what is worth pinning about them
is that they allocate through here rather than around it, which is a statement
about this module and is therefore made next to it.
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document as PDocument
from lxml import etree

from tests.fixtures.builders import build
from tests.support.package_check import validate_package
from tests.support.snapshot import assert_unchanged_except, snapshot
from word_document_server.engine.errors import PackageError
from word_document_server.engine.format import PPR_ORDER
from word_document_server.engine.numbering import (
    LIST_KINDS,
    MAX_LEVELS,
    NUMBERING_PART_NAME,
    apply_list,
    create_list_definition,
    list_definitions,
    numbering_root,
    paragraph_list_info,
    restart_list,
)
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.xmlns import NAMESPACES, qn
from word_document_server.tools.v2.registry import v2_tools
from word_document_server.utils.document_utils import insert_numbered_list_near_text

W_ABSTRACT_NUM = qn("w:abstractNum")
W_ABSTRACT_NUM_ID = qn("w:abstractNumId")
W_ILVL = qn("w:ilvl")
W_NUM = qn("w:num")
W_NUM_ID = qn("w:numId")
W_NSID = qn("w:nsid")
W_P = qn("w:p")
W_PPR = qn("w:pPr")
W_VAL = qn("w:val")


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------


def _pkg(name: str) -> DocxPackage:
    return DocxPackage.open(build(name))


def paragraph_from(children: str = "") -> etree._Element:
    """Return a standalone ``w:p`` whose children are the given XML fragment."""
    declarations = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in NAMESPACES.items())
    return etree.fromstring(f"<w:p {declarations}>{children}</w:p>")


def names(element: etree._Element | None) -> list[str]:
    """Local names of the element's children, in order."""
    return [] if element is None else [etree.QName(child).localname for child in element]


def _without_numbering(blob: bytes) -> bytes:
    """The same package with no numbering part, nor content type, nor relationship.

    python-docx builds every fixture from a template that already ships a
    ``word/numbering.xml``, so the only way to exercise a document that has none
    -- which is most documents an agent is handed -- is to take one away.
    """
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in items:
            if info.filename == "word/numbering.xml":
                continue
            if info.filename in ("[Content_Types].xml", "word/_rels/document.xml.rels"):
                root = etree.fromstring(data)
                for child in list(root):
                    target = child.get("PartName") or child.get("Target") or ""
                    if target.lstrip("/").endswith("numbering.xml"):
                        root.remove(child)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            archive.writestr(info, data)
    return out.getvalue()


def _num_ids(root: etree._Element) -> list[int]:
    return [int(el.get(W_NUM_ID)) for el in root.iter(W_NUM)]


def _abstract_ids(root: etree._Element) -> list[int]:
    return [int(el.get(W_ABSTRACT_NUM_ID)) for el in root.iter(W_ABSTRACT_NUM)]


def _list_paragraph(pkg: DocxPackage) -> etree._Element:
    """The first paragraph of the package that carries a ``w:numPr``."""
    for paragraph in pkg.document.iter(W_P):
        if paragraph_list_info(paragraph) is not None:
            return paragraph
    raise AssertionError("the fixture holds no numbered paragraph")


def call(name: str, *args: Any, **kwargs: Any) -> dict[str, Any]:
    """Call a V2 tool the way an MCP client reaches it, and return its report."""
    return asyncio.run(v2_tools()[name](*args, **kwargs))


# --------------------------------------------------------------------------------------
# The part
# --------------------------------------------------------------------------------------


def test_a_package_without_a_numbering_part_gets_one(tmp_path: Path) -> None:
    pkg = DocxPackage.open(_without_numbering(build("simple")))
    assert numbering_root(pkg) is None
    assert list_definitions(pkg) == []

    num_id = create_list_definition(pkg, "bullet")
    apply_list(pkg.document.iter(W_P).__next__(), num_id)

    path = pkg.save(tmp_path / "created.docx")
    with zipfile.ZipFile(path) as archive:
        content_types = archive.read("[Content_Types].xml").decode()
        assert "word/numbering.xml" in archive.namelist()
    assert NUMBERING_PART_NAME in content_types
    assert validate_package(path) == []


def test_an_existing_numbering_part_is_used_rather_than_a_second_one(tmp_path: Path) -> None:
    pkg = _pkg("complex_numbering")
    before = list(pkg.package.iter_parts())

    create_list_definition(pkg, "decimal")
    create_list_definition(pkg, "bullet")

    after = [str(part.partname) for part in pkg.package.iter_parts()]
    assert len(after) == len(before)
    assert after.count(NUMBERING_PART_NAME) == 1
    assert validate_package(pkg.save(tmp_path / "reused.docx")) == []


def test_reading_a_package_without_a_numbering_part_never_creates_one() -> None:
    pkg = DocxPackage.open(_without_numbering(build("simple")))

    assert numbering_root(pkg) is None
    assert paragraph_list_info(pkg.document.iter(W_P).__next__()) is None

    assert numbering_root(pkg) is None, "a read must not have created the part"
    with pytest.raises(PackageError):
        restart_list(pkg, 1)


# --------------------------------------------------------------------------------------
# Allocation
# --------------------------------------------------------------------------------------


def test_a_new_definition_reuses_no_id_the_document_already_uses() -> None:
    pkg = _pkg("complex_numbering")
    root = numbering_root(pkg)
    taken_nums, taken_abstracts = set(_num_ids(root)), set(_abstract_ids(root))
    assert {900, 901, 902} <= taken_nums  # sanity: the fixture's own lists

    first = create_list_definition(pkg, "decimal")
    second = create_list_definition(pkg, "bullet")

    assert first not in taken_nums
    assert second not in taken_nums | {first}
    assert not taken_abstracts & set(_abstract_ids(root)) - taken_abstracts
    assert len(set(_abstract_ids(root))) == len(_abstract_ids(root)), "abstract ids collide"
    assert len(set(_num_ids(root))) == len(_num_ids(root)), "numIds collide"


def test_each_definition_carries_an_nsid_of_its_own() -> None:
    pkg = _pkg("complex_numbering")
    root = numbering_root(pkg)
    before = [el.get(W_VAL) for el in root.iter(W_NSID)]

    create_list_definition(pkg, "decimal")
    create_list_definition(pkg, "bullet")

    after = [el.get(W_VAL) for el in root.iter(W_NSID)]
    assert len(after) == len(before) + 2
    assert len(set(after)) == len(after), "two lists share an nsid"
    assert all(len(value) == 8 and int(value, 16) >= 0 for value in after)


def test_the_same_document_allocates_the_same_ids_twice() -> None:
    """Determinism: nothing here is random, so a fixture stays reproducible."""
    runs = []
    for _ in range(2):
        pkg = _pkg("complex_numbering")
        num_id = create_list_definition(pkg, "multilevel")
        root = numbering_root(pkg)
        runs.append((num_id, [el.get(W_VAL) for el in root.iter(W_NSID)]))
    assert runs[0] == runs[1]


# --------------------------------------------------------------------------------------
# The shape of a definition
# --------------------------------------------------------------------------------------


def test_every_abstract_definition_precedes_every_num() -> None:
    """``w:numbering`` is a sequence: ``abstractNum*`` then ``num*``."""
    pkg = _pkg("complex_numbering")
    create_list_definition(pkg, "bullet")
    create_list_definition(pkg, "decimal")

    order = names(numbering_root(pkg))
    assert order.count("abstractNum") and order.count("num")
    assert order.index("num") > max(
        index for index, tag in enumerate(order) if tag == "abstractNum"
    )


@pytest.mark.parametrize(
    ("kind", "formats", "labels"),
    [
        ("bullet", ["bullet", "bullet", "bullet"], ["", "o", ""]),
        ("decimal", ["decimal", "lowerLetter", "lowerRoman"], ["%1.", "%2.", "%3."]),
        ("multilevel", ["decimal"] * 3, ["%1.", "%1.%2.", "%1.%2.%3."]),
    ],
)
def test_each_kind_labels_its_levels(kind: str, formats: list[str], labels: list[str]) -> None:
    pkg = _pkg("simple")
    num_id = create_list_definition(pkg, kind)

    levels = next(d for d in list_definitions(pkg) if d["num_id"] == num_id)["levels"]
    assert len(levels) == MAX_LEVELS
    assert [level["level"] for level in levels] == list(range(MAX_LEVELS))
    assert [level["format"] for level in levels[:3]] == formats
    assert [level["text"] for level in levels[:3]] == labels
    assert all(level["start"] == 1 for level in levels)


def test_a_one_level_list_is_declared_as_one() -> None:
    pkg = _pkg("simple")
    single = create_list_definition(pkg, "decimal", levels=1)
    ladder = create_list_definition(pkg, "decimal", levels=3)

    by_id = {d["num_id"]: d for d in list_definitions(pkg)}
    assert by_id[single]["multi_level_type"] == "singleLevel"
    assert len(by_id[single]["levels"]) == 1
    assert by_id[ladder]["multi_level_type"] == "hybridMultilevel"
    assert len(by_id[ladder]["levels"]) == 3


@pytest.mark.parametrize("levels", [0, -1, MAX_LEVELS + 1])
def test_a_definition_of_no_level_or_too_many_is_refused(levels: int) -> None:
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="levels"):
        create_list_definition(pkg, "bullet", levels=levels)


def test_an_unknown_kind_is_refused_by_name() -> None:
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="unknown list kind"):
        create_list_definition(pkg, "roman-numerals")
    assert all(kind in LIST_KINDS for kind in ("bullet", "decimal", "multilevel"))


# --------------------------------------------------------------------------------------
# Applying a list to a paragraph
# --------------------------------------------------------------------------------------


def test_the_numbering_properties_land_where_the_schema_wants_them() -> None:
    """``w:numPr`` after ``w:pStyle`` and before ``w:jc``, ``w:ilvl`` before ``w:numId``.

    The tool this replaces appended ``w:numPr`` after every other property, which
    Word reads as a document to repair.
    """
    paragraph = paragraph_from(
        '<w:pPr><w:pStyle w:val="ListParagraph"/><w:jc w:val="both"/></w:pPr>'
        "<w:r><w:t>item</w:t></w:r>"
    )

    apply_list(paragraph, 7, 2)

    properties = paragraph.find(W_PPR)
    assert names(properties) == ["pStyle", "numPr", "jc"]
    ranks = [PPR_ORDER.index(f"w:{name}") for name in names(properties)]
    assert ranks == sorted(ranks)
    assert names(properties.find(qn("w:numPr"))) == ["ilvl", "numId"]
    assert paragraph_list_info(paragraph) == {"num_id": 7, "level": 2}


def test_a_paragraph_that_was_already_in_a_list_moves_to_the_new_one() -> None:
    paragraph = paragraph_from()
    apply_list(paragraph, 3, 1)
    apply_list(paragraph, 8, 0)

    assert len(paragraph.findall(f"{W_PPR}/{qn('w:numPr')}")) == 1
    assert paragraph_list_info(paragraph) == {"num_id": 8, "level": 0}


def test_a_python_docx_paragraph_is_accepted_like_an_element() -> None:
    document = PDocument()
    paragraph = document.add_paragraph("item")

    apply_list(paragraph, 4, 1)

    assert paragraph_list_info(paragraph) == {"num_id": 4, "level": 1}
    assert paragraph_list_info(paragraph._p) == {"num_id": 4, "level": 1}


@pytest.mark.parametrize(
    ("num_id", "level"), [(0, 0), (-1, 0), (1, -1), (1, MAX_LEVELS)]
)
def test_an_impossible_id_or_level_leaves_the_paragraph_alone(num_id: int, level: int) -> None:
    paragraph = paragraph_from("<w:r><w:t>item</w:t></w:r>")

    with pytest.raises(ValueError):
        apply_list(paragraph, num_id, level)

    assert paragraph.find(W_PPR) is None
    assert paragraph_list_info(paragraph) is None


def test_something_that_is_not_a_paragraph_is_refused() -> None:
    with pytest.raises(TypeError):
        apply_list(etree.fromstring(f'<w:r xmlns:w="{NAMESPACES["w"]}"/>'), 1)
    with pytest.raises(TypeError):
        paragraph_list_info("a paragraph")


# --------------------------------------------------------------------------------------
# Restarting
# --------------------------------------------------------------------------------------


def test_a_restarted_list_keeps_the_labels_and_starts_again() -> None:
    pkg = _pkg("complex_numbering")
    num_id = create_list_definition(pkg, "decimal", levels=3)

    restarted = restart_list(pkg, num_id)

    by_id = {d["num_id"]: d for d in list_definitions(pkg)}
    assert restarted != num_id
    assert by_id[restarted]["abstract_num_id"] == by_id[num_id]["abstract_num_id"]
    assert by_id[restarted]["levels"] == by_id[num_id]["levels"]
    assert by_id[restarted]["overrides"] == [
        {"level": 0, "start": 1},
        {"level": 1, "start": 1},
        {"level": 2, "start": 1},
    ]
    assert by_id[num_id]["overrides"] == []


def test_a_restart_overrides_the_start_value_the_level_declares() -> None:
    """The fixture's abstract 900 starts its third level at 3, not at 1."""
    pkg = _pkg("complex_numbering")

    restarted = restart_list(pkg, 900)

    by_id = {d["num_id"]: d for d in list_definitions(pkg)}
    assert by_id[restarted]["overrides"] == [
        {"level": 0, "start": 1},
        {"level": 1, "start": 1},
        {"level": 2, "start": 3},
    ]


def test_restarting_a_list_the_document_does_not_define_is_refused() -> None:
    pkg = _pkg("complex_numbering")
    with pytest.raises(PackageError, match="no list numId 4242"):
        restart_list(pkg, 4242)


# --------------------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------------------


def test_list_definitions_reports_the_lists_the_fixture_defines() -> None:
    by_id = {d["num_id"]: d for d in list_definitions(_pkg("complex_numbering"))}

    assert [d["num_id"] for d in list_definitions(_pkg("complex_numbering"))] == sorted(by_id)
    assert by_id[900]["abstract_num_id"] == 900
    assert by_id[900]["multi_level_type"] == "multilevel"
    assert [level["format"] for level in by_id[900]["levels"]] == [
        "decimal",
        "lowerLetter",
        "lowerRoman",
    ]
    # 901 is the same abstract definition, restarted at 5.
    assert by_id[901]["abstract_num_id"] == 900
    assert by_id[901]["overrides"] == [{"level": 0, "start": 5}]
    assert by_id[902]["levels"][0]["format"] == "bullet"


def test_paragraph_list_info_reads_direct_numbering_only() -> None:
    pkg = _pkg("complex_numbering")

    numbered = _list_paragraph(pkg)
    assert paragraph_list_info(numbered)["num_id"] in {d["num_id"] for d in list_definitions(pkg)}

    plain = paragraph_from("<w:r><w:t>plain</w:t></w:r>")
    assert paragraph_list_info(plain) is None
    styled = paragraph_from('<w:pPr><w:pStyle w:val="ListParagraph"/></w:pPr>')
    assert paragraph_list_info(styled) is None, "a style is not direct numbering"


def test_a_numbering_reference_without_a_level_reads_as_level_zero() -> None:
    paragraph = paragraph_from('<w:pPr><w:numPr><w:numId w:val="5"/></w:numPr></w:pPr>')
    assert paragraph_list_info(paragraph) == {"num_id": 5, "level": 0}

    empty = paragraph_from("<w:pPr><w:numPr/></w:pPr>")
    assert empty is not None
    assert paragraph_list_info(empty) == {"num_id": None, "level": 0}


# --------------------------------------------------------------------------------------
# At package scale
# --------------------------------------------------------------------------------------


def test_numbering_a_paragraph_touches_it_and_the_numbering_part_only(tmp_path: Path) -> None:
    path = tmp_path / "simple.docx"
    path.write_bytes(build("simple"))
    before = snapshot(path.read_bytes())

    pkg = DocxPackage.open(str(path))
    paragraphs = list(pkg.document.iter(W_P))
    num_id = create_list_definition(pkg, "decimal")
    apply_list(paragraphs[1], num_id)
    pkg.save(path)

    # No table and no content control in this fixture: the snapshot index space
    # and the V2 one are the same one here.
    assert_unchanged_except(
        before, snapshot(path.read_bytes()), paragraphs=[1], parts=["word/numbering.xml"]
    )
    assert validate_package(path) == []


# --------------------------------------------------------------------------------------
# doc_apply_list, the V2 tool over this module
# --------------------------------------------------------------------------------------


def _three_paragraphs(path: Path) -> Path:
    document = PDocument()
    for text in ("Alpha", "Bravo", "Charlie"):
        document.add_paragraph(text)
    document.save(str(path))
    return path


def test_doc_apply_list_numbers_every_paragraph_it_is_given(tmp_path: Path) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")

    report = call(
        "doc_apply_list", str(path), [{"paragraph": 0}, {"find": "Charlie"}], kind="decimal"
    )

    assert report["status"] == "ok"
    assert report["saved"] is True
    assert [change["paragraph"] for change in report["changes"]] == [0, 2]
    assert all(change["before"] == change["after"] for change in report["changes"])

    pkg = DocxPackage.open(str(path))
    info = [paragraph_list_info(p) for p in pkg.document.iter(W_P)]
    assert info == [
        {"num_id": report["num_id"], "level": 0},
        None,
        {"num_id": report["num_id"], "level": 0},
    ]
    assert validate_package(path) == []


def test_two_calls_make_two_lists_that_count_apart(tmp_path: Path) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")

    first = call("doc_apply_list", str(path), [{"paragraph": 0}])
    second = call("doc_apply_list", str(path), [{"paragraph": 2}])

    assert first["num_id"] != second["num_id"]
    definitions = {d["num_id"]: d for d in list_definitions(DocxPackage.open(str(path)))}
    assert definitions[first["num_id"]]["abstract_num_id"] != (
        definitions[second["num_id"]]["abstract_num_id"]
    )
    assert validate_package(path) == []


def test_doc_apply_list_says_when_a_paragraph_leaves_a_list(tmp_path: Path) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")
    first = call("doc_apply_list", str(path), [{"paragraph": 1}])

    report = call("doc_apply_list", str(path), [{"paragraph": 1}], kind="bullet")

    assert report["status"] == "ok"
    assert any(str(first["num_id"]) in warning for warning in report["warnings"])
    assert validate_package(path) == []


def test_two_locators_on_one_paragraph_number_it_once(tmp_path: Path) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")

    report = call("doc_apply_list", str(path), [{"paragraph": 1}, {"find": "Bravo"}])

    assert [change["paragraph"] for change in report["changes"]] == [1]
    assert report["warnings"], "the caller is told the second locator was a repeat"
    assert validate_package(path) == []


@pytest.mark.parametrize(
    ("locators", "kind", "code"),
    [
        ([], "bullet", "invalid_argument"),
        ({"paragraph": 0}, "bullet", "invalid_argument"),
        ([{"paragraph": 0}], "roman", "invalid_argument"),
        ([{"find": "nowhere"}], "bullet", "not_found"),
        ([{"paragraph": 0}, {"paragraph": 99}], "bullet", "not_found"),
    ],
)
def test_a_refused_call_writes_nothing(
    tmp_path: Path, locators: Any, kind: str, code: str
) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")
    before = path.read_bytes()

    report = call("doc_apply_list", str(path), locators, kind=kind)

    assert report == {"status": "error", "code": code, "message": report["message"]}
    assert path.read_bytes() == before


def test_a_dry_run_reports_the_change_and_writes_nothing(tmp_path: Path) -> None:
    path = _three_paragraphs(tmp_path / "list.docx")
    before = path.read_bytes()

    report = call("doc_apply_list", str(path), [{"paragraph": 0}], dry_run=True)

    assert report["dry_run"] is True
    assert report["saved"] is False
    assert len(report["changes"]) == 1
    assert path.read_bytes() == before


# --------------------------------------------------------------------------------------
# The older tools allocate through this module
# --------------------------------------------------------------------------------------


def test_insert_numbered_list_near_text_no_longer_points_at_numid_1_or_2(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy.docx"
    document = PDocument()
    document.add_paragraph("Anchor")
    document.save(str(path))

    assert "inserted after" in insert_numbered_list_near_text(
        str(path), target_text="Anchor", list_items=["one", "two"], bullet_type="number"
    )
    assert "inserted after" in insert_numbered_list_near_text(
        str(path), target_text="Anchor", list_items=["dot"], bullet_type="bullet"
    )

    pkg = DocxPackage.open(str(path))
    used = [
        info["num_id"]
        for info in (paragraph_list_info(p) for p in pkg.document.iter(W_P))
        if info is not None
    ]
    assert len(used) == 3
    assert set(used) & {1, 2} == set(), "the hard-coded ids are someone else's lists"
    assert len(set(used)) == 2, "one list per call, shared by the items of that call"
    defined = {d["num_id"] for d in list_definitions(pkg)}
    assert set(used) <= defined
    assert validate_package(path) == []


def test_a_list_of_no_item_leaves_no_definition_behind(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    document = PDocument()
    document.add_paragraph("Anchor")
    document.save(str(path))
    before = len(list_definitions(DocxPackage.open(str(path))))

    result = insert_numbered_list_near_text(str(path), target_text="Anchor", list_items=[])

    assert "with 0 items inserted" in result
    assert len(list_definitions(DocxPackage.open(str(path)))) == before
    assert validate_package(path) == []


def test_the_list_style_is_applied_only_when_the_document_defines_it(tmp_path: Path) -> None:
    path = tmp_path / "nostyle.docx"
    document = PDocument()
    document.add_paragraph("Anchor")
    document.styles["List Paragraph"].delete()
    document.save(str(path))

    insert_numbered_list_near_text(str(path), target_text="Anchor", list_items=["one"])

    pkg = DocxPackage.open(str(path))
    item = _list_paragraph(pkg)
    assert item.find(f"{W_PPR}/{qn('w:pStyle')}") is None
    assert validate_package(path) == []
