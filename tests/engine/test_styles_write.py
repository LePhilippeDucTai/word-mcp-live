"""Tests for the write half of :mod:`word_document_server.engine.styles`.

Writing a style is not writing formatting: a style is a definition every
paragraph that names it obeys, so a defect here is a defect repeated everywhere
the style is used, and it does not show up in the paragraph that revealed it.
What is pinned below is therefore mostly about the *shape* of what is written
rather than about one visible outcome.

*The schema order is not negotiable.*
    ``w:style``, ``w:rPr`` and ``w:pPr`` are ``xsd:sequence``: a child in the
    wrong place makes Word repair the document -- or refuse it -- whatever the
    properties say.  Every writer here is checked against the ranks of
    :data:`~...styles.STYLE_ORDER`,
    :data:`~...format.RPR_ORDER` and :data:`~...format.PPR_ORDER`, whatever
    order the caller listed the properties in.

*What a reader reports, a writer takes.*
    The point of one vocabulary is that ``get_style`` -> ``create_style`` is a
    round trip and not a translation.  So the round trip is asserted on a style
    read from a fixture, theme references included -- a theme reference is
    written back as a reference and never as the colour it currently resolves to
    (D-028).

*A refusal leaves nothing behind.*
    A spec with one bad value must leave the style sheet exactly as it was: a
    half-written style is one Word would offer to a human.

*Deleting is two operations, and the second one is invisible.*
    The places that use the style, and the other styles that inherit from it.
    A caller can check the first for itself; nothing outside this module can
    check the second, so both are pinned.

The last section leaves the engine layer on purpose: ``core/styles.py`` and the
three V2 tools are this module's callers, and what is worth pinning about them
is that they go through it rather than around it.
"""

from __future__ import annotations

import asyncio
import zipfile

import pytest
from lxml import etree

from tests.fixtures.builders import build
from tests.support.package_check import validate_package
from tests.support.snapshot import assert_unchanged_except, snapshot
from word_document_server.engine.errors import LocatorError, PackageError
from word_document_server.engine.format import PPR_ORDER, RPR_ORDER
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.styles import (
    STYLE_ORDER,
    STYLES_PARTNAME,
    clone_style,
    create_style,
    delete_style,
    derive_style_id,
    find_style_usage,
    get_style,
    list_styles,
    styles_root,
    update_style,
)
from word_document_server.engine.xmlns import qn
from word_document_server.tools.v2.registry import v2_tools

W_STYLE = qn("w:style")
W_STYLE_ID = qn("w:styleId")
W_VAL = qn("w:val")

#: Part of the package the tests below are allowed to change.
STYLES_PART = "word/styles.xml"


def _pkg(name: str) -> DocxPackage:
    return DocxPackage.open(build(name))


def _element(pkg: DocxPackage, style_id: str) -> etree._Element:
    """The live ``w:style`` element of `style_id`."""
    root = styles_root(pkg)
    assert root is not None
    found = [
        style for style in root.findall(W_STYLE) if style.get(W_STYLE_ID) == style_id
    ]
    assert len(found) == 1, f"expected exactly one w:style {style_id!r}, found {len(found)}"
    return found[0]


def _local_names(element: etree._Element) -> list[str]:
    """``w:``-prefixed names of the children of `element`, in document order."""
    return [f"w:{etree.QName(child).localname}" for child in element]


def _assert_in_schema_order(element: etree._Element, order: tuple[str, ...]) -> None:
    """Fail unless the children of `element` follow `order`."""
    names = _local_names(element)
    unknown = [name for name in names if name not in order]
    assert not unknown, f"children outside the transcribed sequence: {unknown}"
    ranks = [order.index(name) for name in names]
    assert ranks == sorted(ranks), (
        f"children of <{etree.QName(element).localname}> are out of schema order: {names}"
    )


def _run(coro):
    """Drive a V2 tool call to completion."""
    return asyncio.run(coro)


def _tool(name: str):
    tools = v2_tools()
    assert name in tools, f"{name} is not registered; registered: {sorted(tools)}"
    return tools[name]


# --------------------------------------------------------------------------
# Identity: the id, and what makes one taken
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("Fixture Body", "FixtureBody"),
        ("heading 1", "heading1"),
        ("Note (long)", "Notelong"),
        ("Résumé détaillé", "Résumédétaillé"),
        ("Already", "Already"),
    ],
)
def test_derive_style_id_keeps_letters_and_digits_only(name, expected):
    assert derive_style_id(name) == expected


def test_derive_style_id_refuses_a_name_it_cannot_derive_one_from():
    with pytest.raises(ValueError, match="nothing"):
        derive_style_id("--- ---")


def test_create_style_derives_the_id_from_the_name():
    pkg = _pkg("simple")
    info = create_style(pkg, {"name": "Probe Body"})

    assert info.style_id == "ProbeBody"
    assert info.name == "Probe Body"
    assert info.family == "paragraph"
    # A style this server writes is a custom style, and says so: without
    # w:customStyle Word believes the document is claiming one of its own.
    assert info.builtin is False


def test_create_style_writes_the_id_it_is_given():
    """Word's own styles are the reason `style_id` exists: `heading 1` -> `Heading1`."""
    pkg = _pkg("simple")
    info = create_style(
        pkg, {"name": "heading 42", "style_id": "Heading42", "builtin": True}
    )

    assert info.style_id == "Heading42"
    assert info.builtin is True
    assert _element(pkg, "Heading42").get(qn("w:customStyle")) is None


def test_create_style_refuses_an_id_the_document_already_defines():
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        create_style(pkg, {"name": "FixtureRoot"})
    assert caught.value.code == "already_exists"


def test_create_style_refuses_a_name_that_differs_only_in_case():
    """Word's style pane treats two names that differ in case as one name."""
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        create_style(pkg, {"name": "fixture ROOT"})
    assert caught.value.code == "already_exists"


def test_a_refused_spec_leaves_the_style_sheet_as_it_was():
    pkg = _pkg("simple")
    before = [info.style_id for info in list_styles(pkg)]

    with pytest.raises(ValueError, match="size_pt"):
        create_style(pkg, {"name": "Probe Body", "run_props": {"size_pt": 11.3}})

    assert [info.style_id for info in list_styles(pkg)] == before


def test_a_refused_reference_leaves_the_style_sheet_as_it_was():
    pkg = _pkg("simple")
    before = [info.style_id for info in list_styles(pkg)]

    with pytest.raises(LocatorError) as caught:
        create_style(pkg, {"name": "Probe Body", "based_on": "NoSuchStyle"})

    assert caught.value.code == "not_found"
    assert [info.style_id for info in list_styles(pkg)] == before


# --------------------------------------------------------------------------
# Schema order
# --------------------------------------------------------------------------


def test_style_children_are_written_in_schema_order():
    """Whatever order the spec listed them in."""
    pkg = _pkg("style_inheritance")
    create_style(
        pkg,
        {
            "paragraph_props": {"alignment": "center"},
            "run_props": {"bold": True},
            "q_format": True,
            "ui_priority": 12,
            "link": "FixtureBranchChar",
            "next": "Normal",
            "based_on": "FixtureRoot",
            "name": "Probe Order",
        },
    )

    style = _element(pkg, "ProbeOrder")
    _assert_in_schema_order(style, STYLE_ORDER)
    assert _local_names(style) == [
        "w:name",
        "w:basedOn",
        "w:next",
        "w:link",
        "w:uiPriority",
        "w:qFormat",
        "w:pPr",
        "w:rPr",
    ]


def test_run_and_paragraph_properties_are_written_in_schema_order():
    pkg = _pkg("simple")
    create_style(
        pkg,
        {
            "name": "Probe Props",
            # Deliberately the reverse of the schema order in both blocks.
            "run_props": {
                "underline": "single",
                "size_pt": 11.0,
                "color": "C00000",
                "bold": True,
                "font": "Consolas",
            },
            "paragraph_props": {
                "outline_level": 3,
                "alignment": "both",
                "indent": {"left": 720},
                "spacing": {"after": 120},
                "numbering": {"num_id": 4, "level": 0},
                "keep_next": True,
            },
        },
    )

    style = _element(pkg, "ProbeProps")
    _assert_in_schema_order(style.find(qn("w:rPr")), RPR_ORDER)
    _assert_in_schema_order(style.find(qn("w:pPr")), PPR_ORDER)


def test_a_new_style_goes_after_the_last_one():
    """``CT_Styles`` is docDefaults, latentStyles, then the styles."""
    pkg = _pkg("simple")
    create_style(pkg, {"name": "Probe Last"})

    root = styles_root(pkg)
    names = _local_names(root)
    assert names[-1] == "w:style"
    assert root.findall(W_STYLE)[-1].get(W_STYLE_ID) == "ProbeLast"


def test_the_written_package_is_valid(tmp_path):
    pkg = _pkg("style_inheritance")
    create_style(
        pkg,
        {
            "name": "Probe Valid",
            "based_on": "FixtureLeaf",
            "run_props": {"bold": True, "color": "112233", "size_pt": 13.5},
            "paragraph_props": {"alignment": "center", "spacing": {"before": 60}},
        },
    )
    target = pkg.save(tmp_path / "written.docx")

    assert validate_package(target) == []


def test_creating_a_style_touches_nothing_but_the_style_sheet(tmp_path):
    source = build("style_inheritance")
    pkg = DocxPackage.open(source)
    create_style(pkg, {"name": "Probe Isolated", "run_props": {"italic": True}})
    target = pkg.save(tmp_path / "written.docx")

    assert_unchanged_except(
        snapshot(source), snapshot(target.read_bytes()), parts={STYLES_PART}
    )


# --------------------------------------------------------------------------
# One vocabulary, both ways
# --------------------------------------------------------------------------


def test_what_get_style_reports_can_be_written_back():
    """The point of sharing a vocabulary: no translation table in between."""
    pkg = _pkg("style_inheritance")
    source = get_style(pkg, "FixtureLeaf")

    create_style(
        pkg,
        {
            "name": "Probe Copy",
            "run_props": source.resolved["run"],
            "paragraph_props": source.resolved["paragraph"],
        },
    )

    written = get_style(pkg, "ProbeCopy")
    assert written.run_props == source.resolved["run"]
    assert written.paragraph_props == source.resolved["paragraph"]


def test_a_theme_reference_is_written_as_a_reference(tmp_path):
    """D-028: the resolved colour is a rendering hint, never what goes in."""
    pkg = _pkg("simple")
    create_style(
        pkg,
        {
            "name": "Probe Theme",
            "run_props": {
                "font": {"value": "Aptos", "theme": "minorHAnsi"},
                "color": {"value": "4472C4", "theme": "accent1", "tint": "99"},
            },
        },
    )

    fonts = _element(pkg, "ProbeTheme").find(qn("w:rPr")).find(qn("w:rFonts"))
    assert fonts.get(qn("w:asciiTheme")) == "minorHAnsi"
    assert fonts.get(qn("w:hAnsiTheme")) == "minorHAnsi"
    assert fonts.get(qn("w:ascii")) == "Aptos"

    color = _element(pkg, "ProbeTheme").find(qn("w:rPr")).find(qn("w:color"))
    assert color.get(W_VAL) == "4472C4"
    assert color.get(qn("w:themeColor")) == "accent1"
    assert color.get(qn("w:themeTint")) == "99"

    assert get_style(pkg, "ProbeTheme").run_props["color"] == {
        "value": "4472C4",
        "theme": "accent1",
        "tint": "99",
        "shade": None,
    }


def test_a_theme_colour_with_no_literal_still_writes_the_required_val():
    """``w:val`` is required on ``w:color``; Word caches ``auto`` beside a theme."""
    pkg = _pkg("simple")
    create_style(
        pkg, {"name": "Probe Themed", "run_props": {"color": {"theme": "accent2"}}}
    )

    color = _element(pkg, "ProbeThemed").find(qn("w:rPr")).find(qn("w:color"))
    assert color.get(W_VAL) == "auto"
    assert color.get(qn("w:themeColor")) == "accent2"


def test_the_east_asian_font_is_a_property_of_its_own():
    pkg = _pkg("simple")
    create_style(
        pkg,
        {
            "name": "Probe Scripts",
            "run_props": {"font": "Consolas", "font_east_asia": "MS Mincho"},
        },
    )

    fonts = _element(pkg, "ProbeScripts").find(qn("w:rPr")).find(qn("w:rFonts"))
    assert fonts.get(qn("w:ascii")) == "Consolas"
    assert fonts.get(qn("w:eastAsia")) == "MS Mincho"


def test_an_unknown_property_is_refused_rather_than_ignored():
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="unknown run properties: bolt"):
        create_style(pkg, {"name": "Probe Typo", "run_props": {"bolt": True}})
    with pytest.raises(ValueError, match="unknown paragraph properties: align"):
        create_style(pkg, {"name": "Probe Typo", "paragraph_props": {"align": "left"}})


def test_an_unknown_spec_key_is_refused_rather_than_ignored():
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="unknown style spec keys: bold"):
        create_style(pkg, {"name": "Probe Typo", "bold": True})


def test_char_style_needs_the_package_to_be_checked_against():
    """A python-docx Document cannot answer "is this a character style?"."""
    from docx import Document as PDocument

    with pytest.raises(ValueError, match="DocxPackage"):
        create_style(
            PDocument(),
            {"name": "Probe Char Style", "run_props": {"char_style": "Strong"}},
        )


def test_a_character_style_has_no_paragraph_properties():
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="only a paragraph style"):
        create_style(
            pkg,
            {
                "name": "Probe Char",
                "family": "character",
                "paragraph_props": {"alignment": "left"},
            },
        )


def test_a_style_is_created_as_a_paragraph_or_a_character_style():
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="paragraph, character"):
        create_style(pkg, {"name": "Probe Table", "family": "table"})


def test_false_on_a_toggle_writes_the_explicit_off():
    """It is how a style cancels what the style it is based on turned on."""
    pkg = _pkg("style_inheritance")
    create_style(
        pkg,
        {
            "name": "Probe Off",
            "based_on": "FixtureBranch",
            "run_props": {"bold": False},
            "paragraph_props": {"keep_next": False},
        },
    )

    style = _element(pkg, "ProbeOff")
    assert style.find(qn("w:rPr")).find(qn("w:b")).get(W_VAL) == "0"
    assert style.find(qn("w:pPr")).find(qn("w:keepNext")).get(W_VAL) == "0"
    assert get_style(pkg, "ProbeOff").resolved["run"]["bold"] is False


# --------------------------------------------------------------------------
# References
# --------------------------------------------------------------------------


def test_a_reference_may_be_given_by_name_and_is_stored_as_the_id():
    pkg = _pkg("style_inheritance")
    info = create_style(pkg, {"name": "Probe Named", "based_on": "Fixture Branch"})

    assert info.based_on == "FixtureBranch"
    assert _element(pkg, "ProbeNamed").find(qn("w:basedOn")).get(W_VAL) == "FixtureBranch"


def test_a_reference_naming_nothing_is_refused():
    """get_style reports such a chain as a warning; the writers do not make one."""
    pkg = _pkg("simple")
    for field in ("based_on", "next", "link"):
        with pytest.raises(LocatorError) as caught:
            create_style(pkg, {"name": "Probe Dangling", "style_id": "PD", field: "Ghost"})
        assert caught.value.code == "not_found"


def test_a_style_may_name_itself_as_its_next():
    """A body style whose next is itself is the common case, not a cycle."""
    pkg = _pkg("simple")
    info = create_style(pkg, {"name": "Probe Self", "next": "ProbeSelf"})

    assert info.next_style == "ProbeSelf"


# --------------------------------------------------------------------------
# Updating
# --------------------------------------------------------------------------


def test_update_changes_only_what_it_names():
    pkg = _pkg("paragraph_styles")
    before = get_style(pkg, "FixtureBody")

    update_style(pkg, "FixtureBody", {"run_props": {"bold": True}})

    after = get_style(pkg, "FixtureBody")
    assert after.run_props["bold"] is True
    # The size and the paragraph properties the style already had are untouched.
    assert after.run_props["size_pt"] == before.run_props["size_pt"]
    assert after.paragraph_props == before.paragraph_props


def test_none_removes_a_property_so_what_the_style_inherits_applies_again():
    pkg = _pkg("style_inheritance")
    assert get_style(pkg, "FixtureLeaf").run_props["italic"] is True

    update_style(pkg, "FixtureLeaf", {"run_props": {"italic": None}})

    detail = get_style(pkg, "FixtureLeaf")
    assert "italic" not in detail.run_props
    # What it inherits from the chain is back in charge.
    assert detail.resolved["run"]["bold"] is True


def test_one_setting_of_spacing_leaves_the_others_alone():
    """The way Word inherits w:spacing, and the way get_style reports it."""
    pkg = _pkg("paragraph_styles")
    assert get_style(pkg, "FixtureBody").paragraph_props["spacing"] == {
        "before": 120,
        "after": 120,
    }

    update_style(pkg, "FixtureBody", {"paragraph_props": {"spacing": {"after": 240}}})

    assert get_style(pkg, "FixtureBody").paragraph_props["spacing"] == {
        "before": 120,
        "after": 240,
    }


def test_update_can_remove_a_whole_property_block():
    pkg = _pkg("paragraph_styles")
    update_style(pkg, "FixtureNote", {"paragraph_props": {"indent": None}})

    assert "indent" not in get_style(pkg, "FixtureNote").paragraph_props


def test_update_rewires_the_metadata():
    pkg = _pkg("style_inheritance")
    info = update_style(
        pkg,
        "FixtureLeaf",
        {"name": "Fixture Leaf Renamed", "q_format": True, "ui_priority": 7, "next": None},
    )

    assert info.name == "Fixture Leaf Renamed"
    assert info.q_format is True
    assert info.ui_priority == 7
    assert info.next_style is None
    # Renaming does not change the id, so every w:pStyle still resolves.
    assert info.style_id == "FixtureLeaf"
    assert len(find_style_usage(pkg, "FixtureLeaf")) == 1


def test_update_refuses_to_change_the_id_or_the_family():
    pkg = _pkg("style_inheritance")
    with pytest.raises(ValueError, match="cannot be changed"):
        update_style(pkg, "FixtureLeaf", {"style_id": "Other"})
    with pytest.raises(ValueError, match="orphan"):
        update_style(pkg, "FixtureLeaf", {"family": "character"})


def test_update_refuses_a_name_another_style_already_has():
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        update_style(pkg, "FixtureLeaf", {"name": "Fixture Root"})
    assert caught.value.code == "already_exists"


def test_update_of_an_undefined_style_is_not_found():
    pkg = _pkg("simple")
    with pytest.raises(LocatorError) as caught:
        update_style(pkg, "Ghost", {"run_props": {"bold": True}})
    assert caught.value.code == "not_found"


def test_updating_a_style_touches_nothing_but_the_style_sheet(tmp_path):
    source = build("style_inheritance")
    pkg = DocxPackage.open(source)
    update_style(pkg, "FixtureLeaf", {"run_props": {"bold": False, "size_pt": 9.0}})
    target = pkg.save(tmp_path / "updated.docx")

    assert_unchanged_except(
        snapshot(source), snapshot(target.read_bytes()), parts={STYLES_PART}
    )


# --------------------------------------------------------------------------
# Cloning
# --------------------------------------------------------------------------


def test_clone_keeps_what_the_minimal_model_does_not_decode():
    """The whole reason to clone rather than describe the style again."""
    pkg = _pkg("character_styles")
    assert "noProof" not in str(get_style(pkg, "FixtureCode").run_props)

    clone_style(pkg, "FixtureCode", "Fixture Code Wide")

    properties = _element(pkg, "FixtureCodeWide").find(qn("w:rPr"))
    assert properties.find(qn("w:noProof")) is not None
    assert properties.find(qn("w:rFonts")).get(qn("w:ascii")) == "Consolas"


def test_clone_is_a_custom_style_and_drops_the_link():
    """A link is reciprocal: the original's twin still points at the original."""
    pkg = _pkg("style_inheritance")
    info = clone_style(pkg, "FixtureBranch", "Fixture Branch Copy")

    assert info.style_id == "FixtureBranchCopy"
    assert info.builtin is False
    assert info.link is None
    # The rest of the wiring comes along.
    assert info.based_on == "FixtureRoot"
    assert info.next_style == "FixtureRoot"


def test_clone_applies_its_overrides_to_the_copy_only():
    pkg = _pkg("style_inheritance")
    clone_style(
        pkg, "FixtureLeaf", "Fixture Leaf Loud", {"run_props": {"size_pt": 20.0}}
    )

    assert get_style(pkg, "FixtureLeafLoud").run_props["size_pt"] == 20.0
    assert "size_pt" not in get_style(pkg, "FixtureLeaf").run_props


def test_clone_refuses_a_name_already_taken():
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        clone_style(pkg, "FixtureLeaf", "Fixture Root")
    assert caught.value.code == "already_exists"


def test_a_clone_declares_its_namespaces_once(tmp_path):
    """A deep copy that carried its own xmlns would bloat every clone."""
    pkg = _pkg("character_styles")
    clone_style(pkg, "FixtureCode", "Fixture Code Copy")
    target = pkg.save(tmp_path / "cloned.docx")

    with zipfile.ZipFile(target) as archive:
        styles = archive.read("word/styles.xml").decode("utf-8")
    start = styles.index('w:styleId="FixtureCodeCopy"')
    assert "xmlns:w=" not in styles[start : start + 400]
    assert validate_package(target) == []


# --------------------------------------------------------------------------
# Deleting
# --------------------------------------------------------------------------


def test_delete_refuses_a_style_in_use():
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        delete_style(pkg, "FixtureLeaf")

    assert caught.value.code == "style_in_use"
    assert "reassign_to" in str(caught.value)
    # And the style is still there.
    assert get_style(pkg, "FixtureLeaf").info.style_id == "FixtureLeaf"


def test_delete_reassigns_every_place_that_used_the_style():
    pkg = _pkg("style_inheritance")
    before = find_style_usage(pkg, "FixtureLeaf")
    assert before

    deletion = delete_style(pkg, "FixtureLeaf", "FixtureRoot")

    assert deletion.style_id == "FixtureLeaf"
    assert deletion.reassigned_to == "FixtureRoot"
    assert len(deletion.usages) == len(before)
    assert find_style_usage(pkg, "FixtureLeaf") == []
    assert len(find_style_usage(pkg, "FixtureRoot")) == 1 + len(before)
    with pytest.raises(PackageError):
        get_style(pkg, "FixtureLeaf")


def test_delete_reassigns_a_run_style_too():
    pkg = _pkg("character_styles")
    assert [usage.kind for usage in find_style_usage(pkg, "FixtureEmphasis")] == [
        "run",
        "run",
    ]

    delete_style(pkg, "FixtureEmphasis", "Fixture Code")

    assert find_style_usage(pkg, "FixtureEmphasis") == []
    assert len(find_style_usage(pkg, "FixtureCode")) == 3


def test_delete_repoints_the_styles_that_named_it():
    """The half of a deletion nothing outside the style sheet can check."""
    pkg = _pkg("style_inheritance")
    assert get_style(pkg, "FixtureLeaf").info.based_on == "FixtureBranch"

    deletion = delete_style(pkg, "FixtureBranch", "FixtureRoot")

    assert set(deletion.references) >= {"FixtureBranchChar", "FixtureLeaf"}
    assert get_style(pkg, "FixtureLeaf").info.based_on == "FixtureRoot"
    # No chain is left pointing at a style the document no longer defines.
    assert get_style(pkg, "FixtureLeaf").warnings == ()


def test_delete_without_a_reassignment_drops_the_references_instead():
    pkg = _pkg("simple")
    create_style(pkg, {"name": "Probe Base"})
    create_style(pkg, {"name": "Probe Child", "based_on": "ProbeBase", "next": "ProbeBase"})

    deletion = delete_style(pkg, "ProbeBase")

    assert deletion.reassigned_to is None
    assert deletion.usages == ()
    assert deletion.references == ("ProbeChild",)
    child = get_style(pkg, "ProbeChild").info
    assert child.based_on is None
    assert child.next_style is None


def test_delete_refuses_the_default_style_of_a_family():
    pkg = _pkg("simple")
    with pytest.raises(ValueError, match="default"):
        delete_style(pkg, "Normal")


def test_delete_refuses_a_reassignment_of_another_family():
    pkg = _pkg("style_inheritance")
    with pytest.raises(LocatorError) as caught:
        delete_style(pkg, "FixtureLeaf", "FixtureBranchChar")
    assert caught.value.code == "wrong_style_type"


def test_delete_needs_the_package_to_walk_every_story():
    with pytest.raises(TypeError, match="DocxPackage"):
        delete_style(object(), "FixtureLeaf")  # type: ignore[arg-type]


def test_a_document_stays_valid_after_a_deletion(tmp_path):
    pkg = _pkg("style_inheritance")
    delete_style(pkg, "FixtureBranch", "FixtureRoot")
    target = pkg.save(tmp_path / "deleted.docx")

    assert validate_package(target) == []


# --------------------------------------------------------------------------
# The part itself
# --------------------------------------------------------------------------


def test_a_package_with_no_style_relationship_gets_a_style_sheet():
    pkg = _pkg("simple")
    document_part = pkg.document_part
    for rId, relationship in list(document_part.rels.items()):
        if relationship.reltype.endswith("/styles"):
            document_part.drop_rel(rId)
    assert styles_root(pkg) is None

    info = create_style(pkg, {"name": "Only Style"})

    assert info.style_id == "OnlyStyle"
    assert [existing.style_id for existing in list_styles(pkg)] == ["OnlyStyle"]
    assert pkg.find_part(STYLES_PARTNAME) is not None


def test_reading_a_style_sheet_that_is_not_there_says_so():
    pkg = _pkg("simple")
    document_part = pkg.document_part
    for rId, relationship in list(document_part.rels.items()):
        if relationship.reltype.endswith("/styles"):
            document_part.drop_rel(rId)

    with pytest.raises(PackageError, match="no style part"):
        update_style(pkg, "Normal", {"run_props": {"bold": True}})


# --------------------------------------------------------------------------
# The callers: core/styles.py
# --------------------------------------------------------------------------


def _without_headings(pkg: DocxPackage) -> DocxPackage:
    """`pkg` with its nine heading styles removed."""
    root = styles_root(pkg)
    for style in list(root.findall(W_STYLE)):
        if (style.get(W_STYLE_ID) or "") in {f"Heading{level}" for level in range(1, 10)}:
            root.remove(style)
    return pkg


def test_ensure_heading_style_creates_real_heading_styles():
    from word_document_server.core.styles import ensure_heading_style

    pkg = _without_headings(_pkg("simple"))
    assert not [info for info in list_styles(pkg) if info.style_id == "Heading1"]

    ensure_heading_style(pkg)

    detail = get_style(pkg, "Heading3")
    assert detail.info.name == "heading 3"
    assert detail.info.based_on == "Normal"
    assert detail.info.next_style == "Normal"
    assert detail.info.q_format is True
    assert detail.info.ui_priority == 9
    # The outline level is what makes a heading a heading rather than a big
    # bold paragraph: the navigation pane and every TOC read it.
    assert detail.paragraph_props["outline_level"] == 2
    assert detail.run_props["bold"] is True


def test_ensure_heading_style_leaves_the_headings_a_document_already_has():
    from word_document_server.core.styles import ensure_heading_style

    pkg = _pkg("simple")
    before = get_style(pkg, "Heading1")

    ensure_heading_style(pkg)

    after = get_style(pkg, "Heading1")
    assert after.run_props == before.run_props
    assert after.paragraph_props == before.paragraph_props


def test_ensure_heading_style_works_on_a_python_docx_document(tmp_path):
    """content_tools and document_tools hold a Document, not a package."""
    from docx import Document as PDocument

    from word_document_server.core.styles import ensure_heading_style

    source = tmp_path / "headless.docx"
    pkg = _without_headings(_pkg("simple"))
    pkg.save(source)

    document = PDocument(str(source))
    ensure_heading_style(document)
    document.add_heading("A real heading", level=1)
    document.save(str(source))

    assert get_style(DocxPackage.open(source), "Heading1").paragraph_props[
        "outline_level"
    ] == 0


def test_core_create_style_writes_the_style_it_reports():
    from docx.enum.style import WD_STYLE_TYPE

    from word_document_server.core.styles import create_style as core_create_style

    pkg = _pkg("style_inheritance")
    info = core_create_style(
        pkg,
        "Legacy Note",
        WD_STYLE_TYPE.PARAGRAPH,
        base_style="Fixture Root",
        font_properties={"bold": True, "size": 13, "color": "red", "name": "Verdana"},
        paragraph_properties={"alignment": "center", "spacing": 1.5},
    )

    assert info.style_id == "LegacyNote"
    detail = get_style(pkg, "LegacyNote")
    assert detail.info.based_on == "FixtureRoot"
    assert detail.run_props == {
        "bold": True,
        "size_pt": 13.0,
        "font": {"value": "Verdana", "theme": None},
        # A font named without a script goes to the complex-script slot too,
        # which is what apply_rpr writes and what Word does.
        "font_cs": {"value": "Verdana", "theme": None},
        "color": {"value": "FF0000", "theme": None, "tint": None, "shade": None},
    }
    assert detail.paragraph_props == {
        "alignment": "center",
        "spacing": {"line": 360, "line_rule": "auto"},
    }


def test_core_create_style_makes_character_styles_too():
    from word_document_server.core.styles import create_style as core_create_style

    pkg = _pkg("simple")
    info = core_create_style(pkg, "Legacy Mark", "character", font_properties={"italic": True})

    assert info.family == "character"


def test_core_create_style_leaves_an_existing_style_alone():
    """Redefining a style reformats every paragraph that names it."""
    from docx.enum.style import WD_STYLE_TYPE

    from word_document_server.core.styles import create_style as core_create_style

    pkg = _pkg("paragraph_styles")
    before = get_style(pkg, "FixtureBody").run_props

    info = core_create_style(
        pkg, "Fixture Body", WD_STYLE_TYPE.PARAGRAPH, font_properties={"bold": True}
    )

    assert info.style_id == "FixtureBody"
    assert get_style(pkg, "FixtureBody").run_props == before


def test_create_custom_style_refuses_a_colour_it_cannot_read(tmp_path):
    from word_document_server.tools.format_tools import create_custom_style

    path = tmp_path / "styled.docx"
    path.write_bytes(build("simple"))

    result = _run(create_custom_style(str(path), "Probe Colour", color="mauve"))

    assert "Invalid color" in result
    assert not [info for info in list_styles(DocxPackage.open(path)) if info.style_id == "ProbeColour"]


def test_create_custom_style_touches_nothing_but_the_style_sheet(tmp_path):
    from word_document_server.tools.format_tools import create_custom_style

    path = tmp_path / "styled.docx"
    source = build("combined")
    path.write_bytes(source)

    result = _run(create_custom_style(str(path), "Probe Tool", bold=True, font_size=13))

    assert "created successfully" in result
    assert_unchanged_except(
        snapshot(source), snapshot(path.read_bytes()), parts={STYLES_PART}
    )
    assert get_style(DocxPackage.open(path), "ProbeTool").run_props == {
        "bold": True,
        "size_pt": 13.0,
    }


# --------------------------------------------------------------------------
# The callers: the V2 tools
# --------------------------------------------------------------------------


def test_the_three_write_tools_are_registered():
    tools = v2_tools()
    assert {"doc_create_style", "doc_update_style", "doc_delete_style"} <= set(tools)


def test_doc_create_style_reports_the_style_it_wrote(tmp_path):
    path = tmp_path / "doc.docx"
    path.write_bytes(build("style_inheritance"))

    report = _run(
        _tool("doc_create_style")(
            str(path),
            {
                "name": "Tool Body",
                "based_on": "Fixture Root",
                "q_format": True,
                "run_props": {"size_pt": 11.0},
            },
        )
    )

    assert report["status"] == "ok"
    assert report["saved"] is True
    assert report["style"]["style_id"] == "ToolBody"
    assert report["style"]["based_on"] == "FixtureRoot"
    assert get_style(DocxPackage.open(path), "ToolBody").run_props == {"size_pt": 11.0}


def test_doc_create_style_dry_run_writes_nothing(tmp_path):
    path = tmp_path / "doc.docx"
    source = build("simple")
    path.write_bytes(source)

    report = _run(_tool("doc_create_style")(str(path), {"name": "Tool Ghost"}, True))

    assert report["status"] == "ok"
    assert report["dry_run"] is True
    assert report["saved"] is False
    assert path.read_bytes() == source


def test_doc_create_style_maps_a_collision_to_a_stable_code(tmp_path):
    path = tmp_path / "doc.docx"
    path.write_bytes(build("style_inheritance"))

    report = _run(_tool("doc_create_style")(str(path), {"name": "Fixture Root"}))

    assert report == {
        "status": "error",
        "code": "already_exists",
        "message": report["message"],
    }


def test_doc_create_style_refuses_a_spec_key_it_does_not_know(tmp_path):
    path = tmp_path / "doc.docx"
    path.write_bytes(build("simple"))

    report = _run(_tool("doc_create_style")(str(path), {"name": "Tool Typo", "bold": True}))

    assert report["status"] == "error"
    assert report["code"] == "invalid_argument"
    assert "bold" in report["message"]


def test_doc_update_style_says_how_many_places_it_affects(tmp_path):
    path = tmp_path / "doc.docx"
    path.write_bytes(build("style_inheritance"))

    report = _run(
        _tool("doc_update_style")(str(path), "Fixture Root", {"run_props": {"bold": True}})
    )

    assert report["status"] == "ok"
    assert report["usages"] == 1
    assert report["warnings"] == [
        "1 place(s) use 'FixtureRoot' and are affected by this change"
    ]
    assert get_style(DocxPackage.open(path), "FixtureRoot").run_props["bold"] is True


def test_doc_delete_style_refuses_a_style_in_use(tmp_path):
    path = tmp_path / "doc.docx"
    source = build("style_inheritance")
    path.write_bytes(source)

    report = _run(_tool("doc_delete_style")(str(path), "FixtureLeaf"))

    assert report["status"] == "error"
    assert report["code"] == "style_in_use"
    assert path.read_bytes() == source


def test_doc_delete_style_reports_what_moved(tmp_path):
    path = tmp_path / "doc.docx"
    path.write_bytes(build("style_inheritance"))

    report = _run(_tool("doc_delete_style")(str(path), "FixtureLeaf", "Fixture Root"))

    assert report["status"] == "ok"
    assert report["style_id"] == "FixtureLeaf"
    assert report["reassigned_to"] == "FixtureRoot"
    assert report["changes"] == [
        {
            "story": "document",
            "paragraph": 3,
            "before": "Leaf, based on branch.",
            "after": "Leaf, based on branch.",
        }
    ]
    assert find_style_usage(DocxPackage.open(path), "FixtureLeaf") == []
