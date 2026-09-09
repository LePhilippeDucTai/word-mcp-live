"""Tests for :mod:`word_document_server.engine.ranges`.

The cases below are written against hand-built paragraphs whose XML is visible
in the test: a range operation is judged on the tree it leaves behind, not only
on the text it reads back, because the whole point of this layer is what it does
*not* disturb.  ``tests/engine/test_ranges_properties.py`` then runs the same
operations over every fixture and checks the invariants at package scale.
"""

from __future__ import annotations

import pytest
from docx import Document
from lxml import etree

from word_document_server.engine.errors import (
    InvalidText,
    LocatorError,
    UnsupportedRange,
)
from word_document_server.engine.ranges import (
    delete_range,
    insert_text,
    replace_range,
    resolve,
    split_run,
)
from word_document_server.engine.textmodel import visible_text
from word_document_server.engine.xmlns import NAMESPACES, qn

W_R = qn("w:r")
W_T = qn("w:t")
W_RPR = qn("w:rPr")
XML_SPACE = qn("xml:space")


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------


def paragraph_from(children: str) -> etree._Element:
    """Return a standalone ``w:p`` whose children are the given XML fragment."""
    declarations = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in NAMESPACES.items())
    return etree.fromstring(f"<w:p {declarations}>{children}</w:p>")


def runs(paragraph: etree._Element) -> list[etree._Element]:
    """Every ``w:r`` of the paragraph, in document order."""
    return list(paragraph.iter(W_R))


def names(element: etree._Element) -> list[str]:
    """Local names of the element's children, in order."""
    return [etree.QName(child).localname for child in element]


def properties(run: etree._Element) -> str:
    """Canonical form of a run's ``w:rPr``, "" when it has none."""
    found = run.find(W_RPR)
    if found is None:
        return ""
    return etree.tostring(found, method="c14n", exclusive=True).decode()


def shape(paragraph: etree._Element) -> list[tuple[str, str]]:
    """``(local name, visible text)`` of every run, in document order."""
    return [
        (etree.QName(run).localname, "".join(t.text or "" for t in run.findall(W_T)))
        for run in runs(paragraph)
    ]


# --------------------------------------------------------------------------------------
# split_run
# --------------------------------------------------------------------------------------


def test_split_clones_the_properties_and_the_run_attributes() -> None:
    paragraph = paragraph_from(
        '<w:r w:rsidR="00AA00BB"><w:rPr><w:b/><w:i/></w:rPr>'
        '<w:t xml:space="preserve">abcdef</w:t></w:r>'
    )
    left, right = split_run(paragraph[0], 3)

    assert left is paragraph[0]
    assert right is paragraph[1]
    assert shape(paragraph) == [("r", "abc"), ("r", "def")]
    assert properties(left) == properties(right)
    assert names(right) == ["rPr", "t"]
    assert left.get(qn("w:rsidR")) == right.get(qn("w:rsidR")) == "00AA00BB"
    # The clone is a copy, not the same element shared by two runs.
    assert left.find(W_RPR) is not right.find(W_RPR)


def test_split_marks_both_halves_as_space_preserving() -> None:
    paragraph = paragraph_from("<w:r><w:t>a b</w:t></w:r>")
    left, right = split_run(paragraph[0], 2)

    assert left.find(W_T).get(XML_SPACE) == "preserve"
    assert right.find(W_T).get(XML_SPACE) == "preserve"
    assert visible_text(paragraph) == "a b"


def test_split_keeps_the_order_of_the_children() -> None:
    paragraph = paragraph_from(
        "<w:r><w:t>ab</w:t><w:tab/><w:t>cd</w:t><w:br/></w:r>"
    )
    left, right = split_run(paragraph[0], 3)

    assert names(left) == ["t", "tab"]
    assert names(right) == ["t", "br"]
    assert visible_text(paragraph) == "ab\tcd\n"


@pytest.mark.parametrize("offset", [0, 6])
def test_split_at_the_bounds_leaves_the_tree_untouched(offset: int) -> None:
    paragraph = paragraph_from("<w:r><w:t>abcdef</w:t></w:r>")
    before = etree.tostring(paragraph)
    left, right = split_run(paragraph[0], offset)

    assert etree.tostring(paragraph) == before
    assert (left, right) == ((None, paragraph[0]) if offset == 0 else (paragraph[0], None))


def test_split_twice_at_the_same_offset_splits_once() -> None:
    paragraph = paragraph_from("<w:r><w:rPr><w:b/></w:rPr><w:t>abcdef</w:t></w:r>")
    split_run(paragraph[0], 3)
    once = etree.tostring(paragraph)
    split_run(paragraph[0], 3)

    assert etree.tostring(paragraph) == once


def test_zero_width_content_goes_to_the_side_the_caller_names() -> None:
    leading = "<w:r><w:drawing/><w:t>ab</w:t></w:r>"
    paragraph = paragraph_from(leading)
    left, right = split_run(paragraph[0], 0, zero_width="left")
    assert names(left) == ["drawing"]
    assert names(right) == ["t"]

    paragraph = paragraph_from(leading)
    assert split_run(paragraph[0], 0, zero_width="right") == (None, paragraph[0])

    trailing = "<w:r><w:t>ab</w:t><w:drawing/></w:r>"
    paragraph = paragraph_from(trailing)
    left, right = split_run(paragraph[0], 2, zero_width="right")
    assert names(left) == ["t"]
    assert names(right) == ["drawing"]

    paragraph = paragraph_from(trailing)
    assert split_run(paragraph[0], 2, zero_width="left") == (paragraph[0], None)


def test_split_of_a_run_with_no_content_is_a_no_op() -> None:
    paragraph = paragraph_from("<w:r><w:rPr><w:b/></w:rPr></w:r>")
    assert split_run(paragraph[0], 0) == (None, paragraph[0])


@pytest.mark.parametrize("offset", [-1, 7])
def test_split_rejects_an_offset_outside_the_run(offset: int) -> None:
    paragraph = paragraph_from("<w:r><w:t>abcdef</w:t></w:r>")
    with pytest.raises(LocatorError) as error:
        split_run(paragraph[0], offset)
    assert error.value.code == "out-of-range"


def test_split_rejects_what_is_not_a_run_in_a_paragraph() -> None:
    paragraph = paragraph_from("<w:r><w:t>abc</w:t></w:r>")
    with pytest.raises(TypeError):
        split_run(paragraph, 0)
    detached = etree.fromstring(
        f'<w:r xmlns:w="{NAMESPACES["w"]}"><w:t>abc</w:t></w:r>'
    )
    with pytest.raises(TypeError):
        split_run(detached, 1)


def test_a_one_character_element_can_only_be_cut_on_its_edges() -> None:
    paragraph = paragraph_from('<w:r><w:sym w:font="Symbol" w:char="F0B7"/></w:r>')
    assert split_run(paragraph[0], 1) == (paragraph[0], None)
    assert split_run(paragraph[0], 0) == (None, paragraph[0])


# --------------------------------------------------------------------------------------
# resolve
# --------------------------------------------------------------------------------------


SPANNING = (
    '<w:r><w:t xml:space="preserve">Hello </w:t></w:r>'
    '<w:bookmarkStart w:id="1" w:name="Middle"/>'
    "<w:r><w:rPr><w:b/></w:rPr><w:t>brave</w:t></w:r>"
    '<w:bookmarkEnd w:id="1"/>'
    '<w:r><w:t xml:space="preserve"> world</w:t></w:r>'
)


def test_resolve_returns_the_runs_the_range_covers_whole() -> None:
    paragraph = paragraph_from(SPANNING)
    pieces = resolve(paragraph, 4, 13)

    assert pieces.text == "o brave w"
    assert pieces.start == 4 and pieces.end == 13
    assert not pieces.is_empty
    assert [
        "".join(t.text or "" for t in run.findall(W_T)) for run in pieces.runs
    ] == ["o ", "brave", " w"]
    assert visible_text(paragraph) == "Hello brave world"


def test_resolve_reports_the_markers_strictly_inside() -> None:
    paragraph = paragraph_from(SPANNING)
    pieces = resolve(paragraph, 4, 13)
    assert [etree.QName(marker).localname for marker in pieces.markers] == [
        "bookmarkStart",
        "bookmarkEnd",
    ]

    # A marker sitting exactly on a boundary belongs to neither side.
    pieces = resolve(paragraph_from(SPANNING), 6, 11)
    assert pieces.markers == ()


def test_resolve_only_splits_and_never_changes_the_text() -> None:
    paragraph = paragraph_from(SPANNING)
    before = visible_text(paragraph)
    resolve(paragraph, 3, 9)
    assert visible_text(paragraph) == before
    assert shape(paragraph) == [
        ("r", "Hel"),
        ("r", "lo "),
        ("r", "bra"),
        ("r", "ve"),
        ("r", " world"),
    ]


@pytest.mark.parametrize("start,end", [(-1, 2), (2, 1), (0, 18)])
def test_resolve_rejects_offsets_outside_the_paragraph(start: int, end: int) -> None:
    with pytest.raises(LocatorError) as error:
        resolve(paragraph_from(SPANNING), start, end)
    assert error.value.code == "out-of-range"


@pytest.mark.parametrize("operation", ["delete", "replace"])
def test_resolve_refuses_an_empty_range_for_a_destructive_operation(operation: str) -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(SPANNING), 4, 4, operation=operation)
    assert error.value.reason == "empty-range"


@pytest.mark.parametrize("operation", ["read", "insert"])
def test_resolve_accepts_an_empty_range_for_an_insertion(operation: str) -> None:
    pieces = resolve(paragraph_from(SPANNING), 4, 4, operation=operation)
    assert pieces.is_empty
    assert pieces.runs == () and pieces.markers == () and pieces.text == ""


COMPLEX_FIELD = (
    '<w:r><w:t xml:space="preserve">Page </w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    "<w:r><w:t>12</w:t></w:r>"
    '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    '<w:r><w:t xml:space="preserve"> of 30</w:t></w:r>'
)


def test_resolve_refuses_a_range_that_cuts_a_field() -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(COMPLEX_FIELD), 3, 6, operation="delete")
    assert error.value.reason == "crosses-field"


def test_resolve_refuses_to_remove_a_whole_field_but_lets_it_be_read() -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(COMPLEX_FIELD), 0, 13, operation="replace")
    assert error.value.reason == "contains-field"
    assert resolve(paragraph_from(COMPLEX_FIELD), 0, 13, operation="read").text == "Page 12 of 30"


def test_resolve_refuses_an_insertion_inside_a_field() -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(COMPLEX_FIELD), 6, 6, operation="insert")
    assert error.value.reason == "inside-field"


DELETION = (
    '<w:r><w:t xml:space="preserve">Kept </w:t></w:r>'
    '<w:del w:id="9" w:author="A" w:date="2024-01-01T00:00:00Z">'
    '<w:r><w:delText>gone</w:delText></w:r>'
    "</w:del>"
    '<w:r><w:t xml:space="preserve"> tail</w:t></w:r>'
)


@pytest.mark.parametrize("operation", ["read", "delete", "replace"])
def test_resolve_refuses_a_range_that_straddles_deleted_content(operation: str) -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(DELETION), 3, 8, operation=operation)
    assert error.value.reason == "hidden-content"


def test_deleted_content_on_a_boundary_is_not_in_the_way() -> None:
    paragraph = paragraph_from(DELETION)
    assert visible_text(paragraph) == "Kept  tail"
    assert delete_range(paragraph, 0, 5) == "Kept "
    assert visible_text(paragraph) == " tail"
    assert paragraph.find(qn("w:del")) is not None
    assert paragraph.find(qn("w:del"))[0].find(qn("w:delText")).text == "gone"


IMAGE = (
    '<w:r><w:t xml:space="preserve">before </w:t></w:r>'
    "<w:r><w:drawing/></w:r>"
    '<w:r><w:t xml:space="preserve"> after</w:t></w:r>'
)


def test_resolve_refuses_to_delete_an_object_it_cannot_rewrite() -> None:
    with pytest.raises(UnsupportedRange) as error:
        resolve(paragraph_from(IMAGE), 3, 10, operation="delete")
    assert error.value.reason == "opaque-content"


def test_an_object_can_be_read_over_and_survives_a_neighbouring_deletion() -> None:
    assert resolve(paragraph_from(IMAGE), 3, 10, operation="read").text == "ore  af"

    paragraph = paragraph_from(IMAGE)
    # The drawing sits exactly at offset 7, where the range ends: a boundary
    # object is outside the range.
    assert delete_range(paragraph, 0, 7) == "before "
    assert paragraph.find(qn("w:drawing")) is None  # it lives inside its run
    assert len(list(paragraph.iter(qn("w:drawing")))) == 1
    assert visible_text(paragraph) == " after"


# --------------------------------------------------------------------------------------
# delete_range
# --------------------------------------------------------------------------------------


def test_delete_returns_the_text_it_removed_and_keeps_the_rest_formatted() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    kept = [properties(run) for run in runs(paragraph)]

    assert delete_range(paragraph, 2, 8) == "ldital"
    assert visible_text(paragraph) == "boic"
    assert shape(paragraph) == [("r", "bo"), ("r", "ic")]
    assert [properties(run) for run in runs(paragraph)] == kept


def test_delete_keeps_the_markers_of_the_range_at_its_start() -> None:
    paragraph = paragraph_from(SPANNING)
    delete_range(paragraph, 4, 13)

    assert visible_text(paragraph) == "Hellorld"
    assert names(paragraph) == ["r", "bookmarkStart", "bookmarkEnd", "r"]
    assert paragraph.find(qn("w:bookmarkStart")).get(qn("w:name")) == "Middle"


def test_delete_removes_the_runs_it_empties() -> None:
    paragraph = paragraph_from(SPANNING)
    delete_range(paragraph, 6, 11)
    assert shape(paragraph) == [("r", "Hello "), ("r", " world")]


HYPERLINK = (
    '<w:r><w:t xml:space="preserve">see </w:t></w:r>'
    '<w:hyperlink r:id="rId7" w:history="1">'
    '<w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>the site</w:t></w:r>'
    "</w:hyperlink>"
    "<w:r><w:t>.</w:t></w:r>"
)


def test_delete_empties_a_hyperlink_without_dropping_its_relationship() -> None:
    paragraph = paragraph_from(HYPERLINK)
    assert delete_range(paragraph, 4, 12) == "the site"

    link = paragraph.find(qn("w:hyperlink"))
    assert link is not None
    assert link.get(qn("r:id")) == "rId7"
    assert len(link) == 0
    assert visible_text(paragraph) == "see ."


def test_delete_inside_a_tracked_insertion_keeps_the_revision() -> None:
    paragraph = paragraph_from(
        '<w:ins w:id="3" w:author="A" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:t>added</w:t></w:r></w:ins>"
    )
    assert delete_range(paragraph, 0, 5) == "added"
    insertion = paragraph.find(qn("w:ins"))
    assert insertion is not None and len(insertion) == 0
    assert visible_text(paragraph) == ""


def test_delete_leaves_the_paragraph_properties_alone() -> None:
    paragraph = paragraph_from(
        '<w:pPr><w:pStyle w:val="Heading1"/><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
        "<w:r><w:t>title</w:t></w:r>"
    )
    before = etree.tostring(paragraph.find(qn("w:pPr")))
    delete_range(paragraph, 0, 5)
    assert etree.tostring(paragraph.find(qn("w:pPr"))) == before


# --------------------------------------------------------------------------------------
# insert_text
# --------------------------------------------------------------------------------------


TWO_RUNS = (
    '<w:r><w:rPr><w:b/></w:rPr><w:t>bold</w:t></w:r>'
    '<w:r><w:rPr><w:i/></w:rPr><w:t>italic</w:t></w:r>'
)


@pytest.mark.parametrize("side,expected", [("left", "b"), ("right", "i")])
def test_insert_takes_the_properties_of_the_side_it_is_told_to(side: str, expected: str) -> None:
    paragraph = paragraph_from(TWO_RUNS)
    (inserted,) = insert_text(paragraph, 4, "X", rpr_from=side)

    assert visible_text(paragraph) == "boldXitalic"
    assert names(inserted.find(W_RPR)) == [expected]


def test_insert_accepts_an_explicit_run_or_run_properties() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    wanted = etree.fromstring(
        f'<w:rPr xmlns:w="{NAMESPACES["w"]}"><w:u w:val="single"/></w:rPr>'
    )
    (inserted,) = insert_text(paragraph, 4, "X", rpr_from=wanted)
    assert names(inserted.find(W_RPR)) == ["u"]
    assert inserted.find(W_RPR) is not wanted

    paragraph = paragraph_from(TWO_RUNS)
    (inserted,) = insert_text(paragraph, 0, "X", rpr_from=runs(paragraph)[1])
    assert names(inserted.find(W_RPR)) == ["i"]


def test_insert_falls_back_to_the_other_side_when_there_is_no_run() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    (inserted,) = insert_text(paragraph, 0, "X", rpr_from="left")

    assert names(inserted.find(W_RPR)) == ["b"]
    assert visible_text(paragraph) == "Xbolditalic"


def test_insert_translates_tabs_and_breaks() -> None:
    paragraph = paragraph_from("<w:r><w:t>ab</w:t></w:r>")
    (inserted,) = insert_text(paragraph, 2, "x\ty\nz")

    assert names(inserted) == ["t", "tab", "t", "br", "t"]
    assert visible_text(paragraph) == "abx\ty\nz"
    assert all(carrier.get(XML_SPACE) == "preserve" for carrier in inserted.findall(W_T))


def test_insert_normalises_carriage_returns_to_line_breaks() -> None:
    paragraph = paragraph_from("<w:r><w:t>ab</w:t></w:r>")
    (inserted,) = insert_text(paragraph, 2, "x\r\ny\rz")

    assert names(inserted) == ["t", "br", "t", "br", "t"]
    assert visible_text(paragraph) == "abx\ny\nz"


@pytest.mark.parametrize("text", ["a\x07b", "\x00", "cell\x0bsep", "lone\ud800"])
def test_insert_refuses_control_characters(text: str) -> None:
    paragraph = paragraph_from("<w:r><w:t>ab</w:t></w:r>")
    before = etree.tostring(paragraph)
    with pytest.raises(InvalidText):
        insert_text(paragraph, 1, text)
    assert etree.tostring(paragraph) == before


def test_inserting_nothing_changes_nothing() -> None:
    paragraph = paragraph_from("<w:r><w:t>ab</w:t></w:r>")
    before = etree.tostring(paragraph)
    assert insert_text(paragraph, 1, "") == ()
    assert etree.tostring(paragraph) == before


def test_insert_into_an_empty_paragraph_and_at_the_very_end() -> None:
    paragraph = paragraph_from('<w:pPr><w:pStyle w:val="Normal"/></w:pPr>')
    (inserted,) = insert_text(paragraph, 0, "first")
    assert names(paragraph) == ["pPr", "r"]
    assert visible_text(paragraph) == "first"

    insert_text(paragraph, 5, " and last")
    assert visible_text(paragraph) == "first and last"
    assert inserted.getparent() is paragraph


def test_insert_at_a_hyperlink_boundary_follows_the_side_it_was_given() -> None:
    paragraph = paragraph_from(HYPERLINK)
    (inside,) = insert_text(paragraph, 12, "!", rpr_from="left")
    assert inside.getparent() is paragraph.find(qn("w:hyperlink"))

    paragraph = paragraph_from(HYPERLINK)
    (outside,) = insert_text(paragraph, 12, "!", rpr_from="right")
    assert outside.getparent() is paragraph
    assert visible_text(paragraph) == "see the site!."


def test_insert_at_a_field_boundary_lands_outside_the_field() -> None:
    paragraph = paragraph_from(COMPLEX_FIELD)
    insert_text(paragraph, 13, "!", rpr_from="left")
    assert visible_text(paragraph) == "Page 12 of 30!"

    paragraph = paragraph_from(COMPLEX_FIELD)
    # Offset 5 opens the field: the right-hand run is its cached result, so the
    # text has to go before the ``begin`` rather than into the result.
    (inserted,) = insert_text(paragraph, 5, "#", rpr_from="right")
    assert visible_text(paragraph) == "Page #12 of 30"
    assert inserted.getnext().find(qn("w:fldChar")).get(qn("w:fldCharType")) == "begin"


def test_insert_at_a_simple_field_boundary_stays_out_of_it() -> None:
    paragraph = paragraph_from(
        '<w:r><w:t xml:space="preserve">Page </w:t></w:r>'
        '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
        '<w:r><w:t xml:space="preserve"> of</w:t></w:r>'
    )
    (before,) = insert_text(paragraph, 5, "[", rpr_from="right")
    (after,) = insert_text(paragraph, 7, "]", rpr_from="left")

    assert visible_text(paragraph) == "Page [1] of"
    assert before.getparent() is paragraph
    assert after.getparent() is paragraph
    assert len(paragraph.find(qn("w:fldSimple"))) == 1


def test_insert_has_nowhere_to_write_inside_a_field_that_owns_the_paragraph() -> None:
    only_field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>Result</w:t></w:r>"
    )
    paragraph = paragraph_from(only_field)
    with pytest.raises(UnsupportedRange) as error:
        insert_text(paragraph, 6, "!")
    assert error.value.reason == "inside-field"


def test_insert_next_to_a_deletion_stays_outside_it() -> None:
    paragraph = paragraph_from(DELETION)
    (inserted,) = insert_text(paragraph, 5, "|", rpr_from="left")

    assert inserted.getparent() is paragraph
    assert visible_text(paragraph) == "Kept | tail"
    assert paragraph.find(qn("w:del"))[0].find(qn("w:delText")).text == "gone"


# --------------------------------------------------------------------------------------
# replace_range
# --------------------------------------------------------------------------------------


def test_replace_keeps_the_formatting_of_the_first_run_of_the_range() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    (inserted,) = replace_range(paragraph, 2, 8, "NEW")

    assert visible_text(paragraph) == "boNEWic"
    assert names(inserted.find(W_RPR)) == ["b"]
    assert shape(paragraph) == [("r", "bo"), ("r", "NEW"), ("r", "ic")]


def test_replace_with_nothing_is_a_deletion() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    assert replace_range(paragraph, 2, 8, "") == ()
    assert visible_text(paragraph) == "boic"


def test_replace_writes_inside_the_markers_the_range_was_framed_by() -> None:
    paragraph = paragraph_from(SPANNING)
    replace_range(paragraph, 6, 11, "bold")

    assert visible_text(paragraph) == "Hello bold world"
    assert names(paragraph) == ["r", "bookmarkStart", "r", "bookmarkEnd", "r"]


def test_replace_keeps_a_marker_that_was_inside_the_range() -> None:
    paragraph = paragraph_from(SPANNING)
    replace_range(paragraph, 4, 13, "-")

    assert visible_text(paragraph) == "Hell-orld"
    assert names(paragraph) == ["r", "bookmarkStart", "bookmarkEnd", "r", "r"]


def test_replace_checks_the_text_before_touching_the_paragraph() -> None:
    paragraph = paragraph_from(TWO_RUNS)
    before = etree.tostring(paragraph)
    with pytest.raises(InvalidText):
        replace_range(paragraph, 2, 8, "bad\x07text")
    assert etree.tostring(paragraph) == before


def test_replace_refuses_the_same_ranges_a_deletion_refuses() -> None:
    with pytest.raises(UnsupportedRange) as error:
        replace_range(paragraph_from(IMAGE), 3, 10, "x")
    assert error.value.reason == "opaque-content"


# --------------------------------------------------------------------------------------
# Accepted inputs
# --------------------------------------------------------------------------------------


def test_a_python_docx_paragraph_is_accepted() -> None:
    document = Document()
    document.add_paragraph("hello world")
    paragraph = document.paragraphs[0]

    assert delete_range(paragraph, 0, 6) == "hello "
    assert visible_text(paragraph) == "world"


@pytest.mark.parametrize("bad", ["not a paragraph", 3])
def test_anything_that_is_not_a_paragraph_is_refused(bad: object) -> None:
    with pytest.raises(TypeError):
        resolve(bad, 0, 1)
