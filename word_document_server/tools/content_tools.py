"""
Content tools for Word Document Server.

These tools add various types of content to Word documents,
including headings, paragraphs, tables, images, and page breaks.
"""
import os
import re
from typing import List, Optional, Dict, Any
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from word_document_server.utils.document_utils import get_effective_text

from word_document_server.utils.file_utils import check_file_writeable, ensure_docx_extension, get_file_lock
from word_document_server.utils.document_utils import find_and_replace_text, replace_text_everywhere, insert_header_near_text, insert_numbered_list_near_text, insert_line_or_paragraph_near_text, replace_paragraph_block_below_header, replace_block_between_manual_anchors
from word_document_server.utils.document_utils import (
    BLOCK_TAGS,
    attach_ppr_child,
    body_element,
    body_paragraphs,
    paragraph_style_label,
    paragraph_style_names,
    styles_root,
)
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import visible_text
from word_document_server.core.styles import ensure_heading_style, ensure_table_style


_W_P = qn("w:p")
_W_PPR = qn("w:pPr")
_W_SECT_PR = qn("w:sectPr")

#: ``w:pStyle`` values -- or style names -- that name a heading and its level.
#: Word stores the built-in headings under the name ``"heading 1"`` and the id
#: ``"Heading1"``; both spellings, and the ``"Heading 1"`` a caller uses, match.
_HEADING_STYLE = re.compile(r"heading\s*([1-9])\s*$", re.IGNORECASE)

#: Gallery name Word looks for to recognise a content control as a table of contents.
_TOC_GALLERY = "Table of Contents"

_SETTINGS_PARTNAME = "/word/settings.xml"
_SETTINGS_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"
)

#: Local names of the ``w:settings`` children that follow ``w:updateFields`` in
#: the schema sequence (ECMA-376 §17.15.1.78). ``w:settings`` is an
#: ``xsd:sequence``: a ``w:updateFields`` appended at the end of a part that
#: already holds a ``w:compat`` or a ``w:rsids`` is out of order, and Word
#: offers to repair the document instead of refreshing the field.
_SETTINGS_AFTER_UPDATE_FIELDS = frozenset(
    {
        "hdrShapeDefaults",
        "footnotePr",
        "endnotePr",
        "compat",
        "docVars",
        "rsids",
        "mathPr",
        "attachedSchema",
        "themeFontLang",
        "clrSchemeMapping",
        "doNotIncludeSubdocsInStats",
        "doNotAutoCompressPictures",
        "forceUpgrade",
        "captions",
        "readModeInkLockDown",
        "smartTagType",
        "schemaLibrary",
        "shapeDefaults",
        "doNotEmbedSmartTags",
        "decimalSymbol",
        "listSeparator",
    }
)


async def add_heading(filename: str, text: str, level: int = 1,
                      font_name: Optional[str] = None, font_size: Optional[int] = None,
                      bold: Optional[bool] = None, italic: Optional[bool] = None,
                      border_bottom: bool = False) -> str:
    """Add a heading to a Word document with optional formatting.

    Args:
        filename: Path to the Word document
        text: Heading text
        level: Heading level (1-9, where 1 is the highest level)
        font_name: Font family (e.g., 'Helvetica')
        font_size: Font size in points (e.g., 14)
        bold: True/False for bold text
        italic: True/False for italic text
        border_bottom: True to add bottom border (for section headers)
    """
    filename = ensure_docx_extension(filename)

    # Ensure level is converted to integer
    try:
        level = int(level)
    except (ValueError, TypeError):
        return "Invalid parameter: level must be an integer between 1 and 9"

    # Validate level range
    if level < 1 or level > 9:
        return f"Invalid heading level: {level}. Level must be between 1 and 9."

    if not os.path.exists(filename):
        return f"Document {filename} does not exist"

    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        # Suggest creating a copy
        return f"Cannot modify document: {error_message}. Consider creating a copy first or creating a new document."

    try:
        async with get_file_lock(filename):
            doc = Document(filename)

            # Ensure heading styles exist
            ensure_heading_style(doc)

            # Try to add heading with style
            try:
                heading = doc.add_heading(text, level=level)
            except Exception as style_error:
                # If style-based approach fails, use direct formatting
                heading = doc.add_paragraph(text)
                heading.style = doc.styles['Normal']
                if heading.runs:
                    run = heading.runs[0]
                    run.bold = True
                    # Adjust size based on heading level
                    if level == 1:
                        run.font.size = Pt(16)
                    elif level == 2:
                        run.font.size = Pt(14)
                    else:
                        run.font.size = Pt(12)

            # Apply formatting to all runs in the heading
            if any([font_name, font_size, bold is not None, italic is not None]):
                for run in heading.runs:
                    if font_name:
                        run.font.name = font_name
                    if font_size:
                        run.font.size = Pt(font_size)
                    if bold is not None:
                        run.font.bold = bold
                    if italic is not None:
                        run.font.italic = italic

            # Add bottom border if requested
            if border_bottom:
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                pPr = heading._element.get_or_add_pPr()
                pBdr = OxmlElement('w:pBdr')

                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '4')  # 0.5pt border
                bottom.set(qn('w:space'), '0')
                bottom.set(qn('w:color'), '000000')

                pBdr.append(bottom)
                pPr.append(pBdr)

            doc.save(filename)
        return f"Heading '{text}' (level {level}) added to {filename}"
    except Exception as e:
        return f"Failed to add heading: {str(e)}"


async def add_paragraph(filename: str, text: str, style: Optional[str] = None,
                        font_name: Optional[str] = None, font_size: Optional[int] = None,
                        bold: Optional[bool] = None, italic: Optional[bool] = None,
                        color: Optional[str] = None) -> str:
    """Add a paragraph to a Word document with optional formatting.

    Args:
        filename: Path to the Word document
        text: Paragraph text
        style: Optional paragraph style name
        font_name: Font family (e.g., 'Helvetica', 'Times New Roman')
        font_size: Font size in points (e.g., 14, 36)
        bold: True/False for bold text
        italic: True/False for italic text
        color: RGB color as hex string (e.g., '000000' for black)
    """
    filename = ensure_docx_extension(filename)

    if not os.path.exists(filename):
        return f"Document {filename} does not exist"

    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        # Suggest creating a copy
        return f"Cannot modify document: {error_message}. Consider creating a copy first or creating a new document."

    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            paragraph = doc.add_paragraph(text)

            if style:
                try:
                    paragraph.style = style
                except KeyError:
                    # Style doesn't exist, use normal and report it
                    paragraph.style = doc.styles['Normal']
                    doc.save(filename)
                    return f"Style '{style}' not found, paragraph added with default style to {filename}"

            # Apply formatting to all runs in the paragraph
            if any([font_name, font_size, bold is not None, italic is not None, color]):
                for run in paragraph.runs:
                    if font_name:
                        run.font.name = font_name
                    if font_size:
                        run.font.size = Pt(font_size)
                    if bold is not None:
                        run.font.bold = bold
                    if italic is not None:
                        run.font.italic = italic
                    if color:
                        # Remove any '#' prefix if present
                        color_hex = color.lstrip('#')
                        run.font.color.rgb = RGBColor.from_string(color_hex)

            doc.save(filename)
        return f"Paragraph added to {filename}"
    except Exception as e:
        return f"Failed to add paragraph: {str(e)}"


async def add_table(filename: str, rows: int, cols: int, data: Optional[List[List[str]]] = None) -> str:
    """Add a table to a Word document.
    
    Args:
        filename: Path to the Word document
        rows: Number of rows in the table
        cols: Number of columns in the table
        data: Optional 2D array of data to fill the table
    """
    filename = ensure_docx_extension(filename)
    
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    
    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        # Suggest creating a copy
        return f"Cannot modify document: {error_message}. Consider creating a copy first or creating a new document."
    
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            table = doc.add_table(rows=rows, cols=cols)

            # Try to set the table style
            try:
                table.style = 'Table Grid'
            except KeyError:
                # If style doesn't exist, add basic borders
                pass

            # Fill table with data if provided
            if data:
                for i, row_data in enumerate(data):
                    if i >= rows:
                        break
                    for j, cell_text in enumerate(row_data):
                        if j >= cols:
                            break
                        table.cell(i, j).text = str(cell_text)

            doc.save(filename)
        return f"Table ({rows}x{cols}) added to {filename}"
    except Exception as e:
        return f"Failed to add table: {str(e)}"


async def add_picture(filename: str, image_path: str, width: Optional[float] = None) -> str:
    """Add an image to a Word document.
    
    Args:
        filename: Path to the Word document
        image_path: Path to the image file
        width: Optional width in inches (proportional scaling)
    """
    filename = ensure_docx_extension(filename)
    
    # Validate document existence
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    
    # Get absolute paths for better diagnostics
    abs_filename = os.path.abspath(filename)
    abs_image_path = os.path.abspath(image_path)
    
    # Validate image existence with improved error message
    if not os.path.exists(abs_image_path):
        return f"Image file not found: {abs_image_path}"
    
    # Check image file size
    try:
        image_size = os.path.getsize(abs_image_path) / 1024  # Size in KB
        if image_size <= 0:
            return f"Image file appears to be empty: {abs_image_path} (0 KB)"
    except Exception as size_error:
        return f"Error checking image file: {str(size_error)}"
    
    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(abs_filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first or creating a new document."
    
    try:
        async with get_file_lock(abs_filename):
            doc = Document(abs_filename)
            # Additional diagnostic info
            diagnostic = f"Attempting to add image ({abs_image_path}, {image_size:.2f} KB) to document ({abs_filename})"

            try:
                if width:
                    doc.add_picture(abs_image_path, width=Inches(width))
                else:
                    doc.add_picture(abs_image_path)
                doc.save(abs_filename)
            except Exception as inner_error:
                # More detailed error for the specific operation
                error_type = type(inner_error).__name__
                error_msg = str(inner_error)
                return f"Failed to add picture: {error_type} - {error_msg or 'No error details available'}\nDiagnostic info: {diagnostic}"
        return f"Picture {image_path} added to {filename}"
    except Exception as outer_error:
        # Fallback error handling
        error_type = type(outer_error).__name__
        error_msg = str(outer_error)
        return f"Document processing error: {error_type} - {error_msg or 'No error details available'}"


async def add_page_break(filename: str) -> str:
    """Add a page break to the document.
    
    Args:
        filename: Path to the Word document
    """
    filename = ensure_docx_extension(filename)
    
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    
    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first."
    
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            doc.add_page_break()
            doc.save(filename)
        return f"Page break added to {filename}."
    except Exception as e:
        return f"Failed to add page break: {str(e)}"


def _heading_level(label: Optional[str]) -> Optional[int]:
    """Heading level a paragraph style label names, or ``None``."""
    if not label:
        return None
    found = _HEADING_STYLE.match(label.strip())
    return int(found.group(1)) if found else None


def _sub(parent, tag: str, attrs: Optional[Dict[str, str]] = None, text: Optional[str] = None):
    """Create a child element of `parent`, in the tree so prefixes are reused."""
    element = etree.SubElement(parent, qn(tag))
    for name, value in (attrs or {}).items():
        element.set(qn(name), value)
    if text is not None:
        element.text = text
    return element


def _build_table_of_contents(body, title, max_level, entries, style_ids):
    """Build the table-of-contents content control at the end of `body`, and return it.

    The result is what Word itself writes: a ``w:sdt`` of the "Table of
    Contents" gallery holding a ``TOC`` complex field. The field's cached result
    -- the entries between ``separate`` and ``end`` -- is what a reader sees
    until the field is refreshed, which is why it is filled with the headings
    found rather than left empty.
    """
    control = etree.SubElement(body, qn("w:sdt"))
    properties = _sub(control, "w:sdtPr")
    gallery = _sub(properties, "w:docPartObj")
    _sub(gallery, "w:docPartGallery", {"w:val": _TOC_GALLERY})
    _sub(gallery, "w:docPartUnique")
    content = _sub(control, "w:sdtContent")

    if title:
        paragraph = _sub(content, "w:p")
        heading_style = next(
            (candidate for candidate in ("TOCHeading", "Heading1") if candidate in style_ids),
            None,
        )
        if heading_style:
            _sub(_sub(paragraph, "w:pPr"), "w:pStyle", {"w:val": heading_style})
        _sub(_sub(paragraph, "w:r"), "w:t", {"xml:space": "preserve"}, title)

    written = []
    for level, text in entries:
        paragraph = _sub(content, "w:p")
        entry_style = f"TOC{level}"
        if entry_style in style_ids:
            _sub(_sub(paragraph, "w:pPr"), "w:pStyle", {"w:val": entry_style})
        written.append((paragraph, text))

    first = written[0][0]
    _sub(_sub(first, "w:r"), "w:fldChar", {"w:fldCharType": "begin"})
    _sub(
        _sub(first, "w:r"),
        "w:instrText",
        {"xml:space": "preserve"},
        f' TOC \\o "1-{max_level}" \\h \\z \\u ',
    )
    _sub(_sub(first, "w:r"), "w:fldChar", {"w:fldCharType": "separate"})
    for paragraph, text in written:
        _sub(_sub(paragraph, "w:r"), "w:t", {"xml:space": "preserve"}, text)
    _sub(_sub(written[-1][0], "w:r"), "w:fldChar", {"w:fldCharType": "end"})
    return control


def _place_table_of_contents(body, control, style_names) -> None:
    """Move `control` to the head of the body, or just after its first heading."""
    blocks = [child for child in body if child.tag in BLOCK_TAGS and child is not control]
    if not blocks:
        body.insert(0, control)
        return
    first = blocks[0]
    if first.tag == _W_P and _heading_level(paragraph_style_label(first, style_names)) is not None:
        first.addnext(control)
    else:
        first.addprevious(control)


def _mark_fields_for_update(pkg) -> None:
    """Ask Word to refresh every field of the document when it opens it.

    A ``TOC`` field only ever shows its cached result until something updates
    it; ``w:updateFields`` in ``settings.xml`` is how a producer that cannot
    paginate -- which is every producer that is not Word -- gets real page
    numbers in front of the reader.

    ``ensure_part`` is idempotent in both directions: a document that already
    has a settings part keeps it untouched, and one that has none gets an empty
    ``w:settings`` related from the document part.
    """
    root = pkg.ensure_part(
        _SETTINGS_PARTNAME,
        _SETTINGS_CONTENT_TYPE,
        RT.SETTINGS,
        f'<w:settings {nsdecls("w")}/>',
    )

    element = root.find(qn("w:updateFields"))
    if element is None:
        element = etree.SubElement(root, qn("w:updateFields"))
        for index, child in enumerate(root):
            if child is element or not isinstance(child.tag, str):
                continue
            if etree.QName(child).localname in _SETTINGS_AFTER_UPDATE_FIELDS:
                root.insert(index, element)
                break
    element.set(qn("w:val"), "true")


async def add_table_of_contents(filename: str, title: str = "Table of Contents", max_level: int = 3) -> str:
    """Add a table of contents to a Word document based on heading styles.

    The table of contents is inserted as a real Word field -- a ``w:sdt`` of the
    "Table of Contents" gallery wrapping a ``TOC \\o "1-N" \\h \\z \\u``
    complex field -- at the head of the body, or just after the document's
    first heading. Nothing else in the document is touched: the tool used to
    rebuild the whole file inside a blank ``Document()``, which kept paragraph
    text and style names and dropped comments, footnotes, headers, images,
    hyperlinks, bookmarks, fields, tracked changes and table merges.

    Args:
        filename: Path to the Word document
        title: Optional title for the table of contents
        max_level: Maximum heading level to include (1-9)
    """
    filename = ensure_docx_extension(filename)

    if not os.path.exists(filename):
        return f"Document {filename} does not exist"

    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first."

    try:
        # Ensure max_level is within valid range
        max_level = max(1, min(int(max_level), 9))

        async with get_file_lock(filename):
            pkg = DocxPackage.open(filename)
            body = body_element(pkg)
            style_names = paragraph_style_names(styles_root(pkg))

            entries = []
            for element in body_paragraphs(body):
                level = _heading_level(paragraph_style_label(element, style_names))
                if level is not None and level <= max_level:
                    entries.append((level, visible_text(element)))

            if not entries:
                return f"No headings found in document {filename}. Table of contents not created."

            control = _build_table_of_contents(
                body, title, max_level, entries, set(style_names)
            )
            _place_table_of_contents(body, control, style_names)
            _mark_fields_for_update(pkg)
            pkg.save(filename)

        return f"Table of contents with {len(entries)} entries added to {filename}"
    except Exception as e:
        return f"Failed to add table of contents: {str(e)}"


def _carry_over_section_properties(paragraph) -> None:
    """Move a paragraph's ``w:sectPr`` onto the paragraph before it.

    A ``w:sectPr`` inside a ``w:pPr`` describes the section that *ends* at that
    paragraph -- it is the section break itself. Removing the paragraph without
    moving it takes the page size, the margins, the columns and the header and
    footer references of everything above it along, which is a page-layout loss
    no caller asked for.

    When the preceding paragraph already carries a ``w:sectPr`` of its own, its
    section ends there and the one being moved would cover no content at all, so
    it is left behind; the same reasoning applies when the deleted paragraph is
    the first of the body. In both cases the body's own final ``w:sectPr``, which
    is not attached to any paragraph, still describes the document's layout.
    """
    properties = paragraph.find(_W_PPR)
    if properties is None:
        return
    section = properties.find(_W_SECT_PR)
    if section is None:
        return
    previous = next(iter(paragraph.itersiblings(_W_P, preceding=True)), None)
    if previous is None or previous.find(f"{_W_PPR}/{_W_SECT_PR}") is not None:
        return
    properties.remove(section)
    attach_ppr_child(previous, section)


async def delete_paragraph(filename: str, paragraph_index: int) -> str:
    """Delete a paragraph from a document.

    A paragraph that carries a section break hands its ``w:sectPr`` over to the
    paragraph before it rather than taking the section with it; see
    :func:`_carry_over_section_properties`.

    Args:
        filename: Path to the Word document
        paragraph_index: Index of the paragraph to delete (0-based)
    """
    filename = ensure_docx_extension(filename)

    if not os.path.exists(filename):
        return f"Document {filename} does not exist"

    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first."

    try:
        async with get_file_lock(filename):
            pkg = DocxPackage.open(filename)
            body = body_element(pkg)
            paragraphs = body_paragraphs(body)

            # Validate paragraph index
            if paragraph_index < 0 or paragraph_index >= len(paragraphs):
                return f"Invalid paragraph index. Document has {len(paragraphs)} paragraphs (0-{len(paragraphs)-1})."

            paragraph = paragraphs[paragraph_index]
            _carry_over_section_properties(paragraph)
            body.remove(paragraph)

            pkg.save(filename)
        return f"Paragraph at index {paragraph_index} deleted successfully."
    except Exception as e:
        return f"Failed to delete paragraph: {str(e)}"


async def search_and_replace(filename: str, find_text: str, replace_text: str) -> str:
    """Search for text and replace all occurrences.
    
    Args:
        filename: Path to the Word document
        find_text: Text to search for
        replace_text: Text to replace with
    """
    filename = ensure_docx_extension(filename)
    
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    
    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first."
    
    try:
        async with get_file_lock(filename):
            pkg = DocxPackage.open(filename)

            # Perform find and replace on the OOXML engine: matches split across
            # runs, hyperlinks, tracked insertions, content controls, tables,
            # headers, footers and notes are all found and rewritten in place.
            report = replace_text_everywhere(pkg, find_text, replace_text)

            if report.replaced > 0:
                pkg.save(filename)
        if report.replaced > 0:
            message = f"Replaced {report.replaced} occurrence(s) of '{find_text}' with '{replace_text}'."
        else:
            message = f"No occurrences of '{find_text}' found."
        if report.skipped:
            message += f", {report.skipped} skipped (inside fields)"
        return message
    except Exception as e:
        return f"Failed to search and replace: {str(e)}"

async def insert_header_near_text_tool(filename: str, target_text: str = None, header_title: str = "", position: str = 'after', header_style: str = 'Heading 1', target_paragraph_index: int = None) -> str:
    """Insert a header (with specified style) before or after the target paragraph. Specify by text or paragraph index."""
    async with get_file_lock(filename):
        return insert_header_near_text(filename, target_text, header_title, position, header_style, target_paragraph_index)

async def insert_numbered_list_near_text_tool(filename: str, target_text: str = None, list_items: list = None, position: str = 'after', target_paragraph_index: int = None, bullet_type: str = 'bullet') -> str:
    """Insert a bulleted or numbered list before or after the target paragraph. Specify by text or paragraph index."""
    async with get_file_lock(filename):
        return insert_numbered_list_near_text(filename, target_text, list_items, position, target_paragraph_index, bullet_type)

async def insert_line_or_paragraph_near_text_tool(filename: str, target_text: str = None, line_text: str = "", position: str = 'after', line_style: str = None, target_paragraph_index: int = None) -> str:
    """Insert a new line or paragraph (with specified or matched style) before or after the target paragraph. Specify by text or paragraph index."""
    async with get_file_lock(filename):
        return insert_line_or_paragraph_near_text(filename, target_text, line_text, position, line_style, target_paragraph_index)

async def replace_paragraph_block_below_header_tool(filename: str, header_text: str, new_paragraphs: list, detect_block_end_fn=None) -> str:
    """Reemplaza el bloque de párrafos debajo de un encabezado, evitando modificar TOC."""
    async with get_file_lock(filename):
        return replace_paragraph_block_below_header(filename, header_text, new_paragraphs, detect_block_end_fn)

async def replace_block_between_manual_anchors_tool(filename: str, start_anchor_text: str, new_paragraphs: list, end_anchor_text: str = None, match_fn=None, new_paragraph_style: str = None) -> str:
    """Replace all content between start_anchor_text and end_anchor_text (or next logical header if not provided)."""
    async with get_file_lock(filename):
        return replace_block_between_manual_anchors(filename, start_anchor_text, new_paragraphs, end_anchor_text, match_fn, new_paragraph_style)
