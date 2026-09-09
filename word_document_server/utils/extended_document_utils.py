"""
Extended document utilities for Word Document Server.
"""
from typing import Dict, List, Any, Optional, Tuple
from docx import Document
from docx.oxml.ns import qn

from word_document_server.utils.document_utils import get_effective_text
from word_document_server.engine.find import find as engine_find
from word_document_server.engine.package import DocxPackage
from word_document_server.engine.textmodel import visible_text

#: Characters of the paragraph's own text reported as ``context`` before the
#: ellipsis, unchanged from the python-docx implementation this replaced.
CONTEXT_CHARS = 100


def _table_location(story_root, paragraph) -> Optional[str]:
    """Describe a paragraph sitting in a table cell, or return ``None``.

    The shape (``"Table 0, Row 1, Column 1"``) is the one the python-docx
    implementation produced. The table number counts every ``w:tbl`` of the
    story in document order, so a nested table has a number of its own instead
    of being invisible the way ``Document.tables`` made it. Row and column
    number the direct ``w:tr`` and ``w:tc`` children, i.e. the grid as written:
    a horizontally merged cell occupies one column here, not the several
    ``Row.cells`` expands it into.
    """
    cell = next(iter(paragraph.iterancestors(qn('w:tc'))), None)
    if cell is None:
        return None
    row = next(iter(cell.iterancestors(qn('w:tr'))), None)
    table = next(iter(cell.iterancestors(qn('w:tbl'))), None)
    if row is None or table is None:
        return None
    tables = list(story_root.iter(qn('w:tbl')))
    rows = table.findall(qn('w:tr'))
    cells = row.findall(qn('w:tc'))
    if table not in tables or row not in rows or cell not in cells:
        return None
    return f"Table {tables.index(table)}, Row {rows.index(row)}, Column {cells.index(cell)}"


def get_paragraph_text(doc_path: str, paragraph_index: int) -> Dict[str, Any]:
    """
    Get text from a specific paragraph in a Word document.
    
    Args:
        doc_path: Path to the Word document
        paragraph_index: Index of the paragraph to extract (0-based)
    
    Returns:
        Dictionary with paragraph text and metadata
    """
    import os
    if not os.path.exists(doc_path):
        return {"error": f"Document {doc_path} does not exist"}
    
    try:
        doc = Document(doc_path)
        
        # Check if paragraph index is valid
        if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
            return {"error": f"Invalid paragraph index: {paragraph_index}. Document has {len(doc.paragraphs)} paragraphs."}
        
        paragraph = doc.paragraphs[paragraph_index]
        
        return {
            "index": paragraph_index,
            "text": get_effective_text(paragraph),
            "style": paragraph.style.name if paragraph.style else "Normal",
            "is_heading": paragraph.style.name.startswith("Heading") if paragraph.style else False
        }
    except Exception as e:
        return {"error": f"Failed to get paragraph text: {str(e)}"}


def find_text(doc_path: str, text_to_find: str, match_case: bool = True, whole_word: bool = False) -> Dict[str, Any]:
    """
    Find all occurrences of specific text in a Word document.

    The search runs on the OOXML engine, on the *visible* text of every story of
    the package -- body, headers, footers, footnotes, endnotes -- so an
    occurrence is found wherever it reads: split across runs, inside a
    hyperlink, a tracked insertion, a content control or a table cell. Text
    hidden under a tracked deletion is not matched, because it is not what the
    document says.

    Each occurrence reports:

    ``story``
        the story it belongs to, named the way the package names it
        (``"document"``, ``"header1"``, ``"footnotes"``, ...).
    ``paragraph_index``
        the paragraph's V2 index in that story, or ``None`` when the paragraph
        is out of that index space (in a table cell or in a text box). A cell
        paragraph carries ``location`` instead.
    ``position``, ``start``, ``end``
        character offsets in the paragraph's visible text; ``position`` repeats
        ``start`` under its historical name.
    ``text``
        the matched text itself.
    ``context``
        the first :data:`CONTEXT_CHARS` characters of the paragraph, ellipsised.
    ``match_context``
        the text immediately around the match, centred on it.

    Args:
        doc_path: Path to the Word document
        text_to_find: Text to search for
        match_case: Whether to perform case-sensitive search
        whole_word: Whether to match whole words only

    Returns:
        Dictionary with search results
    """
    import os
    if not os.path.exists(doc_path):
        return {"error": f"Document {doc_path} does not exist"}

    if not text_to_find:
        return {"error": "Search text cannot be empty"}

    try:
        pkg = DocxPackage.open(doc_path)
        story_roots = dict(pkg.stories())
        matches = engine_find(
            pkg,
            text_to_find,
            case=match_case,
            whole_word=whole_word,
            stories=list(story_roots),
        )

        occurrences: List[Dict[str, Any]] = []
        for match in matches:
            paragraph_text = visible_text(match.paragraph)
            ellipsis = "..." if len(paragraph_text) > CONTEXT_CHARS else ""
            occurrence: Dict[str, Any] = {
                "story": match.story,
                "paragraph_index": match.index,
                "position": match.start,
                "start": match.start,
                "end": match.end,
                "text": match.text,
                "context": paragraph_text[:CONTEXT_CHARS] + ellipsis,
                "match_context": match.context,
            }
            location = _table_location(story_roots[match.story], match.paragraph)
            if location is not None:
                occurrence["location"] = location
            occurrences.append(occurrence)

        return {
            "query": text_to_find,
            "match_case": match_case,
            "whole_word": whole_word,
            "occurrences": occurrences,
            "total_count": len(occurrences),
        }
    except Exception as e:
        return {"error": f"Failed to search for text: {str(e)}"}


def get_highlighted_text(doc_path: str, color: str = None) -> Dict[str, Any]:
    """
    Extract all highlighted text from a Word document,
    including text inside table cells.

    Args:
        doc_path: Path to the Word document
        color: Optional highlight color filter (e.g. "yellow", "green", "cyan").
               If None, returns all highlighted text regardless of color.

    Returns:
        Dictionary with highlighted sections grouped by location
    """
    import os
    from docx.oxml.ns import qn

    if not os.path.exists(doc_path):
        return {"error": f"Document {doc_path} does not exist"}

    try:
        doc = Document(doc_path)
        results = {
            "filter_color": color,
            "highlights": [],
            "total_runs": 0,
            "summary": {}
        }

        def _check_paragraphs(paragraphs, location_prefix):
            for p_idx, para in enumerate(paragraphs):
                for run in para.runs:
                    hl = run._element.find(qn('w:rPr'))
                    if hl is not None:
                        hl_elem = hl.find(qn('w:highlight'))
                        if hl_elem is not None:
                            hl_color = hl_elem.get(qn('w:val'))
                            if color and hl_color != color:
                                continue
                            results["highlights"].append({
                                "location": f"{location_prefix}, Paragraph {p_idx}",
                                "color": hl_color,
                                "text": run.text
                            })
                            results["total_runs"] += 1
                            results["summary"][hl_color] = results["summary"].get(hl_color, 0) + 1

        # Top-level paragraphs
        _check_paragraphs(doc.paragraphs, "Body")

        # Table cells
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    _check_paragraphs(
                        cell.paragraphs,
                        f"Table {t_idx}, Row {r_idx}, Col {c_idx}"
                    )

        return results
    except Exception as e:
        return {"error": f"Failed to extract highlights: {str(e)}"}
